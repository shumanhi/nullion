from __future__ import annotations

import base64
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
import time
import urllib.error
from urllib.parse import parse_qs, urlparse
from zipfile import ZipFile
from zoneinfo import ZoneInfo

import pytest

from nullion.reminders import ReminderRecord
from nullion.runtime_store import RuntimeStore
from nullion.scheduler import ScheduleKind, ScheduledTask
from nullion import tools
from nullion.tools import (
    ToolInvocation,
    ToolRegistry,
    register_browser_plugin,
    register_calendar_plugin,
    register_cron_tools,
    register_email_plugin,
    register_media_plugin,
    register_market_quote_tool,
    register_reminder_tools,
    register_search_plugin,
    register_weather_forecast_tool,
    register_workspace_plugin,
)


def _invoke(tool_name: str, arguments: dict[str, object], *, principal_id: str = "web:operator") -> ToolInvocation:
    return ToolInvocation(
        invocation_id=f"inv-{tool_name}",
        tool_name=tool_name,
        principal_id=principal_id,
        arguments=arguments,
    )


def _write_png(path, color: str = "green"):
    from PIL import Image

    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (24, 24), color).save(path)
    return path


def _write_content_png(path, *, size: tuple[int, int] = (320, 220)):
    from PIL import Image

    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = size
    image = Image.new("RGB", size)
    pixels = image.load()
    for y in range(height):
        for x in range(width):
            pixels[x, y] = ((x * 7 + y * 3) % 256, (x * 2 + y * 11) % 256, (x * 13 + y * 5) % 256)
    image.save(path)
    return path


def _write_valid_renderer_pdf(path, *, image_paths=(), page_count=1):
    """Renderer doubles must produce readable PDFs with real embedded media."""
    import io
    from PIL import Image
    from pypdf import PdfReader, PdfWriter
    writer=PdfWriter()
    page=writer.add_blank_page(width=612,height=792)
    for source in image_paths:
        buffer=io.BytesIO()
        with Image.open(source) as image:
            image.convert('RGB').save(buffer,'PDF')
        page.merge_page(PdfReader(buffer).pages[0])
    for _ in range(1,page_count):
        writer.add_blank_page(width=612,height=792)
    writer.write(path)


def _write_browser_screenshot_artifact(tmp_path):
    return _write_png(tmp_path / "artifacts" / "screenshot-123456789abc.png", color="orange")


@dataclass
class RuntimeDouble:
    store: RuntimeStore


def test_media_provider_registers_svg_fallback_when_image_model_is_unconfigured(tmp_path, monkeypatch) -> None:
    from nullion.providers import resolve_plugin_provider_kwargs

    monkeypatch.delenv("NULLION_IMAGE_GENERATE_ENABLED", raising=False)
    monkeypatch.delenv("NULLION_IMAGE_GENERATE_PROVIDER", raising=False)
    monkeypatch.delenv("NULLION_IMAGE_GENERATE_MODEL", raising=False)
    monkeypatch.delenv("NULLION_IMAGE_GENERATE_COMMAND", raising=False)
    registry = ToolRegistry(plugin_registration_allowed=True, filesystem_allowed_roots=[tmp_path])

    register_media_plugin(
        registry,
        **resolve_plugin_provider_kwargs(plugin_name="media_plugin", provider_name="local_media_provider"),
    )
    output = tmp_path / "peacock.png"
    result = registry.invoke(
        _invoke("image_generate", {"prompt": "a peacock", "output_path": str(output), "size": "1024x1024"})
    )

    assert result.status == "completed"
    assert result.output["provider"] == "local_svg_fallback"
    fallback_path = tmp_path / "peacock.svg"
    assert result.output["path"] == str(fallback_path)
    assert result.output["raster_path"] == str(output)
    assert fallback_path.read_text(encoding="utf-8").startswith("<svg")
    assert output.is_file()


def test_browser_image_collect_registers_with_browser_web_scope_without_direct_fetch(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("NULLION_WEB_ACCESS_ENABLED", "true")
    monkeypatch.setenv("NULLION_FILE_ACCESS_ENABLED", "true")
    monkeypatch.setenv("NULLION_DIRECT_WEB_FETCH_ENABLED", "false")

    registry = tools.build_default_tool_registry(workspace_root=tmp_path)
    names = {definition["name"] for definition in registry.list_tool_definitions()}

    assert "browser_image_collect" in names
    assert "web_fetch" not in names


def test_media_provider_uses_svg_fallback_when_configured_image_model_fails(tmp_path, monkeypatch) -> None:
    from nullion import providers
    from nullion.providers import resolve_plugin_provider_kwargs

    monkeypatch.setenv("NULLION_IMAGE_GENERATE_ENABLED", "true")
    monkeypatch.setenv("NULLION_IMAGE_GENERATE_PROVIDER", "gemini")
    monkeypatch.setenv("NULLION_IMAGE_GENERATE_MODEL", "gemini-2.5-flash-image")
    monkeypatch.setenv("GEMINI_API_KEY", "gemini-key")
    monkeypatch.setattr(providers, "_model_image_generate", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("boom")))
    registry = ToolRegistry(plugin_registration_allowed=True, filesystem_allowed_roots=[tmp_path])

    register_media_plugin(
        registry,
        **resolve_plugin_provider_kwargs(plugin_name="media_plugin", provider_name="local_media_provider"),
    )
    output = tmp_path / "image.png"
    result = registry.invoke(_invoke("image_generate", {"prompt": "a peacock", "output_path": str(output)}))

    assert result.status == "completed"
    assert result.output["provider"] == "local_svg_fallback"
    assert result.output["fallback_error"] == "boom"
    assert (tmp_path / "image.svg").is_file()
    assert (tmp_path / "image.png").is_file()


def test_spreadsheet_create_fails_when_requested_images_are_skipped(tmp_path) -> None:
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)
    missing_image = tmp_path / "missing.png"
    output_path = tmp_path / "out.xlsx"

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name"],
                "rows": [{"Name": "Item", "image_path": str(missing_image)}],
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(missing_image)]
    assert not output_path.exists()


def test_file_write_rejects_placeholder_source_urls(tmp_path) -> None:
    handler = tools._build_file_write_handler(workspace_root=tmp_path)
    output_path = tmp_path / "products.csv"

    result = handler(
        _invoke(
            "file_write",
            {
                "path": str(output_path),
                "content": "Name,Link\nItem,https://shop.test/dp/example-product-1\n",
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "placeholder_source_urls"
    assert result.output["placeholder_urls"] == ["https://shop.test/dp/example-product-1"]
    assert not output_path.exists()


def test_generated_deliverables_use_workspace_artifact_root_for_runtime_home_paths(tmp_path, monkeypatch) -> None:
    nullion_home = tmp_path / "nullion-home"
    monkeypatch.setenv("NULLION_HOME", str(nullion_home))
    monkeypatch.setenv("NULLION_DATA_DIR", str(nullion_home))
    artifact_root = nullion_home / "workspaces" / "workspace_admin" / "artifacts"

    file_result = tools._build_file_write_handler(workspace_root=nullion_home)(
        _invoke(
            "file_write",
            {
                "path": str(nullion_home / "stage_values.csv"),
                "content": "Region,Value\nNorth,12\n",
            },
        )
    )
    assert file_result.status == "completed"
    assert Path(str(file_result.output["path"])) == artifact_root / "stage_values.csv"
    assert not (nullion_home / "stage_values.csv").exists()

    spreadsheet_result = tools._build_spreadsheet_create_handler(workspace_root=nullion_home)(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(nullion_home / "stage_values.xlsx"),
                "columns": ["Region", "Value"],
                "rows": [{"Region": "North", "Value": 12}],
            },
        )
    )
    assert spreadsheet_result.status == "completed"
    assert Path(str(spreadsheet_result.output["path"])) == artifact_root / "stage_values.xlsx"
    assert not (nullion_home / "stage_values.xlsx").exists()

    pdf_result = tools._build_pdf_create_handler(workspace_root=nullion_home)(
        _invoke(
            "pdf_create",
            {
                "output_path": str(nullion_home / "stage_values.pdf"),
                "text_pages": ["Stage values summary"],
            },
        )
    )
    assert pdf_result.status == "completed"
    assert Path(str(pdf_result.output["path"])) == artifact_root / "stage_values.pdf"
    assert not (nullion_home / "stage_values.pdf").exists()


def test_file_search_prioritizes_specific_workspace_roots_before_broad_root(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("NULLION_FILE_SEARCH_MAX_ENTRIES", "20")
    workspace_root = tmp_path / "runtime-home"
    specific_root = workspace_root / "workspaces" / "workspace_admin"
    artifact = specific_root / "artifacts" / "stage_file_search_delivery_guard.pdf"
    artifact.parent.mkdir(parents=True)
    artifact.write_bytes(b"%PDF-1.4\n")
    for index in range(30):
        noise = workspace_root / f"aaa-noise-{index:02d}"
        noise.mkdir(parents=True)
        (noise / "unrelated.txt").write_text("noise", encoding="utf-8")

    registry = ToolRegistry(plugin_registration_allowed=True, filesystem_allowed_roots=[workspace_root])
    register_workspace_plugin(registry, workspace_root=workspace_root, allowed_roots=(specific_root,))

    result = registry.invoke(
        _invoke(
            "file_search",
            {
                "pattern": artifact.name,
                "limit": 5,
            },
        )
    )

    assert result.status == "completed"
    assert result.output["matches"] == [str(artifact.resolve())]
    assert result.output["match_details"] == [
        {
            "path": str(artifact.resolve()),
            "match_type": "filename",
        }
    ]


def test_file_write_rejects_truncated_html_data_images(tmp_path) -> None:
    handler = tools._build_file_write_handler(workspace_root=tmp_path)
    output_path = tmp_path / "products.html"

    result = handler(
        _invoke(
            "file_write",
            {
                "path": str(output_path),
                "content": '<html><body><img src="data:image/jpeg;base64,/9j/4AAQSkZJRgABAQIAJQAlAAD"></body></html>',
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "invalid_embedded_html_images"
    assert result.output["invalid_images"] == ["image 1: invalid base64 data URI"]
    assert not output_path.exists()


def test_file_write_rejects_incomplete_html_document(tmp_path) -> None:
    handler = tools._build_file_write_handler(workspace_root=tmp_path)
    output_path = tmp_path / "truncated.html"

    result = handler(
        _invoke(
            "file_write",
            {
                "path": str(output_path),
                "content": "<html><body><main>unfinished",
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "incomplete_html_document"
    assert result.output["missing_document_parts"] == ["</body>", "</html>"]
    assert not output_path.exists()


def test_file_write_rejects_html_missing_typed_required_content_tokens(tmp_path) -> None:
    handler = tools._build_file_write_handler(workspace_root=tmp_path)
    output_path = tmp_path / "partial.html"

    result = handler(
        ToolInvocation(
            invocation_id="inv-required-content",
            tool_name="file_write",
            principal_id="web:operator",
            arguments={
                "path": str(output_path),
                "content": (
                    '<html><head><meta name="NVDA"></head><body>'
                    "<section>MSFT market card</section></body></html>"
                ),
            },
            flow_context={
                "requires_artifact_delivery": True,
                "required_artifact_extensions": [".html"],
                "required_artifact_content_tokens": ["MSFT", "NVDA"],
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "artifact_required_content_missing"
    assert result.output["missing_required_content_tokens"] == ["NVDA"]
    assert not output_path.exists()


def test_file_write_inlines_local_html_image_artifacts(tmp_path) -> None:
    from PIL import Image

    handler = tools._build_file_write_handler(workspace_root=tmp_path)
    image_path = tmp_path / "artifacts" / "product.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (128, 128))
    image.putdata(
        [
            (
                (x * 37 + y * 17) % 256,
                (x * y * 11 + 23) % 256,
                ((x ^ y) * 29 + x * 5) % 256,
            )
            for y in range(128)
            for x in range(128)
        ]
    )
    image.save(image_path)
    output_path = tmp_path / "products.html"

    result = handler(
        _invoke(
            "file_write",
            {
                "path": str(output_path),
                "content": '<html><body><img src="artifacts/product.png" alt="Product"></body></html>',
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_html_images"][0]["path"] == str(image_path)
    content = output_path.read_text(encoding="utf-8")
    assert "artifacts/product.png" not in content
    assert 'src="data:image/png;base64,' in content
    payload = content.split('src="data:image/png;base64,', 1)[1].split('"', 1)[0]
    assert base64.b64decode(payload, validate=True) == image_path.read_bytes()


def test_file_write_can_preserve_linked_html_image_references(tmp_path) -> None:
    image_path = _write_content_png(tmp_path / "assets" / "source.png")
    handler = tools._build_file_write_handler(workspace_root=tmp_path)
    output_path = tmp_path / "linked.html"

    result = handler(
        _invoke(
            "file_write",
            {
                "path": str(output_path),
                "content": '<html><body><img src="assets/source.png" alt="Source"></body></html>',
                "inline_local_html_images": False,
                "disallow_html_data_images": True,
            },
        )
    )

    assert result.status == "completed"
    assert result.output["html_image_mode"] == "linked"
    content = output_path.read_text(encoding="utf-8")
    assert 'src="assets/source.png"' in content
    assert "data:image" not in content
    assert image_path.exists()


def test_file_write_rejects_data_images_for_linked_html(tmp_path) -> None:
    handler = tools._build_file_write_handler(workspace_root=tmp_path)
    output_path = tmp_path / "linked.html"

    result = handler(
        _invoke(
            "file_write",
            {
                "path": str(output_path),
                "content": '<html><body><img src="data:image/svg+xml;base64,PHN2Zy8+" alt="Inline"></body></html>',
                "inline_local_html_images": False,
                "disallow_html_data_images": True,
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "html_data_images_disallowed"
    assert not output_path.exists()


def test_file_write_requires_raster_data_image_for_html_embedded_media_contract(tmp_path) -> None:
    handler = tools._build_file_write_handler(workspace_root=tmp_path)
    output_path = tmp_path / "visual.html"

    result = handler(
        ToolInvocation(
            invocation_id="inv-file-write",
            tool_name="file_write",
            principal_id="web:operator",
            arguments={
                "path": str(output_path),
                "content": '<html><body><img src="data:image/svg+xml;base64,PHN2Zy8+" alt="Icon"></body></html>',
            },
            flow_context={"required_embedded_media_extensions": [".html"]},
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "html_embedded_raster_image_required"
    assert not output_path.exists()


def test_spreadsheet_create_requires_embedded_media_when_turn_contract_demands_it(tmp_path) -> None:
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)
    output_path = tmp_path / "visual.xlsx"

    result = handler(
        ToolInvocation(
            invocation_id="inv-spreadsheet-create",
            tool_name="spreadsheet_create",
            principal_id="web:operator",
            arguments={
                "output_path": str(output_path),
                "columns": ["Name", "Score"],
                "rows": [{"Name": "Alpha", "Score": 9}],
            },
            flow_context={"required_embedded_media_extensions": [".xlsx"]},
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "artifact_media_required_by_turn_contract"
    assert result.output["required_arguments"] == ["image_paths", "screenshot_paths"]
    assert not output_path.exists()


def test_spreadsheet_create_overwrites_existing_media_when_media_is_not_required(tmp_path) -> None:
    from openpyxl import load_workbook

    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)
    image_path = _write_content_png(tmp_path / "source.png")
    output_path = tmp_path / "visual.xlsx"

    first = handler(
        ToolInvocation(
            invocation_id="inv-spreadsheet-create-media",
            tool_name="spreadsheet_create",
            principal_id="web:operator",
            arguments={
                "output_path": str(output_path),
                "columns": ["Name", "Score"],
                "rows": [{"Name": "Alpha", "Score": 9}],
                "image_paths": [str(image_path)],
            },
        )
    )

    assert first.status == "completed"
    with ZipFile(output_path) as package:
        first_media = [name for name in package.namelist() if name.startswith("xl/media/")]
    assert first_media

    second = handler(
        ToolInvocation(
            invocation_id="inv-spreadsheet-create-plain",
            tool_name="spreadsheet_create",
            principal_id="web:operator",
            arguments={
                "output_path": str(output_path),
                "columns": ["Name", "Score"],
                "rows": [{"Name": "Beta", "Score": 10}],
            },
        )
    )

    assert second.status == "completed"
    assert second.output["artifact_paths"] == [str(output_path)]
    workbook = load_workbook(output_path, data_only=False)
    assert workbook.active["A2"].value == "Beta"
    assert workbook.active["B2"].value == 10
    with ZipFile(output_path) as package:
        second_media = [name for name in package.namelist() if name.startswith("xl/media/")]
    assert second_media == []


def test_spreadsheet_create_rejects_stale_media_for_structured_visual_regeneration(tmp_path) -> None:
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)
    image_path = _write_content_png(tmp_path / "source.png")
    output_path = tmp_path / "visual-regeneration.xlsx"
    first = handler(
        ToolInvocation(
            invocation_id="inv-spreadsheet-create-media-first",
            tool_name="spreadsheet_create",
            principal_id="web:operator",
            arguments={
                "output_path": str(output_path),
                "columns": ["Name", "Score"],
                "rows": [{"Name": "Old", "Score": 1}],
                "image_paths": [str(image_path)],
            },
        )
    )
    assert first.status == "completed"

    regenerated = handler(
        ToolInvocation(
            invocation_id="inv-spreadsheet-create-media-regenerated",
            tool_name="spreadsheet_create",
            principal_id="web:operator",
            arguments={
                "output_path": str(output_path),
                "columns": ["Name", "Score"],
                "rows": [{"Name": "New", "Score": 2}],
                "charts": [
                    {
                        "type": "bar",
                        "categories_column": "Name",
                        "values_column": "Score",
                    }
                ],
            },
            flow_context={"required_embedded_media_extensions": [".xlsx"]},
        )
    )

    assert regenerated.status == "failed"
    assert regenerated.output["reason"] == "stale_embedded_media_cannot_satisfy_regeneration"
    assert regenerated.output["required_arguments"] == ["image_paths", "screenshot_paths"]


def test_spreadsheet_create_builds_requested_multi_sheet_workbook(tmp_path) -> None:
    from openpyxl import load_workbook

    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)
    output_path = tmp_path / "regional-summary.xlsx"

    result = handler(
        ToolInvocation(
            invocation_id="inv-spreadsheet-create-multi-sheet",
            tool_name="spreadsheet_create",
            principal_id="web:operator",
            arguments={
                "output_path": str(output_path),
                "expected_rows": 5,
                "sheets": [
                    {
                        "sheet_name": "Values",
                        "columns": ["Region", "Value"],
                        "rows": [
                            {"Region": "North", "Value": 12},
                            {"Region": "South", "Value": 19},
                            {"Region": "East", "Value": 7},
                            {"Region": "West", "Value": 15},
                        ],
                    },
                    {
                        "sheet_name": "Summary",
                        "columns": ["Metric", "Value"],
                        "rows": [{"Metric": "Total", "Value": 53}],
                    },
                ],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["sheet_names"] == ["Values", "Summary"]
    workbook = load_workbook(output_path, read_only=True, data_only=False)
    assert workbook.sheetnames == ["Values", "Summary"]
    assert list(workbook["Values"].values) == [
        ("Region", "Value"),
        ("North", 12),
        ("South", 19),
        ("East", 7),
        ("West", 15),
    ]
    assert list(workbook["Summary"].values) == [("Metric", "Value"), ("Total", 53)]
    workbook.close()


def test_file_write_allows_linked_html_inside_mixed_embedded_media_contract(tmp_path) -> None:
    image_path = _write_content_png(tmp_path / "assets" / "source.png")
    handler = tools._build_file_write_handler(workspace_root=tmp_path)
    output_path = tmp_path / "linked.html"

    result = handler(
        ToolInvocation(
            invocation_id="inv-file-write",
            tool_name="file_write",
            principal_id="web:operator",
            arguments={
                "path": str(output_path),
                "content": '<html><body><img src="assets/source.png" alt="Source"></body></html>',
                "inline_local_html_images": False,
                "disallow_html_data_images": True,
            },
            flow_context={"required_embedded_media_extensions": [".html"]},
        )
    )

    assert result.status == "completed"
    assert result.output["html_image_mode"] == "linked"
    assert image_path.exists()
    assert output_path.read_text(encoding="utf-8").count("data:image/") == 0


def test_file_write_allows_svg_only_html_without_embedded_media_contract(tmp_path) -> None:
    handler = tools._build_file_write_handler(workspace_root=tmp_path)
    output_path = tmp_path / "plain.html"

    result = handler(
        _invoke(
            "file_write",
            {
                "path": str(output_path),
                "content": '<html><body><img src="data:image/svg+xml;base64,PHN2Zy8+" alt="Icon"></body></html>',
            },
        )
    )

    assert result.status == "completed"
    assert output_path.exists()


def test_document_create_can_write_self_contained_html_output(tmp_path) -> None:
    image_path = _write_content_png(tmp_path / "assets" / "source.png")
    handler = tools._build_document_create_handler(workspace_root=tmp_path)
    output_path = tmp_path / "report.html"

    result = handler(
        _invoke(
            "document_create",
            {
                "output_path": str(output_path),
                "title": "Museum report",
                "paragraphs": ["Source URL: https://museum.example/item"],
                "image_paths": [str(image_path)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["html_image_mode"] == "self_contained"
    content = output_path.read_text(encoding="utf-8")
    assert "Museum report" in content
    assert "https://museum.example/item" in content
    assert "data:image/png;base64," in content


def test_document_create_rejects_html_missing_typed_required_content_tokens(tmp_path) -> None:
    handler = tools._build_document_create_handler(workspace_root=tmp_path)
    output_path = tmp_path / "partial-report.html"

    result = handler(
        ToolInvocation(
            invocation_id="inv-document-required-content",
            tool_name="document_create",
            principal_id="web:operator",
            arguments={
                "output_path": str(output_path),
                "title": "Market report",
                "sections": [{"heading": "MSFT", "body": "Market card"}],
            },
            flow_context={
                "requires_artifact_delivery": True,
                "required_artifact_extensions": [".html"],
                "required_artifact_content_tokens": ["MSFT", "NVDA"],
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "artifact_required_content_missing"
    assert result.output["missing_required_content_tokens"] == ["NVDA"]
    assert not output_path.exists()


def test_document_create_can_write_linked_html_output(tmp_path) -> None:
    image_path = _write_content_png(tmp_path / "assets" / "source.png")
    handler = tools._build_document_create_handler(workspace_root=tmp_path)
    output_path = tmp_path / "report.html"

    result = handler(
        _invoke(
            "document_create",
            {
                "output_path": str(output_path),
                "title": "Museum report",
                "sections": [{"heading": "Images", "image_paths": [str(image_path)]}],
                "inline_local_html_images": False,
                "disallow_html_data_images": True,
            },
        )
    )

    assert result.status == "completed"
    assert result.output["html_image_mode"] == "linked"
    content = output_path.read_text(encoding="utf-8")
    assert "data:image" not in content
    assert 'src="assets/source.png"' in content


def test_spreadsheet_create_rejects_placeholder_image_objects_in_case_insensitive_image_column(tmp_path) -> None:
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)
    output_path = tmp_path / "placeholder.xlsx"

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name", "Link", "Image"],
                "rows": [
                    {
                        "Name": "Item",
                        "Link": "https://example.com/item",
                        "Image": {"generated_placeholder": "Row 1"},
                    }
                ],
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["embedded_images"] == []
    assert result.output["skipped_images"] == ['{"generated_placeholder": "Row 1"}']
    assert not output_path.exists()


def test_spreadsheet_create_embeds_svg_companion_raster(tmp_path) -> None:
    svg_path = tmp_path / "fallback.svg"
    png_path = tmp_path / "fallback.png"
    output_path = tmp_path / "with-image.xlsx"
    svg_path.write_text("<svg xmlns='http://www.w3.org/2000/svg'></svg>", encoding="utf-8")
    _write_content_png(png_path)
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name"],
                "rows": [{"Name": "Item", "image_path": str(svg_path)}],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_images"] == [str(png_path)]
    assert result.output["skipped_images"] == []
    with ZipFile(output_path) as archive:
        assert any(name.startswith("xl/media/") for name in archive.namelist())


def test_spreadsheet_create_rasterizes_svg_without_companion(tmp_path) -> None:
    svg_path = tmp_path / "generated-marker.svg"
    png_path = tmp_path / "generated-marker.png"
    output_path = tmp_path / "with-generated-marker.xlsx"
    svg_path.write_text(
        (
            "<svg xmlns='http://www.w3.org/2000/svg' width='320' height='180'>"
            "<rect width='320' height='180' fill='#dbeafe'/>"
            "<text x='24' y='90'>Cafe marker</text>"
            "</svg>"
        ),
        encoding="utf-8",
    )
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name"],
                "rows": [{"Name": "Item", "image_path": str(svg_path)}],
            },
        )
    )

    assert result.status == "completed"
    assert png_path.exists()
    assert result.output["embedded_images"] == [str(png_path)]
    assert result.output["skipped_images"] == []
    with ZipFile(output_path) as archive:
        assert any(name.startswith("xl/media/") for name in archive.namelist())


def test_spreadsheet_create_treats_capital_image_column_as_embed_path(tmp_path) -> None:
    image_path = _write_content_png(tmp_path / "image.png")
    output_path = tmp_path / "capital-image.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name", "Link", "Image"],
                "rows": [
                    {
                        "Name": "Item",
                        "Link": "https://example.com/item",
                        "Image": str(image_path),
                    }
                ],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["columns"] == 3
    assert result.output["embedded_images"] == [str(image_path)]
    with ZipFile(output_path) as archive:
        assert any(name.startswith("xl/media/") for name in archive.namelist())


def test_spreadsheet_create_embeds_local_image_path_cells_without_reserved_header(tmp_path) -> None:
    image_path = _write_content_png(tmp_path / "image.png")
    output_path = tmp_path / "image-file-column.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Product", "Image File"],
                "rows": [{"Product": "QA item", "Image File": str(image_path)}],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_images"] == [str(image_path)]
    with ZipFile(output_path) as archive:
        assert any(name.startswith("xl/media/") for name in archive.namelist())


def test_spreadsheet_create_rejects_remote_image_url_cells_before_plain_file_save(tmp_path) -> None:
    output_path = tmp_path / "remote-image-url.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Product", "Image URL"],
                "rows": [
                    {
                        "Product": "QA item",
                        "Image URL": "https://cdn.example.test/products/qa-item.jpg",
                    }
                ],
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "remote_image_paths_not_supported"
    assert result.output["remote_image_urls"] == ["https://cdn.example.test/products/qa-item.jpg"]
    assert "local artifact file paths" in (result.error or "")
    assert not output_path.exists()


def test_spreadsheet_create_rejects_structured_expected_row_count_mismatch(tmp_path) -> None:
    output_path = tmp_path / "wrong-row-count.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Product", "Price"],
                "expected_rows": 1,
                "rows": [
                    {"Product": "First item", "Price": "$12.00"},
                    {"Product": "Second item", "Price": "$13.00"},
                ],
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "row_count_mismatch"
    assert result.output["rows"] == 2
    assert result.output["expected_rows"] == 1
    assert not output_path.exists()


def test_spreadsheet_create_allows_reused_images_for_distinct_linked_rows(tmp_path) -> None:
    image_path = _write_content_png(tmp_path / "same-product-image.png")
    output_path = tmp_path / "duplicate-row-images.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Product", "Listing Link"],
                "rows": [
                    {"Product": "First item", "Listing Link": "https://example.com/item/1"},
                    {"Product": "Second item", "Listing Link": "https://example.com/item/2"},
                ],
                "image_paths": [str(image_path), str(image_path)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_images"] == [str(image_path), str(image_path)]
    assert output_path.exists()
    with ZipFile(output_path) as archive:
        assert len([name for name in archive.namelist() if name.startswith("xl/media/")]) == 2


def test_spreadsheet_create_rejects_aggregate_per_row_links(tmp_path) -> None:
    output_path = tmp_path / "aggregate-links.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name", "Link"],
                "rows": [
                    {"Name": "Item A", "Link": "https://example.com/search?q=a"},
                    {"Name": "Item B", "Link": "https://example.com/s?q=b"},
                ],
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "non_row_specific_links"
    assert "https://example.com/search?q=a" in result.output["aggregate_row_links"]
    assert "https://example.com/s?q=b" in result.output["aggregate_row_links"]


def test_spreadsheet_create_allows_site_homepage_rows(tmp_path) -> None:
    from openpyxl import load_workbook

    output_path = tmp_path / "site-homepages.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name", "Link"],
                "rows": [
                    {"Name": "Museum A", "Link": "https://museum-a.test"},
                    {"Name": "Museum B", "Link": "https://museum-b.test/"},
                ],
            },
        )
    )

    assert result.status == "completed"
    workbook = load_workbook(output_path)
    assert workbook.active["B2"].hyperlink.target == "https://museum-a.test"
    assert workbook.active["B3"].hyperlink.target == "https://museum-b.test/"


def test_spreadsheet_create_can_write_csv_output(tmp_path) -> None:
    output_path = tmp_path / "sources.csv"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name", "Link"],
                "rows": [
                    {"Name": "Museum A", "Link": "https://museum-a.test"},
                    {"Name": "Museum B", "Link": {"url": "https://museum-b.test/item", "title": "Object"}},
                ],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["format"] == "csv"
    assert output_path.read_text(encoding="utf-8").splitlines() == [
        "Name,Link",
        "Museum A,https://museum-a.test",
        "Museum B,https://museum-b.test/item",
    ]


def test_spreadsheet_create_canonicalizes_nested_redirect_row_links(tmp_path) -> None:
    from openpyxl import load_workbook

    output_path = tmp_path / "canonical-links.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)
    redirect_link = "https://example.com/click?campaign=ad&url=%2Fitem%2F123%3Ftracking%3Dabc#ignored"

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name", "Link"],
                "rows": [{"Name": "Item A", "Link": redirect_link}],
            },
        )
    )

    assert result.status == "completed"
    workbook = load_workbook(output_path)
    cell = workbook.active["B2"]
    assert cell.value == "https://example.com/item/123"
    assert cell.hyperlink.target == "https://example.com/item/123"


def test_spreadsheet_create_canonicalizes_path_embedded_row_links(tmp_path) -> None:
    from openpyxl import load_workbook

    output_path = tmp_path / "canonical-path-links.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)
    redirect_link = "https://redirect.example/click/https://shop.example/item/123?tracking=abc#ignored"

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name", "Link"],
                "rows": [{"Name": "Item A", "Link": redirect_link}],
            },
        )
    )

    assert result.status == "completed"
    workbook = load_workbook(output_path)
    cell = workbook.active["B2"]
    assert cell.value == "https://shop.example/item/123"
    assert cell.hyperlink.target == "https://shop.example/item/123"


def test_spreadsheet_create_rejects_incomplete_mixed_media_link_price_rows(tmp_path) -> None:
    image_path = _write_png(tmp_path / "image.png", color="teal")
    output_path = tmp_path / "incomplete-rows.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Product", "Listing Price", "Listing Link", "Image File"],
                "rows": [
                    {
                        "Product": "Complete row",
                        "Listing Price": "$12.34",
                        "Listing Link": "https://example.com/item/1",
                        "Image File": str(image_path),
                    },
                    {
                        "Product": "Incomplete row",
                        "Listing Price": "-",
                        "Listing Link": None,
                        "image_path": None,
                    },
                ],
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "incomplete_required_row_values"
    assert result.output["missing_image_rows"] == [3]
    assert result.output["missing_link_cells"] == [{"row": 3, "column": "Listing Link"}]
    assert "missing_price_cells" not in result.output
    assert not output_path.exists()


def test_spreadsheet_create_allows_summary_rows_without_media_contract(tmp_path) -> None:
    image_paths = [
            _write_content_png(tmp_path / "first.png"),
            _write_content_png(tmp_path / "second.png"),
            _write_content_png(tmp_path / "third.png"),
    ]
    output_path = tmp_path / "products-with-summary.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Product", "Listing Price", "Listing Link", "Margin"],
                "rows": [
                    {
                        "Product": "First item",
                        "Listing Price": "$12.34",
                        "Listing Link": "https://example.com/item/1",
                        "Margin": "42%",
                    },
                    {
                        "Product": "Second item",
                        "Listing Price": "$23.45",
                        "Listing Link": "https://example.com/item/2",
                        "Margin": "38%",
                    },
                    {
                        "Product": "Third item",
                        "Listing Price": "$34.56",
                        "Listing Link": "https://example.com/item/3",
                        "Margin": "40%",
                    },
                    {
                        "Product": "Portfolio total",
                        "Listing Price": "=SUM(B2:B4)",
                        "Listing Link": None,
                        "Margin": "=AVERAGE(D2:D4)",
                    },
                ],
                "image_paths": [str(path) for path in image_paths],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_images"] == [str(path) for path in image_paths]
    assert output_path.exists()
    with ZipFile(output_path) as archive:
        assert len([name for name in archive.namelist() if name.startswith("xl/media/")]) == 3


def test_spreadsheet_create_embeds_chart_from_structured_columns(tmp_path) -> None:
    image_path = _write_content_png(tmp_path / "item.png")
    output_path = tmp_path / "with-chart.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Product", "Score", "Listing Link"],
                "rows": [
                    {"Product": "First", "Score": 91, "Listing Link": "https://example.com/item/1"},
                    {"Product": "Second", "Score": 84, "Listing Link": "https://example.com/item/2"},
                ],
                "image_paths": [str(image_path), str(image_path)],
                "charts": [
                    {
                        "type": "bar",
                        "title": "Score by product",
                        "categories_column": "Product",
                        "values_column": "Score",
                        "anchor": "G2",
                    }
                ],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_charts"][0]["type"] == "bar"
    with ZipFile(output_path) as archive:
        assert any(name.startswith("xl/charts/chart") for name in archive.namelist())


def test_spreadsheet_create_applies_structured_conditional_formatting(tmp_path) -> None:
    output_path = tmp_path / "with-conditional-format.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Location", "Score", "Payback"],
                "rows": [
                    {"Location": "Gallery Lane", "Score": 86.5, "Payback": "=ROUND(12000/372,2)"},
                    {"Location": "Market Square", "Score": 81.7, "Payback": "=ROUND(9500/392,2)"},
                ],
                "formulas_required": True,
                "conditional_formats": [
                    {"type": "data_bar", "column": "Score", "color": "8E6F63"},
                    {"type": "color_scale", "column": "Payback"},
                ],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["conditional_formats"] == [
        {"type": "data_bar", "column": "Score", "color": "8E6F63", "range": "B2:B3"},
        {"type": "color_scale", "column": "Payback", "color": "", "range": "C2:C3"},
    ]
    with ZipFile(output_path) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<conditionalFormatting" in sheet_xml
    assert "<f>ROUND(12000/372,2)</f>" in sheet_xml or "<ns0:f>ROUND(12000/372,2)</ns0:f>" in sheet_xml


def test_spreadsheet_create_accepts_formula_chart_and_format_on_non_primary_sheet(tmp_path) -> None:
    output_path = tmp_path / "multi-sheet-markets.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "sheets": [
                    {
                        "sheet_name": "Briefing",
                        "columns": ["Topic", "Status"],
                        "rows": [{"Topic": "Calendar", "Status": "Unavailable"}],
                    },
                    {
                        "sheet_name": "Markets",
                        "columns": ["Symbol", "Price", "Previous", "Change"],
                        "rows": [
                            {
                                "Symbol": "SPY",
                                "Price": 700.0,
                                "Previous": 695.0,
                                "Change": "=ROUND((B2-C2)/C2,4)",
                            },
                            {
                                "Symbol": "QQQ",
                                "Price": 620.0,
                                "Previous": 615.0,
                                "Change": "=ROUND((B3-C3)/C3,4)",
                            },
                        ],
                    },
                    {
                        "sheet_name": "Schedule",
                        "columns": ["Time", "Item"],
                        "rows": [{"Time": "14:00", "Item": "Museum"}],
                    },
                ],
                "expected_rows": 4,
                "formulas_required": True,
                "charts": [
                    {
                        "sheet_name": "Markets",
                        "type": "bar",
                        "categories_column": "Symbol",
                        "values_column": "Price",
                    }
                ],
                "conditional_formats": [
                    {
                        "sheet_name": "Markets",
                        "type": "color_scale",
                        "column": "Change",
                    }
                ],
            },
        )
    )

    assert result.status == "completed", result.error
    assert result.output["sheet_names"] == ["Briefing", "Markets", "Schedule"]
    assert result.output["embedded_charts"][0]["sheet_name"] == "Markets"
    assert result.output["conditional_formats"][0]["sheet_name"] == "Markets"
    with ZipFile(output_path) as archive:
        assert any(name.startswith("xl/charts/chart") for name in archive.namelist())
        markets_xml = archive.read("xl/worksheets/sheet2.xml").decode("utf-8")
    assert "<conditionalFormatting" in markets_xml
    assert "ROUND((B2-C2)/C2,4)" in markets_xml


def test_spreadsheet_create_rejects_missing_required_formula_cells(tmp_path) -> None:
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(tmp_path / "missing-formulas.xlsx"),
                "columns": ["Location", "Score"],
                "rows": [{"Location": "Gallery Lane", "Score": 86.5}],
                "formulas_required": True,
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "missing_required_formulas"
    assert "selected columns" in str(result.error)

    omitted_formula_result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(tmp_path / "omitted-formula.xlsx"),
                "columns": ["Location", "Score"],
                "rows": [{"Location": "Gallery Lane", "Score": 86.5, "Formula Note": "=B2+1"}],
                "formulas_required": True,
            },
        )
    )

    assert omitted_formula_result.status == "failed"
    assert omitted_formula_result.output["reason"] == "missing_required_formulas"


def test_spreadsheet_create_writes_cached_formula_values_for_previewers(tmp_path) -> None:
    output_path = tmp_path / "formula-preview.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Activity", "Read", "Draw", "Build", "Score"],
                "rows": [
                    {"Activity": "Rocket", "Read": 2, "Draw": 3, "Build": 4, "Score": "=SUM(B2:D2)"},
                    {"Activity": "Stars", "Read": 1, "Draw": 5, "Build": 2, "Score": "=AVERAGE(B3:D3)"},
                ],
            },
        )
    )

    assert result.status == "completed"
    with ZipFile(output_path) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<ns0:f>SUM(B2:D2)</ns0:f><ns0:v>9</ns0:v>" in sheet_xml
    assert "<ns0:f>AVERAGE(B3:D3)</ns0:f><ns0:v>2.66666666667</ns0:v>" in sheet_xml


def test_spreadsheet_create_does_not_infer_required_price_columns_from_header_text(tmp_path) -> None:
    image_path = _write_content_png(tmp_path / "image.png")
    output_path = tmp_path / "zero-prices.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Product", "Listing Price", "Listing Link", "Image File"],
                "rows": [
                    {
                        "Product": "Placeholder price row",
                        "Listing Price": "$0.00",
                        "Listing Link": "https://example.com/item/1",
                        "Image File": str(image_path),
                    },
                ],
            },
        )
    )

    assert result.status == "completed"
    assert output_path.exists()


def test_browser_image_collect_materializes_embeddable_spreadsheet_images(tmp_path, monkeypatch) -> None:
    source_image = _write_content_png(tmp_path / "source.png")
    image_data_url = "data:image/png;base64," + base64.b64encode(source_image.read_bytes()).decode("ascii")
    monkeypatch.setenv("NULLION_WORKSPACE_STORAGE_ROOT", str(tmp_path / "workspaces"))

    collect = tools._build_browser_image_collect_handler()
    collected = collect(
        _invoke(
            "browser_image_collect",
            {
                "image_urls": [image_data_url],
                "output_stem": "qa-product-image",
                "max_images": 1,
            },
            principal_id="telegram:1",
        )
    )

    assert collected.status == "completed"
    image_paths = collected.output["image_paths"]
    assert len(image_paths) == 1
    assert Path(image_paths[0]).is_file()
    descriptor = collected.output["artifact_descriptors"][0]
    assert descriptor["path"] == image_paths[0]
    assert descriptor["role"] == "source"
    assert descriptor["kind"] == "browser_image"
    assert descriptor["media_type"] == "image/png"

    output_path = tmp_path / "product-images.xlsx"
    spreadsheet = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)
    result = spreadsheet(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Product", "Image"],
                "rows": [{"Product": "QA item", "Image": image_paths[0]}],
            },
            principal_id="telegram:1",
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_images"] == image_paths
    with ZipFile(output_path) as archive:
        assert any(name.startswith("xl/media/") for name in archive.namelist())


def test_browser_image_collect_accepts_modern_browser_image_content_type(tmp_path, monkeypatch) -> None:
    source_image = _write_content_png(tmp_path / "source.png")
    monkeypatch.setenv("NULLION_WORKSPACE_STORAGE_ROOT", str(tmp_path / "workspaces"))

    materialized = tools._materialize_browser_image_artifact(
        invocation=_invoke("browser_image_collect", {}, principal_id="telegram:1"),
        image_payload={
            "source_url": "https://cdn.example.test/rendered-image.avif",
            "final_url": "https://cdn.example.test/rendered-image.avif",
            "content_type": "image/avif",
            "data": source_image.read_bytes(),
        },
        output_stem="qa-modern-image",
    )

    assert Path(materialized["local_path"]).is_file()
    assert materialized["format"] in {"png", "jpg", "jpeg", "gif"}


def test_browser_image_collect_keeps_trying_until_requested_images_are_saved(tmp_path, monkeypatch) -> None:
    source_image = _write_content_png(tmp_path / "source.png")
    image_data_url = "data:image/png;base64," + base64.b64encode(source_image.read_bytes()).decode("ascii")
    monkeypatch.setenv("NULLION_WORKSPACE_STORAGE_ROOT", str(tmp_path / "workspaces"))

    collect = tools._build_browser_image_collect_handler()
    collected = collect(
        _invoke(
            "browser_image_collect",
            {
                "html": f"""
                <html><body>
                  <img src="data:image/svg+xml;base64,PHN2Zy8+">
                  <source srcset="data:image/svg+xml;base64,PHN2Zy8+ 1x">
                  <img data-src="{image_data_url}">
                </body></html>
                """,
                "output_stem": "qa-collected-image",
                "max_images": 1,
            },
            principal_id="telegram:1",
        )
    )

    assert collected.status == "completed"
    assert collected.output["candidate_count"] == 3
    assert collected.output["saved_count"] == 1
    assert len(collected.output["skipped_images"]) == 2
    assert Path(collected.output["image_paths"][0]).is_file()


def test_browser_image_collect_rejects_tiny_images_by_default(tmp_path, monkeypatch) -> None:
    source_image = _write_png(tmp_path / "tiny.png", color="blue")
    image_data_url = "data:image/png;base64," + base64.b64encode(source_image.read_bytes()).decode("ascii")
    monkeypatch.setenv("NULLION_WORKSPACE_STORAGE_ROOT", str(tmp_path / "workspaces"))

    collect = tools._build_browser_image_collect_handler()
    result = collect(
        _invoke(
            "browser_image_collect",
            {"image_urls": [image_data_url], "output_stem": "tiny-source"},
            principal_id="telegram:1",
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "image_collection_failed"
    assert result.output["image_paths"] == []
    assert "too small" in result.output["skipped_images"][0]["reason"]


def test_browser_image_collect_can_opt_into_small_image_assets(tmp_path, monkeypatch) -> None:
    source_image = _write_png(tmp_path / "tiny-logo.png", color="blue")
    image_data_url = "data:image/png;base64," + base64.b64encode(source_image.read_bytes()).decode("ascii")
    monkeypatch.setenv("NULLION_WORKSPACE_STORAGE_ROOT", str(tmp_path / "workspaces"))

    collect = tools._build_browser_image_collect_handler()
    result = collect(
        _invoke(
            "browser_image_collect",
            {
                "image_urls": [image_data_url],
                "output_stem": "tiny-logo",
                "quality_profile": "any",
            },
            principal_id="telegram:1",
        )
    )

    assert result.status == "completed"
    assert len(result.output["image_paths"]) == 1
    assert Path(result.output["image_paths"][0]).is_file()


def test_browser_image_collect_prefers_declared_larger_content_candidates(tmp_path, monkeypatch) -> None:
    tiny_image = _write_png(tmp_path / "tiny.png", color="blue")
    content_image = _write_content_png(tmp_path / "content.png")
    tiny_data_url = "data:image/png;base64," + base64.b64encode(tiny_image.read_bytes()).decode("ascii")
    content_data_url = "data:image/png;base64," + base64.b64encode(content_image.read_bytes()).decode("ascii")
    monkeypatch.setenv("NULLION_WORKSPACE_STORAGE_ROOT", str(tmp_path / "workspaces"))

    collect = tools._build_browser_image_collect_handler()
    result = collect(
        _invoke(
            "browser_image_collect",
            {
                "html": f"""
                <html><body>
                  <img src="{tiny_data_url}" width="24" height="24" alt="small decorative asset">
                  <img src="{content_data_url}" width="320" height="220" alt="content asset">
                </body></html>
                """,
                "output_stem": "ranked-source",
                "max_images": 1,
            },
            principal_id="telegram:1",
        )
    )

    assert result.status == "completed"
    assert result.output["saved_count"] == 1
    assert result.output["images"][0]["width"] == 320
    assert result.output["images"][0]["height"] == 220


def test_browser_image_collect_rejects_placeholder_image_services(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("NULLION_WORKSPACE_STORAGE_ROOT", str(tmp_path / "workspaces"))
    collect = tools._build_browser_image_collect_handler()

    result = collect(
        _invoke(
            "browser_image_collect",
            {
                "image_urls": [
                    "https://httpbin.org/image/png",
                    "https://placehold.co/100x100/png?text=1",
                ],
                "output_stem": "placeholder-product",
            },
            principal_id="telegram:1",
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "image_collection_failed"
    assert result.output["image_paths"] == []
    reasons = [entry["reason"] for entry in result.output["skipped_images"]]
    assert all("Blocked placeholder image URL" in reason for reason in reasons)


def test_spreadsheet_create_rejects_browser_screenshot_as_content_image_by_default(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "screenshot.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name"],
                "rows": [{"Name": "Item", "image_path": str(screenshot)}],
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(screenshot)]
    assert not output_path.exists()


def test_spreadsheet_create_rejects_browser_screenshot_even_if_model_sets_allow_flag(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "allowed-screenshot.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name"],
                "rows": [{"Name": "Item", "image_path": str(screenshot)}],
                "allow_screenshot_images": True,
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(screenshot)]
    assert not output_path.exists()


def test_spreadsheet_create_embeds_browser_screenshot_through_screenshot_paths(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "with-screenshot.xlsx"
    handler = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(output_path),
                "columns": ["Name"],
                "rows": [{"Name": "Item"}],
                "screenshot_paths": [str(screenshot)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["columns"] == 2
    assert result.output["embedded_images"] == []
    assert result.output["embedded_screenshots"] == [str(screenshot)]
    with ZipFile(output_path) as archive:
        assert any(name.startswith("xl/media/") for name in archive.namelist())


def test_document_create_fails_when_requested_images_are_skipped(tmp_path) -> None:
    handler = tools._build_document_create_handler(workspace_root=tmp_path)
    missing_image = tmp_path / "missing.png"
    output_path = tmp_path / "out.docx"

    result = handler(
        _invoke(
            "document_create",
            {
                "output_path": str(output_path),
                "title": "Document",
                "paragraphs": ["One"],
                "image_paths": [str(missing_image)],
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(missing_image)]
    assert not output_path.exists()


def test_document_create_autolinks_plain_http_urls(tmp_path) -> None:
    handler = tools._build_document_create_handler(workspace_root=tmp_path)
    output_path = tmp_path / "linked.docx"

    result = handler(
        _invoke(
            "document_create",
            {
                "output_path": str(output_path),
                "title": "Linked",
                "paragraphs": ["Direct Listing Link: https://shop.test/listing/12345/item-name"],
                "sections": [
                    {
                        "heading": "Entry",
                        "body": "Source: https://shop.test/listing/67890/other-item.",
                        "bullets": ["Buy: https://shop.test/listing/11111/bullet"],
                    }
                ],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["quality_profile"] == "report_quality_v1"
    assert "styled_headings" in result.output["layout_features"]
    with ZipFile(output_path) as archive:
        rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_xml = archive.read("docProps/core.xml").decode("utf-8")
    assert "https://shop.test/listing/12345/item-name" in rels
    assert "https://shop.test/listing/67890/other-item" in rels
    assert "https://shop.test/listing/11111/bullet" in rels
    assert "Linked" in document_xml
    assert "Linked" in settings_xml


def test_document_create_writes_real_tables_and_embedded_images(tmp_path) -> None:
    image_path = _write_content_png(tmp_path / "product.png")
    handler = tools._build_document_create_handler(workspace_root=tmp_path)
    output_path = tmp_path / "table-and-image.docx"

    result = handler(
        _invoke(
            "document_create",
            {
                "output_path": str(output_path),
                "title": "Product memo",
                "paragraphs": ["Summary with image-backed product evidence."],
                "tables": [
                    {
                        "headers": ["Product", "Price", "Link"],
                        "rows": [["First", "$10", "https://example.com/item/1"]],
                    }
                ],
                "image_paths": [str(image_path)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["tables"] == 1
    assert result.output["embedded_images"] == [str(image_path)]
    with ZipFile(output_path) as archive:
        names = archive.namelist()
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "<w:tbl" in document_xml
    assert any(name.startswith("word/media/") for name in names)


def test_visual_artifact_tools_embed_optimized_images_without_losing_source_contract(tmp_path) -> None:
    from PIL import Image

    source = tmp_path / "large-source.png"
    image = Image.effect_noise((2200, 1600), 95).convert("RGB")
    image.save(source, format="PNG")
    assert source.stat().st_size > 2_000_000

    docx_path = tmp_path / "optimized.docx"
    pdf_path = tmp_path / "optimized.pdf"
    pptx_path = tmp_path / "optimized.pptx"
    xlsx_path = tmp_path / "optimized.xlsx"

    docx = tools._build_document_create_handler(workspace_root=tmp_path)(
        _invoke(
            "document_create",
            {
                "output_path": str(docx_path),
                "title": "Optimized document",
                "paragraphs": ["Image evidence."],
                "image_paths": [str(source)],
            },
        )
    )
    pdf = tools._build_pdf_create_handler(workspace_root=tmp_path)(
        _invoke(
            "pdf_create",
            {
                "output_path": str(pdf_path),
                "title": "Optimized PDF",
                "text_pages": ["Image evidence."],
                "image_paths": [str(source)],
            },
        )
    )
    pptx = tools._build_presentation_create_handler(workspace_root=tmp_path)(
        _invoke(
            "presentation_create",
            {
                "output_path": str(pptx_path),
                "title": "Optimized deck",
                "image_paths": [str(source)],
            },
        )
    )
    xlsx = tools._build_spreadsheet_create_handler(workspace_root=tmp_path)(
        _invoke(
            "spreadsheet_create",
            {
                "output_path": str(xlsx_path),
                "title": "Optimized workbook",
                "columns": ["Item"],
                "rows": [{"Item": "Image row", "image_path": str(source)}],
            },
        )
    )

    for result in (docx, pdf, pptx, xlsx):
        assert result.status == "completed", result.error
        assert result.output.get("optimized_image_paths")
        optimized_path = Path(result.output["optimized_image_paths"][0])
        assert optimized_path.is_file()
        assert optimized_path.stat().st_size < source.stat().st_size

    assert docx.output["embedded_images"] == [str(source)]
    assert pdf.output["source_image_paths"] == [str(source)]
    assert pptx.output["embedded_images"] == [str(source)]
    assert xlsx.output["embedded_images"] == [str(source)]
    assert docx_path.stat().st_size < source.stat().st_size
    assert pdf_path.stat().st_size < source.stat().st_size
    assert pptx_path.stat().st_size < source.stat().st_size
    assert xlsx_path.stat().st_size < source.stat().st_size
    for artifact_path, prefix in (
        (docx_path, "word/media/"),
        (pptx_path, "ppt/media/"),
        (xlsx_path, "xl/media/"),
    ):
        with ZipFile(artifact_path) as archive:
            media_sizes = [len(archive.read(name)) for name in archive.namelist() if name.startswith(prefix)]
        assert media_sizes
        assert max(media_sizes) < source.stat().st_size


def test_document_create_embeds_svg_companion_raster(tmp_path) -> None:
    svg_path = tmp_path / "fallback.svg"
    png_path = tmp_path / "fallback.png"
    output_path = tmp_path / "with-image.docx"
    svg_path.write_text("<svg xmlns='http://www.w3.org/2000/svg'></svg>", encoding="utf-8")
    _write_content_png(png_path)
    handler = tools._build_document_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "document_create",
            {
                "output_path": str(output_path),
                "title": "Document",
                "paragraphs": ["One"],
                "image_paths": [str(svg_path)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_images"] == [str(png_path)]
    assert result.output["skipped_images"] == []
    with ZipFile(output_path) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())


def test_document_create_rejects_tiny_raster_image(tmp_path) -> None:
    tiny_image = _write_png(tmp_path / "tiny.png")
    output_path = tmp_path / "tiny.docx"
    handler = tools._build_document_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "document_create",
            {
                "output_path": str(output_path),
                "title": "Document",
                "image_paths": [str(tiny_image)],
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(tiny_image)]
    assert not output_path.exists()


def test_document_create_rejects_browser_screenshot_as_content_image_by_default(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "screenshot.docx"
    handler = tools._build_document_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "document_create",
            {
                "output_path": str(output_path),
                "title": "Document",
                "image_paths": [str(screenshot)],
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(screenshot)]
    assert not output_path.exists()


def test_document_create_rejects_browser_screenshot_even_if_model_sets_allow_flag(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "allowed-screenshot.docx"
    handler = tools._build_document_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "document_create",
            {
                "output_path": str(output_path),
                "title": "Document",
                "image_paths": [str(screenshot)],
                "allow_screenshot_images": True,
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(screenshot)]
    assert not output_path.exists()


def test_document_create_embeds_browser_screenshot_through_screenshot_paths(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "with-screenshot.docx"
    handler = tools._build_document_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "document_create",
            {
                "output_path": str(output_path),
                "title": "Document",
                "screenshot_paths": [str(screenshot)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_images"] == []
    assert result.output["embedded_screenshots"] == [str(screenshot)]
    with ZipFile(output_path) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())


def test_presentation_create_fails_when_requested_images_are_skipped(tmp_path) -> None:
    handler = tools._build_presentation_create_handler(workspace_root=tmp_path)
    missing_image = tmp_path / "missing.png"
    output_path = tmp_path / "out.pptx"

    result = handler(
        _invoke(
            "presentation_create",
            {
                "output_path": str(output_path),
                "title": "Presentation",
                "image_paths": [str(missing_image)],
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(missing_image)]
    assert not output_path.exists()


def test_presentation_create_embeds_svg_companion_raster(tmp_path) -> None:
    svg_path = tmp_path / "fallback.svg"
    png_path = tmp_path / "fallback.png"
    output_path = tmp_path / "with-image.pptx"
    svg_path.write_text("<svg xmlns='http://www.w3.org/2000/svg'></svg>", encoding="utf-8")
    _write_content_png(png_path)
    handler = tools._build_presentation_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "presentation_create",
            {
                "output_path": str(output_path),
                "title": "Presentation",
                "image_paths": [str(svg_path)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_images"] == [str(png_path)]
    assert result.output["skipped_images"] == []
    with ZipFile(output_path) as archive:
        assert any(name.startswith("ppt/media/") for name in archive.namelist())


def test_presentation_create_rejects_tiny_raster_image(tmp_path) -> None:
    tiny_image = _write_png(tmp_path / "tiny.png")
    output_path = tmp_path / "tiny.pptx"
    handler = tools._build_presentation_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "presentation_create",
            {
                "output_path": str(output_path),
                "title": "Presentation",
                "image_paths": [str(tiny_image)],
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(tiny_image)]
    assert not output_path.exists()


def test_presentation_create_uses_report_quality_layout_and_preserves_image_ratio(tmp_path) -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    image_path = tmp_path / "wide.png"
    output_path = tmp_path / "report.pptx"
    Image.new("RGB", (240, 80), "purple").save(image_path)
    handler = tools._build_presentation_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "presentation_create",
            {
                "output_path": str(output_path),
                "title": "Report Deck",
                "slides": [{"title": "First", "body": "Summary", "image_paths": [str(image_path)]}],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["quality_profile"] == "report_quality_v1"
    assert "aspect_ratio_safe_media" in result.output["layout_features"]
    deck = Presentation(str(output_path))
    pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    assert pictures[0].width > pictures[0].height
    ratio = pictures[0].width / pictures[0].height
    assert 2.9 < ratio < 3.1


def test_presentation_create_rejects_browser_screenshot_as_content_image_by_default(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "screenshot.pptx"
    handler = tools._build_presentation_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "presentation_create",
            {
                "output_path": str(output_path),
                "title": "Presentation",
                "image_paths": [str(screenshot)],
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(screenshot)]
    assert not output_path.exists()


def test_presentation_create_rejects_browser_screenshot_even_if_model_sets_allow_flag(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "allowed-screenshot.pptx"
    handler = tools._build_presentation_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "presentation_create",
            {
                "output_path": str(output_path),
                "title": "Presentation",
                "image_paths": [str(screenshot)],
                "allow_screenshot_images": True,
            },
        )
    )

    assert result.status == "failed"
    assert "could not embed all requested image/screenshot paths" in (result.error or "")
    assert result.output["skipped_images"] == [str(screenshot)]
    assert not output_path.exists()


def test_presentation_create_embeds_browser_screenshot_through_screenshot_paths(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "with-screenshot.pptx"
    handler = tools._build_presentation_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "presentation_create",
            {
                "output_path": str(output_path),
                "title": "Presentation",
                "screenshot_paths": [str(screenshot)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_images"] == []
    assert result.output["embedded_screenshots"] == [str(screenshot)]
    with ZipFile(output_path) as archive:
        assert any(name.startswith("ppt/media/") for name in archive.namelist())


def test_pdf_create_rejects_browser_screenshot_as_content_image_by_default(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "screenshot.pdf"
    handler = tools._build_pdf_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(output_path),
                "image_paths": [str(screenshot)],
            },
        )
    )

    assert result.status == "failed"
    assert "Browser screenshot artifacts must be supplied through screenshot_paths" in (result.error or "")
    assert not output_path.exists()


def test_pdf_create_text_pages_use_text_preserving_renderer(tmp_path, monkeypatch) -> None:
    image_path = _write_png(tmp_path / "image.png", color="blue")
    output_path = tmp_path / "report.pdf"
    captured: dict[str, object] = {}

    def fake_save(path, *, title, text_pages, image_paths, page_size_name):
        captured.update(
            {
                "path": path,
                "title": title,
                "text_pages": text_pages,
                "image_paths": image_paths,
                "page_size_name": page_size_name,
            }
        )
        _write_valid_renderer_pdf(path, image_paths=image_paths, page_count=len(text_pages))
        return True

    monkeypatch.setattr(tools, "_save_text_pdf_with_chromium", fake_save)
    handler = tools._build_pdf_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(output_path),
                "title": "Report",
                "text_pages": ["Marketplace: Etsy\nListing Price: $14.04\nDirect Link: https://shop.test/listing/12345/item"],
                "image_paths": [str(image_path)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["text_layer"] is True
    assert result.output["text_pages"] == 1
    assert result.output["quality_profile"] == "report_quality_v1"
    assert "clickable_links" in result.output["layout_features"]
    assert captured["text_pages"] == [
        "Marketplace: Etsy\nListing Price: $14.04\nDirect Link: https://shop.test/listing/12345/item"
    ]
    assert captured["image_paths"] == [image_path]
    assert output_path.is_file()


def test_pdf_create_html_pages_use_designed_pdf_renderer(tmp_path, monkeypatch) -> None:
    output_path = tmp_path / "designed.pdf"
    captured: dict[str, object] = {}

    def fake_save(path, *, title, html_pages, image_paths=(), screenshot_paths=(), page_size_name):
        captured.update(
            {
                "path": path,
                "title": title,
                "html_pages": html_pages,
                "image_paths": image_paths,
                "screenshot_paths": screenshot_paths,
                "page_size_name": page_size_name,
            }
        )
        _write_valid_renderer_pdf(path, page_count=len(html_pages))
        return True

    monkeypatch.setattr(tools, "_save_html_pdf_with_chromium", fake_save)
    handler = tools._build_pdf_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(output_path),
                "title": "Fun Malta Guide",
                "html_pages": [
                    "<section class='hero'><h1>Malta Guide</h1><table><tr><th>Hotel</th><th>Price</th></tr>"
                    "<tr><td>Hotel Valentina</td><td>$145</td></tr></table></section>"
                ],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["html_pages"] == 1
    assert result.output["text_layer"] is True
    assert result.output["quality_profile"] == "designed_pdf_v1"
    assert "custom_html_pages" in result.output["layout_features"]
    assert "styled_tables" in result.output["layout_features"]
    assert captured["html_pages"] == [
        "<section class='hero'><h1>Malta Guide</h1><table><tr><th>Hotel</th><th>Price</th></tr>"
        "<tr><td>Hotel Valentina</td><td>$145</td></tr></table></section>"
    ]
    assert output_path.is_file()


def test_pdf_create_inlines_local_html_page_images(tmp_path, monkeypatch) -> None:
    output_path = tmp_path / "designed-with-inline-image.pdf"
    image_path = _write_png(tmp_path / "assets" / "hero.png", color="teal")
    captured: dict[str, object] = {}

    def fake_save(path, *, title, html_pages, image_paths=(), screenshot_paths=(), page_size_name):
        captured.update(
            {
                "path": path,
                "title": title,
                "html_pages": html_pages,
                "image_paths": image_paths,
                "screenshot_paths": screenshot_paths,
                "page_size_name": page_size_name,
            }
        )
        _write_valid_renderer_pdf(path, image_paths=[image_path], page_count=len(html_pages))
        return True

    monkeypatch.setattr(tools, "_save_html_pdf_with_chromium", fake_save)
    handler = tools._build_pdf_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(output_path),
                "title": "Designed Report",
                "html_pages": ['<h1>Designed report</h1><img src="assets/hero.png" alt="Hero">'],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["embedded_html_images"][0]["path"] == str(image_path.resolve())
    assert "verified_media_embeds" in result.output["layout_features"]
    assert "data:image/png;base64," in captured["html_pages"][0]
    assert "assets/hero.png" not in captured["html_pages"][0]


def test_pdf_create_html_pages_include_verified_media_inputs(tmp_path, monkeypatch) -> None:
    output_path = tmp_path / "designed-with-media.pdf"
    image_path = _write_png(tmp_path / "hero.png", color="purple")
    screenshot_path = _write_browser_screenshot_artifact(tmp_path)
    captured: dict[str, object] = {}

    def fake_save(path, *, title, html_pages, image_paths=(), screenshot_paths=(), page_size_name):
        captured.update(
            {
                "path": path,
                "title": title,
                "html_pages": html_pages,
                "image_paths": image_paths,
                "screenshot_paths": screenshot_paths,
                "page_size_name": page_size_name,
            }
        )
        _write_valid_renderer_pdf(path, image_paths=[*image_paths,*screenshot_paths], page_count=len(html_pages))
        return True

    monkeypatch.setattr(tools, "_save_html_pdf_with_chromium", fake_save)
    handler = tools._build_pdf_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(output_path),
                "title": "Designed Report",
                "html_pages": ["<h1>Designed report</h1><table><tr><td>ok</td></tr></table>"],
                "image_paths": [str(image_path)],
                "screenshot_paths": [str(screenshot_path)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["quality_profile"] == "designed_pdf_v1"
    assert result.output["source_image_paths"] == [str(image_path.resolve())]
    assert result.output["source_screenshot_paths"] == [str(screenshot_path.resolve())]
    assert "verified_media_embeds" in result.output["layout_features"]
    assert captured["image_paths"] == [image_path.resolve()]
    assert captured["screenshot_paths"] == [screenshot_path.resolve()]


def test_html_pdf_renderer_inlines_media_assets_as_data_uris(tmp_path, monkeypatch) -> None:
    import sys

    output_path = tmp_path / "inline-media.pdf"
    image_path = _write_png(tmp_path / "hero.png", color="purple")
    captured: dict[str, object] = {}

    class FakePage:
        def set_content(self, html, wait_until=None):  # noqa: ANN001
            captured["html"] = html
            captured["wait_until"] = wait_until

        def pdf(self, **kwargs):  # noqa: ANN003
            captured["pdf_kwargs"] = kwargs
            Path(kwargs["path"]).write_bytes(b"%PDF-1.4\n% fake\n")

    class FakeBrowser:
        def new_page(self):
            return FakePage()

        def close(self):
            captured["closed"] = True

    class FakeChromium:
        def launch(self, **kwargs):  # noqa: ANN003
            captured["launch_kwargs"] = kwargs
            return FakeBrowser()

    class FakePlaywright:
        chromium = FakeChromium()

    class FakePlaywrightContext:
        def __enter__(self):
            return FakePlaywright()

        def __exit__(self, exc_type, exc, tb):  # noqa: ANN001
            return False

    monkeypatch.setattr(tools, "_pdf_chromium_executable", lambda: "/fake/chromium")
    monkeypatch.setitem(
        sys.modules,
        "playwright.sync_api",
        SimpleNamespace(sync_playwright=lambda: FakePlaywrightContext()),
    )

    ok = tools._save_html_pdf_with_chromium(
        output_path,
        title="Inline Media",
        html_pages=["<h1>Inline media</h1>"],
        image_paths=[image_path],
        page_size_name="letter",
    )

    assert ok is True
    assert output_path.is_file()
    assert "data:image/png;base64," in str(captured["html"])
    assert image_path.as_uri() not in str(captured["html"])
    assert "Embedded Visual Assets" not in str(captured["html"])


def test_html_pdf_renderer_places_primary_media_when_layout_has_no_image(tmp_path, monkeypatch) -> None:
    import sys

    output_path = tmp_path / "primary-media.pdf"
    image_path = _write_png(tmp_path / "hero.png", color="purple")
    captured: dict[str, object] = {}

    class FakePage:
        def set_content(self, html, wait_until=None):  # noqa: ANN001
            captured["html"] = html
            captured["wait_until"] = wait_until

        def pdf(self, **kwargs):  # noqa: ANN003
            captured["pdf_kwargs"] = kwargs
            Path(kwargs["path"]).write_bytes(b"%PDF-1.4\n% fake\n")

    class FakeBrowser:
        def new_page(self):
            return FakePage()

        def close(self):
            pass

    class FakeChromium:
        def launch(self, **kwargs):  # noqa: ANN003
            return FakeBrowser()

    class FakePlaywright:
        chromium = FakeChromium()

    class FakePlaywrightContext:
        def __enter__(self):
            return FakePlaywright()

        def __exit__(self, exc_type, exc, tb):  # noqa: ANN001
            return False

    monkeypatch.setattr(tools, "_pdf_chromium_executable", lambda: "/fake/chromium")
    monkeypatch.setitem(
        sys.modules,
        "playwright.sync_api",
        SimpleNamespace(sync_playwright=lambda: FakePlaywrightContext()),
    )

    assert tools._save_html_pdf_with_chromium(
        output_path,
        title="Primary Media",
        html_pages=["<h1>Flyer title</h1><p>Body copy</p>"],
        image_paths=[image_path],
        page_size_name="letter",
    )

    html_text = str(captured["html"])
    assert 'class="primary-media"' in html_text
    assert html_text.index('class="primary-media"') < html_text.index("<h1>Flyer title</h1>")
    assert "data:image/png;base64," in html_text


def test_html_pdf_renderer_strips_decorative_emoji_text(tmp_path, monkeypatch) -> None:
    import sys

    output_path = tmp_path / "no-emoji.pdf"
    captured: dict[str, object] = {}

    class FakePage:
        def set_content(self, html, wait_until=None):  # noqa: ANN001
            captured["html"] = html

        def evaluate(self, script, *args):  # noqa: ANN001, ANN002
            return 0.92 if args else 0

        def pdf(self, **kwargs):  # noqa: ANN003
            Path(kwargs["path"]).write_bytes(b"%PDF-1.4\n% fake\n")

    class FakeBrowser:
        def new_page(self):
            return FakePage()

        def close(self):
            pass

    class FakeChromium:
        def launch(self, **kwargs):  # noqa: ANN003
            return FakeBrowser()

    class FakePlaywright:
        chromium = FakeChromium()

    class FakePlaywrightContext:
        def __enter__(self):
            return FakePlaywright()

        def __exit__(self, exc_type, exc, tb):  # noqa: ANN001
            return False

    monkeypatch.setattr(tools, "_pdf_chromium_executable", lambda: "/fake/chromium")
    monkeypatch.setitem(
        sys.modules,
        "playwright.sync_api",
        SimpleNamespace(sync_playwright=lambda: FakePlaywrightContext()),
    )

    assert tools._save_html_pdf_with_chromium(
        output_path,
        title="No Emoji",
        html_pages=["<h1>Malta guide 🏛️</h1><p>Beach stay ✨ 🌈</p>"],
        page_size_name="letter",
    )

    html_text = str(captured["html"])
    assert "Malta guide" in html_text
    assert "Beach stay" in html_text
    assert "🏛" not in html_text
    assert "✨" not in html_text
    assert "🌈" not in html_text


def test_html_pdf_renderer_flows_bottom_anchored_elements_before_print(tmp_path, monkeypatch) -> None:
    import sys

    output_path = tmp_path / "flow-safe.pdf"
    captured: dict[str, object] = {}

    class FakePage:
        def set_content(self, html, wait_until=None):  # noqa: ANN001
            captured["html"] = html
            captured["wait_until"] = wait_until

        def evaluate(self, script):  # noqa: ANN001
            captured["script"] = script
            return 1

        def pdf(self, **kwargs):  # noqa: ANN003
            captured["pdf_kwargs"] = kwargs
            Path(kwargs["path"]).write_bytes(b"%PDF-1.4\n% fake\n")

    class FakeBrowser:
        def new_page(self):
            return FakePage()

        def close(self):
            pass

    class FakeChromium:
        def launch(self, **kwargs):  # noqa: ANN003
            return FakeBrowser()

    class FakePlaywright:
        chromium = FakeChromium()

    class FakePlaywrightContext:
        def __enter__(self):
            return FakePlaywright()

        def __exit__(self, exc_type, exc, tb):  # noqa: ANN001
            return False

    monkeypatch.setattr(tools, "_pdf_chromium_executable", lambda: "/fake/chromium")
    monkeypatch.setitem(
        sys.modules,
        "playwright.sync_api",
        SimpleNamespace(sync_playwright=lambda: FakePlaywrightContext()),
    )

    assert tools._save_html_pdf_with_chromium(
        output_path,
        title="Flow Safe",
        html_pages=[
            "<section><h1>Open House</h1><p>Details</p>"
            "<footer style='position:absolute;bottom:0'>Final callout</footer></section>"
        ],
        page_size_name="letter",
    )

    html_text = str(captured["html"])
    assert "nullion-print-flow-safe" in html_text
    assert "--nullion-pdf-content-height:10.16in" in html_text
    assert ".designed-page-inner>:first-child" in html_text
    assert "min-height:auto!important" in html_text
    assert "dataset.nullionFlowSafe" in str(captured["script"])
    assert ".designed-page-inner>*{break-inside:avoid;page-break-inside:avoid;}" in html_text
    assert captured["pdf_kwargs"]["path"] == str(output_path)


def test_html_pdf_renderer_repairs_broken_image_slots_from_verified_media(tmp_path, monkeypatch) -> None:
    import sys

    output_path = tmp_path / "repaired-media.pdf"
    image_path = _write_png(tmp_path / "hero.png", color="purple")
    captured: dict[str, object] = {"evaluations": []}

    class FakePage:
        def set_content(self, html, wait_until=None):  # noqa: ANN001
            captured["html"] = html
            captured["wait_until"] = wait_until

        def evaluate(self, script, *args):  # noqa: ANN001
            captured["evaluations"].append((script, args))
            if "nullionRepairedMedia" in script:
                return 1
            return 0

        def pdf(self, **kwargs):  # noqa: ANN003
            captured["pdf_kwargs"] = kwargs
            Path(kwargs["path"]).write_bytes(b"%PDF-1.4\n% fake\n")

    class FakeBrowser:
        def new_page(self):
            return FakePage()

        def close(self):
            pass

    class FakeChromium:
        def launch(self, **kwargs):  # noqa: ANN003
            return FakeBrowser()

    class FakePlaywright:
        chromium = FakeChromium()

    class FakePlaywrightContext:
        def __enter__(self):
            return FakePlaywright()

        def __exit__(self, exc_type, exc, tb):  # noqa: ANN001
            return False

    monkeypatch.setattr(tools, "_pdf_chromium_executable", lambda: "/fake/chromium")
    monkeypatch.setitem(
        sys.modules,
        "playwright.sync_api",
        SimpleNamespace(sync_playwright=lambda: FakePlaywrightContext()),
    )

    assert tools._save_html_pdf_with_chromium(
        output_path,
        title="Repair Media",
        html_pages=["<h1>Flyer</h1><img src='missing-hero.png' alt='Hero'>"],
        image_paths=[image_path],
        page_size_name="letter",
    )

    repair_calls = [
        args
        for script, args in captured["evaluations"]
        if "nullionRepairedMedia" in script
    ]
    assert repair_calls
    assert repair_calls[0][0][0].startswith("data:image/png;base64,")
    assert captured["pdf_kwargs"]["path"] == str(output_path)


def test_html_pdf_renderer_uses_dynamic_print_scale_for_designed_pages(tmp_path, monkeypatch) -> None:
    import sys

    output_path = tmp_path / "fit-page.pdf"
    captured: dict[str, object] = {"evaluations": []}

    class FakePage:
        def set_content(self, html, wait_until=None):  # noqa: ANN001
            captured["html"] = html

        def evaluate(self, script, *args):  # noqa: ANN001
            captured["evaluations"].append((script, args))
            if "maxHeight" in script:
                return 0.74
            if "nullionPrintScale" in script:
                return 1
            return 0

        def pdf(self, **kwargs):  # noqa: ANN003
            captured["pdf_kwargs"] = kwargs
            Path(kwargs["path"]).write_bytes(b"%PDF-1.4\n% fake\n")

    class FakeBrowser:
        def new_page(self):
            return FakePage()

        def close(self):
            pass

    class FakeChromium:
        def launch(self, **kwargs):  # noqa: ANN003
            return FakeBrowser()

    class FakePlaywright:
        chromium = FakeChromium()

    class FakePlaywrightContext:
        def __enter__(self):
            return FakePlaywright()

        def __exit__(self, exc_type, exc, tb):  # noqa: ANN001
            return False

    monkeypatch.setattr(tools, "_pdf_chromium_executable", lambda: "/fake/chromium")
    monkeypatch.setitem(
        sys.modules,
        "playwright.sync_api",
        SimpleNamespace(sync_playwright=lambda: FakePlaywrightContext()),
    )

    assert tools._save_html_pdf_with_chromium(
        output_path,
        title="Fit Page",
        html_pages=["<h1>Dense flyer</h1><div>Lots of content</div>"],
        page_size_name="letter",
    )

    html_text = str(captured["html"])
    assert 'class="designed-page"><div class="designed-page-inner">' in html_text
    assert any("nullionPrintScale" in script for script, _ in captured["evaluations"])
    assert captured["pdf_kwargs"]["scale"] == 0.98


def test_pdf_create_splits_single_structured_report_page_to_align_media(tmp_path, monkeypatch) -> None:
    image_paths = [
        _write_png(tmp_path / "uno.png", color="blue"),
        _write_png(tmp_path / "dos.png", color="green"),
        _write_png(tmp_path / "tres.png", color="red"),
    ]
    output_path = tmp_path / "catalog.pdf"
    captured: dict[str, object] = {}

    def fake_save(path, *, title, text_pages, image_paths, page_size_name):
        captured.update(
            {
                "text_pages": text_pages,
                "image_paths": image_paths,
                "page_size_name": page_size_name,
            }
        )
        _write_valid_renderer_pdf(path, image_paths=image_paths, page_count=len(text_pages))
        return True

    monkeypatch.setattr(tools, "_save_text_pdf_with_chromium", fake_save)
    handler = tools._build_pdf_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(output_path),
                "title": "Catalogo",
                "text_pages": [
                    "Catalogo de productos\nFuente: tienda\n\n"
                    "* 1. Collar brillante\nPrecio: 10\nEnlace: https://shop.test/uno\n\n"
                    "* 2. Pulsera brillante\nPrecio: 20\nEnlace: https://shop.test/dos\n\n"
                    "* 3. Anillo brillante\nPrecio: 30\nEnlace: https://shop.test/tres"
                ],
                "image_paths": [str(path) for path in image_paths],
                "media_alignment": "align_pages",
            },
        )
    )

    assert result.status == "completed"
    assert result.output["text_pages"] == 3
    assert result.output["page_count"] == 3
    assert "media_aligned_report_pages" in result.output["layout_features"]
    assert captured["image_paths"] == image_paths
    assert captured["text_pages"] == [
        "Catalogo de productos\nFuente: tienda\n\n* 1. Collar brillante\nPrecio: 10\nEnlace: https://shop.test/uno",
        "* 2. Pulsera brillante\nPrecio: 20\nEnlace: https://shop.test/dos",
        "* 3. Anillo brillante\nPrecio: 30\nEnlace: https://shop.test/tres",
    ]


def test_pdf_create_preserves_single_text_page_by_default(tmp_path, monkeypatch) -> None:
    image_paths = [
        _write_png(tmp_path / "one.png", color="blue"),
        _write_png(tmp_path / "two.png", color="green"),
    ]
    output_path = tmp_path / "preserved.pdf"
    captured: dict[str, object] = {}

    def fake_save(path, *, title, text_pages, image_paths, page_size_name):
        captured.update({"text_pages": text_pages, "image_paths": image_paths})
        _write_valid_renderer_pdf(path, image_paths=image_paths, page_count=len(text_pages))
        return True

    monkeypatch.setattr(tools, "_save_text_pdf_with_chromium", fake_save)
    handler = tools._build_pdf_create_handler(workspace_root=tmp_path)
    single_page = "Resumen\n\n1. Uno\nDetalle\n\n2. Dos\nDetalle"

    result = handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(output_path),
                "title": "Preserved",
                "text_pages": [single_page],
                "image_paths": [str(path) for path in image_paths],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["text_pages"] == 1
    assert result.output["page_count"] == 1
    assert "media_aligned_report_pages" not in result.output["layout_features"]
    assert captured["text_pages"] == [single_page]
    assert captured["image_paths"] == image_paths


def test_pdf_text_html_normalizes_mojibake_bullet_markers() -> None:
    html = tools._pdf_text_html(
        "Highlights\n"
        "b7 Private lakehouse stay\n"
        "• b7 Lakefront coffee setup\n"
        "b7 Relaxed meals and scenic drives\n"
        "DAY 1 b7 FRIDAY"
    )

    assert "b7 Private" not in html
    assert "b7 Lakefront" not in html
    assert "• · Lakefront" not in html
    assert "DAY 1 b7 FRIDAY" not in html
    # Markdown reports now use semantic lists rather than simulated bullets.
    assert "<ul>" in html
    assert "<li>Private lakehouse stay</li>" in html
    assert "<li>Lakefront coffee setup</li>" in html
    assert "Private lakehouse stay" in html
    assert "Lakefront coffee setup" in html
    assert "DAY 1 · FRIDAY" in html


def test_pdf_text_html_renders_markdown_tables_as_tables() -> None:
    html = tools._pdf_text_html(
        "Price comparison\n\n"
        "| Hotel | Area | Price |\n"
        "|---|---|---:|\n"
        "| Hotel Valentina | St. Julian's | $145 |\n"
        "| AC Hotel | St. Julian's | $181 |\n"
        "\nRecommendation: pick the best final checkout total."
    )

    assert '<table class="report-table">' in html
    assert "<th>Hotel</th>" in html
    assert "<td>Hotel Valentina</td>" in html
    assert "|---|" not in html


def test_pdf_create_rejects_browser_screenshot_even_if_model_sets_allow_flag(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "allowed-screenshot.pdf"
    handler = tools._build_pdf_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(output_path),
                "image_paths": [str(screenshot)],
                "allow_screenshot_images": True,
            },
        )
    )

    assert result.status == "failed"
    assert "Browser screenshot artifacts must be supplied through screenshot_paths" in (result.error or "")
    assert not output_path.exists()


def test_pdf_create_embeds_browser_screenshot_through_screenshot_paths(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    output_path = tmp_path / "with-screenshot.pdf"
    handler = tools._build_pdf_create_handler(workspace_root=tmp_path)

    result = handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(output_path),
                "screenshot_paths": [str(screenshot)],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["source_image_paths"] == []
    assert result.output["source_screenshot_paths"] == [str(screenshot)]
    assert output_path.stat().st_size > 0


def test_pdf_edit_rejects_appended_browser_screenshot_by_default(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    source_path = tmp_path / "source.pdf"
    output_path = tmp_path / "edited.pdf"
    create_handler = tools._build_pdf_create_handler(workspace_root=tmp_path)
    edit_handler = tools._build_pdf_edit_handler(workspace_root=tmp_path)
    created = create_handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(source_path),
                "text_pages": ["Source"],
            },
        )
    )
    assert created.status == "completed"

    result = edit_handler(
        _invoke(
            "pdf_edit",
            {
                "input_path": str(source_path),
                "output_path": str(output_path),
                "append_image_paths": [str(screenshot)],
            },
        )
    )

    assert result.status == "failed"
    assert "Browser screenshot artifacts must be supplied through append_screenshot_paths" in (result.error or "")
    assert not output_path.exists()


def test_pdf_edit_rejects_browser_screenshot_even_if_model_sets_allow_flag(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    source_path = tmp_path / "source.pdf"
    output_path = tmp_path / "edited.pdf"
    create_handler = tools._build_pdf_create_handler(workspace_root=tmp_path)
    edit_handler = tools._build_pdf_edit_handler(workspace_root=tmp_path)
    created = create_handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(source_path),
                "text_pages": ["Source"],
            },
        )
    )
    assert created.status == "completed"

    result = edit_handler(
        _invoke(
            "pdf_edit",
            {
                "input_path": str(source_path),
                "output_path": str(output_path),
                "append_image_paths": [str(screenshot)],
                "allow_screenshot_images": True,
            },
        )
    )

    assert result.status == "failed"
    assert "Browser screenshot artifacts must be supplied through append_screenshot_paths" in (result.error or "")
    assert not output_path.exists()


def test_pdf_edit_embeds_browser_screenshot_through_append_screenshot_paths(tmp_path) -> None:
    screenshot = _write_browser_screenshot_artifact(tmp_path)
    source_path = tmp_path / "source.pdf"
    output_path = tmp_path / "edited.pdf"
    create_handler = tools._build_pdf_create_handler(workspace_root=tmp_path)
    edit_handler = tools._build_pdf_edit_handler(workspace_root=tmp_path)
    created = create_handler(
        _invoke(
            "pdf_create",
            {
                "output_path": str(source_path),
                "text_pages": ["Source"],
            },
        )
    )
    assert created.status == "completed"

    result = edit_handler(
        _invoke(
            "pdf_edit",
            {
                "input_path": str(source_path),
                "output_path": str(output_path),
                "append_screenshot_paths": [str(screenshot)],
            },
        )
    )

    assert result.status == "completed"
    assert output_path.stat().st_size > 0


def test_connector_request_timeout_is_short_and_configurable(monkeypatch) -> None:
    monkeypatch.delenv("NULLION_CONNECTOR_REQUEST_TIMEOUT_SECONDS", raising=False)
    assert tools._connector_request_timeout_seconds() == 12

    monkeypatch.setenv("NULLION_CONNECTOR_REQUEST_TIMEOUT_SECONDS", "2")
    assert tools._connector_request_timeout_seconds() == 3

    monkeypatch.setenv("NULLION_CONNECTOR_REQUEST_TIMEOUT_SECONDS", "90")
    assert tools._connector_request_timeout_seconds() == 30


def test_list_crons_exposes_typed_zero_result_count(monkeypatch) -> None:
    monkeypatch.setattr("nullion.crons.list_crons", lambda **_kwargs: [])
    registry = ToolRegistry()
    register_cron_tools(
        registry,
        default_delivery_channel="web",
        default_delivery_target="web:operator",
    )

    result = registry.invoke(_invoke("list_crons", {}))

    assert result.status == "completed"
    assert result.output == {
        "crons": [],
        "count": 0,
        "message": "No crons scheduled.",
    }


def test_plugin_registration_exposes_expected_tool_families(tmp_path, monkeypatch) -> None:
    from nullion import crons
    monkeypatch.delenv("NULLION_IMAGE_GENERATE_ENABLED", raising=False)
    monkeypatch.setattr(crons, "_RUNTIME_DB_PATH", tmp_path / "runtime.db")
    registry = ToolRegistry(filesystem_allowed_roots=[tmp_path])
    runtime = RuntimeDouble(store=RuntimeStore())
    sample_file = tmp_path / "report.txt"
    sample_file.write_text("hello", encoding="utf-8")

    register_search_plugin(
        registry,
        web_searcher=lambda query, limit: [{"title": query, "url": "https://example.com", "snippet": "market summary"}],
    )
    register_email_plugin(
        registry,
        email_searcher=lambda query, limit: [{"id": "msg-1", "subject": query}],
        email_reader=lambda message_id: {"id": message_id, "body": "hello"},
    )
    register_calendar_plugin(
        registry,
        calendar_lister=lambda start, end, limit: [{"id": "evt-1", "summary": f"{start}/{end}"}],
    )
    register_media_plugin(
        registry,
        audio_transcriber=lambda path, language=None: {"text": f"transcribed:{path}", "language": language or "auto"},
        image_text_extractor=lambda path: {"text": f"ocr:{path}"},
        image_generator=lambda prompt, output_path, source_path=None: {"prompt": prompt, "source_path": source_path},
    )
    register_browser_plugin(registry, browser_navigator=lambda url: {"url": url, "title": "Example"})
    register_workspace_plugin(registry, workspace_root=tmp_path)
    register_reminder_tools(registry, runtime, default_chat_id="web:operator")
    register_cron_tools(registry, default_delivery_channel="web", default_delivery_target="web:operator")

    expected_tools = {
        "web_search",
        "email_search",
        "email_read",
        "calendar_list",
        "audio_transcribe",
        "image_extract_text",
        "image_generate",
        "browser_navigate",
        "file_search",
        "set_reminder",
        "list_reminders",
        "delete_reminder",
        "update_reminder",
        "create_cron",
        "list_crons",
        "update_cron",
        "delete_cron",
        "toggle_cron",
        "run_cron",
    }
    assert expected_tools.issubset({spec.name for spec in registry.list_specs()})
    assert registry.list_installed_plugins() == [
        "browser_plugin",
        "calendar_plugin",
        "email_plugin",
        "media_plugin",
        "search_plugin",
        "workspace_plugin",
    ]

    assert registry.invoke(_invoke("web_search", {"query": "stocks", "limit": 1})).status == "completed"
    assert registry.invoke(_invoke("email_search", {"query": "urgent", "limit": 1})).output["results"][0]["id"] == "msg-1"
    assert registry.invoke(_invoke("email_read", {"id": "msg-1"})).output["message"]["body"] == "hello"
    calendar_spec = registry.get_spec("calendar_list")
    assert "query" in calendar_spec.input_schema["properties"]
    calendar_result = registry.invoke(
        _invoke("calendar_list", {"start": "2026-04-29", "end": "2026-04-30", "max": 1, "query": "2026-04-29"})
    )
    assert calendar_result.status == "completed"
    assert calendar_result.output["selected_results"][0]["id"] == "evt-1"
    assert registry.invoke(_invoke("audio_transcribe", {"path": str(sample_file)})).output["text"].startswith("transcribed:")
    assert registry.invoke(_invoke("image_extract_text", {"path": str(sample_file)})).output["text"].startswith("ocr:")
    assert registry.invoke(_invoke("image_generate", {"prompt": "draw", "output_path": str(tmp_path / "out.png")})).status == "completed"
    assert registry.invoke(_invoke("browser_navigate", {"url": "https://example.com"})).output["title"] == "Example"
    assert registry.invoke(_invoke("file_search", {"pattern": "report"})).output["matches"] == [str(sample_file)]


def test_web_search_exposes_systemic_http_failure_as_typed_tool_state() -> None:
    registry = ToolRegistry()

    def unavailable_search(query, limit):  # noqa: ARG001
        raise urllib.error.HTTPError(
            "https://search.example.test",
            402,
            "Payment Required",
            {},
            None,
        )

    register_search_plugin(registry, web_searcher=unavailable_search)

    result = registry.invoke(_invoke("web_search", {"query": "market close", "limit": 5}))

    assert result.status == "failed"
    assert result.output == {
        "query": "market close",
        "reason": "web_search_failed",
        "failure_scope": "tool",
        "status_code": 402,
    }


def test_discovery_tools_expose_continuation_metadata_for_fallback_chains(tmp_path: Path) -> None:
    runtime = SimpleNamespace(store=RuntimeStore(tmp_path / "runtime.db"))
    registry = ToolRegistry(plugin_registration_allowed=True)
    register_email_plugin(
        registry,
        email_searcher=lambda query, limit: [{"id": "msg-1", "subject": query}],
        email_reader=lambda message_id: {"id": message_id, "body": "hello"},
    )
    register_search_plugin(
        registry,
        web_searcher=lambda query, limit: [{"title": query, "url": "https://example.com", "snippet": "evidence"}],
    )
    register_browser_plugin(registry, browser_navigator=lambda url: {"url": url, "title": "Example"})
    register_workspace_plugin(registry, workspace_root=tmp_path)
    register_reminder_tools(registry, runtime, default_chat_id="web:operator")
    register_cron_tools(registry, default_delivery_channel="web", default_delivery_target="web:operator")

    assert registry.get_spec("email_search").continuation_tools == ("email_read",)
    assert set(registry.get_spec("web_search").continuation_tools) == {"web_fetch", "browser_navigate"}
    assert set(registry.get_spec("file_search").continuation_tools) == {"file_read"}
    assert set(registry.get_spec("workspace_summary").continuation_tools) == {"file_search", "file_read"}
    assert set(registry.get_spec("list_reminders").continuation_tools) == {"delete_reminder", "update_reminder"}
    assert set(registry.get_spec("list_crons").continuation_tools) == {
        "run_cron",
        "delete_cron",
        "update_cron",
        "toggle_cron",
    }
    definitions = {definition["name"]: definition for definition in registry.list_tool_definitions()}
    assert set(definitions["email_search"]["continuation_tools"]) == {"email_read"}
    assert set(definitions["web_search"]["continuation_tools"]) == {"browser_navigate"}
    assert "continuation_tools" not in definitions["browser_navigate"]


def test_continuation_metadata_filters_unregistered_tools_from_followup_scope(tmp_path: Path) -> None:
    from nullion.chat_operator import _registry_continuation_tool_names_for_tool_names
    from nullion.turn_context_policy import ScopedTurnToolRegistry, TurnToolEvidence, TurnToolScopeDecision

    registry = ToolRegistry(plugin_registration_allowed=True)
    register_browser_plugin(registry, browser_navigator=lambda url: {"url": url, "title": "Example"})

    assert "browser_snapshot" in registry.get_spec("browser_navigate").continuation_tools
    assert _registry_continuation_tool_names_for_tool_names(registry, ["browser_navigate"]) == ()

    scoped = ScopedTurnToolRegistry(
        registry,
        evidence=TurnToolEvidence(has_url_target=True),
        tool_scope_decision=TurnToolScopeDecision(web_action="live_research", valid=True),
    )
    definitions = {definition["name"]: definition for definition in scoped.list_tool_definitions()}
    assert "browser_navigate" in definitions
    assert "browser_snapshot" not in definitions
    assert "continuation_tools" not in definitions["browser_navigate"]


def test_tool_definitions_filter_specs_without_callable_handlers() -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)
    registry.register(
        tools.ToolSpec(
            name="browser_snapshot",
            description="stale browser snapshot spec",
            risk_level=tools.ToolRiskLevel.LOW,
            side_effect_class=tools.ToolSideEffectClass.READ,
            requires_approval=False,
            timeout_seconds=1,
        ),
        lambda invocation: tools.ToolResult(invocation.invocation_id, invocation.tool_name, "completed", {}),
    )
    registry._handlers.pop("browser_snapshot")

    assert registry.can_invoke_tool("browser_snapshot") is False
    assert "browser_snapshot" not in {definition["name"] for definition in registry.list_tool_definitions()}


def test_scope_request_filters_stale_browser_tools_through_recovery_wrapper(tmp_path: Path) -> None:
    from nullion.agent_orchestrator import _BlockedToolRegistry
    from nullion.turn_context_policy import ScopedTurnToolRegistry, TurnToolEvidence, TurnToolScopeDecision

    registry = ToolRegistry(plugin_registration_allowed=True)
    register_browser_plugin(registry, browser_navigator=lambda url: {"url": url, "title": "Example"})
    for callable_tool_name in ("browser_extract_text", "browser_extract_items"):
        registry.register(
            tools.ToolSpec(
                name=callable_tool_name,
                description=f"callable {callable_tool_name} spec",
                risk_level=tools.ToolRiskLevel.LOW,
                side_effect_class=tools.ToolSideEffectClass.READ,
                requires_approval=False,
                timeout_seconds=1,
            ),
            lambda invocation: tools.ToolResult(invocation.invocation_id, invocation.tool_name, "completed", {}),
        )
    for stale_tool_name in ("browser_snapshot", "browser_wait_for"):
        registry.register(
            tools.ToolSpec(
                name=stale_tool_name,
                description=f"stale {stale_tool_name} spec",
                risk_level=tools.ToolRiskLevel.LOW,
                side_effect_class=tools.ToolSideEffectClass.READ,
                requires_approval=False,
                timeout_seconds=1,
            ),
            lambda invocation: tools.ToolResult(invocation.invocation_id, invocation.tool_name, "completed", {}),
        )
        registry._handlers.pop(stale_tool_name)
    assert "browser_snapshot" in {str(getattr(spec, "name", "") or "") for spec in registry.list_specs()}
    assert registry.can_invoke_tool("browser_snapshot") is False

    wrapped = _BlockedToolRegistry(registry, blocked_tool_names=set())
    scoped = ScopedTurnToolRegistry(
        wrapped,
        evidence=TurnToolEvidence(has_url_target=True),
        tool_scope_decision=TurnToolScopeDecision(web_action="live_research", valid=True),
    )

    result, widened = scoped.apply_scope_request(
        _invoke("request_tool_scope", {"capabilities": ["web"]})
    )

    assert result.status == "completed"
    assert "browser_navigate" in result.output["available_tools"]
    assert "browser_extract_text" in result.output["available_tools"]
    assert "browser_snapshot" not in result.output["available_tools"]
    assert "browser_wait_for" not in result.output["available_tools"]
    assert "browser_snapshot" not in {str(getattr(spec, "name", "") or "") for spec in widened.list_specs()}


def test_email_tools_expose_strict_input_schemas() -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)
    register_email_plugin(
        registry,
        email_searcher=lambda query, limit: [{"id": "msg-1", "subject": query}],
        email_reader=lambda message_id: {"id": message_id, "body": "hello"},
    )

    search_schema = registry.get_spec("email_search").input_schema
    read_schema = registry.get_spec("email_read").input_schema

    assert search_schema is not None
    assert search_schema["required"] == ["query"]
    assert search_schema["additionalProperties"] is False
    assert "limit" in search_schema["properties"]
    assert read_schema is not None
    assert read_schema["required"] == ["id"]
    assert read_schema["additionalProperties"] is False
    assert "provider_id" in read_schema["properties"]


def test_email_read_rejects_placeholder_id_without_provider_call() -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)

    def fail_reader(message_id: str) -> dict[str, object]:
        raise AssertionError(f"provider should not be called for invalid id: {message_id}")

    register_email_plugin(
        registry,
        email_searcher=lambda query, limit: [{"id": "msg-1", "subject": query}],
        email_reader=fail_reader,
    )

    result = registry.invoke(_invoke("email_read", {"id": "placeholder"}))

    assert result.status == "failed"
    assert result.output["reason"] == "invalid_message_id"
    assert result.output["required_source_tool"] == "email_search"
    assert "email_search" in (result.error or "")


def test_email_read_provider_bad_request_requires_source_search() -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)

    def fail_reader(message_id: str) -> dict[str, object]:
        raise urllib.error.HTTPError(
            url=f"https://mail.example/messages/{message_id}",
            code=400,
            msg="Bad Request",
            hdrs=None,
            fp=None,
        )

    register_email_plugin(
        registry,
        email_searcher=lambda query, limit: [{"id": "msg-1", "subject": query}],
        email_reader=fail_reader,
    )

    result = registry.invoke(_invoke("email_read", {"id": "stale-message-id"}))

    assert result.status == "failed"
    assert result.output["reason"] == "invalid_message_id"
    assert result.output["required_source_tool"] == "email_search"
    assert result.output["next_step"] == "Search email and pass one returned result id to email_read."
    assert result.output["http_status"] == 400


def test_connector_plugin_registers_email_search_adapter(monkeypatch) -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)
    calls: list[tuple[str, dict[str, object] | None]] = []

    monkeypatch.setattr(
        tools,
        "_default_email_connector_provider_id",
        lambda principal_id, *, require_write=False: "custom_connector_provider",
    )
    monkeypatch.setattr(
        tools,
        "_connector_connection_for_invocation",
        lambda invocation, provider_id: SimpleNamespace(provider_profile="https://api.maton.ai", permission_mode="write"),
    )

    def fake_connector_json_request(invocation, *, provider_id, url, params=None, method="GET", json_payload=None):
        calls.append((url, params))
        if url.endswith("/messages"):
            return {"resultSizeEstimate": 1, "messages": [{"id": "msg-1"}]}
        if url.endswith("/messages/msg-1"):
            assert params == {
                "format": "metadata",
                "fields": "id,threadId,labelIds,snippet,payload(headers)",
            }
            return {
                "id": "msg-1",
                "threadId": "thread-1",
                "payload": {
                    "headers": [
                        {"name": "Subject", "value": "Test"},
                        {"name": "From", "value": "sender@example.com"},
                    ],
                },
                "snippet": "hello",
            }
        raise AssertionError(f"unexpected URL: {url}")

    monkeypatch.setattr(tools, "_connector_json_request", fake_connector_json_request)

    tools.register_connector_plugin(registry)
    result = registry.invoke(_invoke("email_search", {"query": "newer_than:1d", "limit": 1}))

    assert result.status == "completed"
    assert result.output["provider_id"] == "custom_connector_provider"
    assert result.output["resultSizeEstimate"] == 1
    assert result.output["results"][0]["id"] == "msg-1"
    assert calls[0] == (
        "https://api.maton.ai/google-mail/gmail/v1/users/me/messages",
        {"q": "newer_than:1d", "maxResults": 1},
    )


def test_connector_email_search_caps_metadata_detail_fetches(monkeypatch) -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)
    calls: list[tuple[str, dict[str, object] | None]] = []

    monkeypatch.setenv("NULLION_EMAIL_SEARCH_DETAIL_FETCH_LIMIT", "2")
    monkeypatch.setattr(
        tools,
        "_default_email_connector_provider_id",
        lambda principal_id, *, require_write=False: "custom_connector_provider",
    )
    monkeypatch.setattr(
        tools,
        "_connector_connection_for_invocation",
        lambda invocation, provider_id: SimpleNamespace(provider_profile="https://api.maton.ai", permission_mode="write"),
    )

    def fake_connector_json_request(invocation, *, provider_id, url, params=None, method="GET", json_payload=None):
        calls.append((url, params))
        if url.endswith("/messages"):
            return {
                "resultSizeEstimate": 4,
                "messages": [
                    {"id": "msg-1", "threadId": "thread-1"},
                    {"id": "msg-2", "threadId": "thread-2"},
                    {"id": "msg-3", "threadId": "thread-3"},
                    {"id": "msg-4", "threadId": "thread-4"},
                ],
            }
        if url.endswith("/messages/msg-1") or url.endswith("/messages/msg-2"):
            message_id = url.rsplit("/", 1)[-1]
            return {
                "id": message_id,
                "threadId": message_id.replace("msg", "thread"),
                "payload": {"headers": [{"name": "Subject", "value": f"Subject {message_id}"}]},
            }
        raise AssertionError(f"unexpected detail fetch: {url}")

    monkeypatch.setattr(tools, "_connector_json_request", fake_connector_json_request)

    tools.register_connector_plugin(registry)
    result = registry.invoke(_invoke("email_search", {"query": "in:inbox", "limit": 4}))

    assert result.status == "completed"
    assert [item["id"] for item in result.output["results"]] == ["msg-1", "msg-2", "msg-3", "msg-4"]
    detail_calls = [url for url, _params in calls if "/messages/msg-" in url]
    assert detail_calls == [
        "https://api.maton.ai/google-mail/gmail/v1/users/me/messages/msg-1",
        "https://api.maton.ai/google-mail/gmail/v1/users/me/messages/msg-2",
    ]


def test_connector_email_read_tools_use_account_timeout(monkeypatch) -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)

    monkeypatch.delenv("NULLION_ACCOUNT_READ_TOOL_TIMEOUT_SECONDS", raising=False)
    tools.register_connector_plugin(registry)

    assert registry.get_spec("email_search").timeout_seconds == 45
    assert registry.get_spec("email_read").timeout_seconds == 45
    assert registry.get_spec("email_attachment_read").timeout_seconds == 45

    configured = ToolRegistry(plugin_registration_allowed=True)
    monkeypatch.setenv("NULLION_ACCOUNT_READ_TOOL_TIMEOUT_SECONDS", "120")
    tools.register_connector_plugin(configured)

    assert configured.get_spec("email_search").timeout_seconds == 90
    assert configured.get_spec("email_read").timeout_seconds == 90
    assert configured.get_spec("email_attachment_read").timeout_seconds == 90


def test_connector_email_attachment_read_is_not_automatic_continuation(monkeypatch) -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)

    monkeypatch.delenv("NULLION_ACCOUNT_READ_TOOL_TIMEOUT_SECONDS", raising=False)
    tools.register_connector_plugin(registry)

    assert registry.can_invoke_tool("email_attachment_read")
    assert registry.get_spec("email_search").continuation_tools == ("email_read",)
    assert registry.get_spec("email_read").continuation_tools == ()

    definitions = {definition["name"]: definition for definition in registry.list_tool_definitions()}
    assert definitions["email_search"]["continuation_tools"] == ["email_read"]
    assert "continuation_tools" not in definitions["email_read"]


def test_connector_email_read_rejects_placeholder_id_without_connector_call(monkeypatch) -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)

    monkeypatch.setattr(
        tools,
        "_default_email_connector_provider_id",
        lambda principal_id, *, require_write=False: (_ for _ in ()).throw(
            AssertionError("provider lookup should not run for invalid id")
        ),
    )
    monkeypatch.setattr(
        tools,
        "_connector_json_request",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            AssertionError("connector should not be called for invalid id")
        ),
    )

    tools.register_connector_plugin(registry)
    result = registry.invoke(_invoke("email_read", {"id": "placeholder"}))

    assert result.status == "failed"
    assert result.output["reason"] == "invalid_message_id"
    assert result.output["next_step"] == "Search email and pass one returned result id to email_read."


def test_connector_email_read_bad_request_requires_source_search(monkeypatch) -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)

    monkeypatch.setattr(
        tools,
        "_default_email_connector_provider_id",
        lambda principal_id, *, require_write=False: "custom_connector_provider",
    )
    monkeypatch.setattr(
        tools,
        "_connector_connection_for_invocation",
        lambda invocation, provider_id: SimpleNamespace(provider_profile="https://api.maton.ai", permission_mode="write"),
    )

    def fail_connector_json_request(invocation, *, provider_id, url, params=None, method="GET", json_payload=None):
        raise urllib.error.HTTPError(url=url, code=400, msg="Bad Request", hdrs=None, fp=None)

    monkeypatch.setattr(tools, "_connector_json_request", fail_connector_json_request)

    tools.register_connector_plugin(registry)
    result = registry.invoke(_invoke("email_read", {"id": "stale-message-id"}))

    assert result.status == "failed"
    assert result.output["reason"] == "invalid_message_id"
    assert result.output["required_source_tool"] == "email_search"
    assert result.output["http_status"] == 400
    assert result.output["next_step"] == "Search email and pass one returned result id to email_read."


def test_cron_mutation_tools_reject_placeholder_ids(tmp_path, monkeypatch) -> None:
    from nullion import crons

    monkeypatch.setattr(crons, "_RUNTIME_DB_PATH", tmp_path / "runtime.db")
    registry = ToolRegistry(plugin_registration_allowed=True)
    tools.register_cron_tools(registry, cron_runner=lambda job, invocation=None: "ok")

    delete_result = registry.invoke(_invoke("delete_cron", {"id": "?unknown"}))
    run_result = registry.invoke(_invoke("run_cron", {"ids": ["cron-1", "?unknown"]}))
    run_placeholder_name_result = registry.invoke(_invoke("run_cron", {"name": "__list__"}))

    assert delete_result.status == "failed"
    assert delete_result.output["reason"] == "invalid_cron_id"
    assert delete_result.output["required_source_tool"] == "list_crons"
    assert run_result.status == "failed"
    assert run_result.output["reason"] == "invalid_cron_id"
    assert run_result.output["invalid_ids"] == ["?unknown"]
    assert run_placeholder_name_result.status == "failed"
    assert run_placeholder_name_result.output["reason"] == "invalid_cron_name"
    assert run_placeholder_name_result.output["required_source_tool"] == "list_crons"


def test_reminder_mutation_tools_reject_placeholder_task_ids() -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)
    runtime = RuntimeDouble(store=RuntimeStore())
    register_reminder_tools(registry, runtime, default_chat_id="web:operator")

    delete_result = registry.invoke(_invoke("delete_reminder", {"task_id": "?unknown"}))
    update_result = registry.invoke(_invoke("update_reminder", {"task_id": "?unknown", "text": "later"}))

    assert delete_result.status == "failed"
    assert delete_result.output["reason"] == "invalid_reminder_task_id"
    assert delete_result.output["required_source_tool"] == "list_reminders"
    assert update_result.status == "failed"
    assert update_result.output["reason"] == "invalid_reminder_task_id"


def test_connector_plugin_registers_calendar_list_adapter(monkeypatch) -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)
    calls: list[tuple[str, dict[str, object] | None]] = []

    monkeypatch.setattr(
        tools,
        "_default_calendar_connector_provider_id",
        lambda principal_id: "custom_connector_provider",
    )
    monkeypatch.setattr(
        tools,
        "_connector_connection_for_invocation",
        lambda invocation, provider_id: SimpleNamespace(provider_profile="https://api.maton.ai", permission_mode="read"),
    )

    def fake_connector_json_request(invocation, *, provider_id, url, params=None, method="GET", json_payload=None):
        calls.append((url, params))
        assert provider_id == "custom_connector_provider"
        return {
            "items": [
                {
                    "id": "evt-1",
                    "summary": "Planning",
                    "start": {"dateTime": "2026-05-31T09:00:00-04:00"},
                    "end": {"dateTime": "2026-05-31T09:30:00-04:00"},
                },
                {
                    "id": "evt-2",
                    "summary": "Other",
                    "start": {"dateTime": "2026-05-31T10:00:00-04:00"},
                    "end": {"dateTime": "2026-05-31T10:30:00-04:00"},
                }
            ]
        }

    monkeypatch.setattr(tools, "_connector_json_request", fake_connector_json_request)

    tools.register_connector_plugin(registry)
    result = registry.invoke(
        _invoke(
            "calendar_list",
            {
                "start": "2026-05-31T00:00:00-04:00",
                "end": "2026-06-01T00:00:00-04:00",
                "max": 3,
                "query": "Planning",
            },
        )
    )

    assert result.status == "completed"
    assert result.output["provider_id"] == "custom_connector_provider"
    assert result.output["result_count"] == 2
    assert result.output["results"][0]["summary"] == "Planning"
    assert result.output["selected_results"][0]["id"] == "evt-1"
    assert result.output["selected_result_ids"] == ["evt-1"]
    assert calls == [
        (
            "https://api.maton.ai/google-calendar/calendar/v3/calendars/primary/events",
            {
                "timeMin": "2026-05-31T00:00:00-04:00",
                "timeMax": "2026-06-01T00:00:00-04:00",
                "maxResults": 3,
                "singleEvents": "true",
                "orderBy": "startTime",
            },
        )
    ]


def test_connector_plugin_registers_calendar_write_adapters(monkeypatch) -> None:
    registry = ToolRegistry(plugin_registration_allowed=True)
    calls: list[dict[str, object]] = []

    monkeypatch.setattr(
        tools,
        "_default_calendar_connector_provider_id",
        lambda principal_id, require_write=False: "custom_connector_provider",
    )
    monkeypatch.setattr(
        tools,
        "_connector_connection_for_invocation",
        lambda invocation, provider_id: SimpleNamespace(provider_profile="https://api.maton.ai", permission_mode="write"),
    )

    def fake_connector_json_request(invocation, *, provider_id, url, params=None, method="GET", json_payload=None):
        calls.append(
            {
                "tool_name": invocation.tool_name,
                "provider_id": provider_id,
                "url": url,
                "params": params,
                "method": method,
                "json_payload": json_payload,
            }
        )
        if invocation.tool_name == "calendar_respond" and method == "GET":
            return {
                "id": "evt-1",
                "summary": "Horse Meat Disco",
                "attendees": [{"email": "me@example.com", "self": True, "responseStatus": "needsAction"}],
            }
        if method == "DELETE":
            return {}
        return {
            "id": "evt-1",
            "summary": (json_payload or {}).get("summary", "Horse Meat Disco"),
            "start": (json_payload or {}).get("start", {"dateTime": "2026-06-26T22:00:00-04:00"}),
            "end": (json_payload or {}).get("end", {"dateTime": "2026-06-27T02:00:00-04:00"}),
        }

    monkeypatch.setattr(tools, "_connector_json_request", fake_connector_json_request)

    tools.register_connector_plugin(registry)

    create = registry.invoke(
        _invoke(
            "calendar_create",
            {
                "summary": "Horse Meat Disco",
                "start": "2026-06-26T22:00:00-04:00",
                "end": "2026-06-27T02:00:00-04:00",
                "time_zone": "America/New_York",
                "location": "Knockdown Center",
            },
        )
    )
    update = registry.invoke(
        _invoke(
            "calendar_update",
            {
                "event_id": "evt-1",
                "start": "2026-06-26T23:00:00-04:00",
                "end": "2026-06-27T03:00:00-04:00",
            },
        )
    )
    respond = registry.invoke(_invoke("calendar_respond", {"event_id": "evt-1", "response_status": "accepted"}))
    delete = registry.invoke(_invoke("calendar_delete", {"event_id": "evt-1"}))

    assert create.status == "completed"
    assert create.output["event"]["summary"] == "Horse Meat Disco"
    assert update.status == "completed"
    assert respond.status == "completed"
    assert delete.status == "completed"
    assert [call["method"] for call in calls] == ["POST", "PATCH", "GET", "PATCH", "DELETE"]
    assert calls[0]["url"] == "https://api.maton.ai/google-calendar/calendar/v3/calendars/primary/events"
    assert calls[1]["url"] == "https://api.maton.ai/google-calendar/calendar/v3/calendars/primary/events/evt-1"
    assert calls[0]["json_payload"]["start"] == {
        "dateTime": "2026-06-26T22:00:00-04:00",
        "timeZone": "America/New_York",
    }
    assert calls[3]["json_payload"] == {
        "attendees": [{"email": "me@example.com", "self": True, "responseStatus": "accepted"}]
    }
    assert [call["params"] for call in calls] == [
        {"sendUpdates": "all"},
        {"sendUpdates": "all"},
        None,
        {"sendUpdates": "all"},
        {"sendUpdates": "all"},
    ]


def test_media_plugin_keeps_unconfigured_image_generation_available_with_fallback(tmp_path, monkeypatch) -> None:
    from nullion.system_context import build_system_context_snapshot, format_system_context_for_prompt

    monkeypatch.delenv("NULLION_IMAGE_GENERATE_ENABLED", raising=False)
    registry = ToolRegistry(plugin_registration_allowed=True, filesystem_allowed_roots=[tmp_path])
    register_media_plugin(
        registry,
        audio_transcriber=lambda path, language=None: {"text": "audio"},
        image_text_extractor=lambda path: {"text": "ocr"},
    )

    tool_names = {spec.name for spec in registry.list_specs()}
    assert {"audio_transcribe", "image_extract_text", "image_generate"} <= tool_names

    context = format_system_context_for_prompt(build_system_context_snapshot(tool_registry=registry))
    assert "image_generate [plugin path:media_plugin" in context

    result = registry.invoke(
        _invoke(
            "image_generate",
            {"prompt": "fallback sample", "output_path": str(tmp_path / "fallback.png")},
        )
    )
    assert result.status == "completed"
    assert result.output["provider"] == "local_svg_fallback"
    assert (tmp_path / "fallback.svg").is_file()


def test_weather_forecast_prioritizes_zip_from_location_text() -> None:
    calls: list[str] = []

    def fake_json_get(url: str, timeout: int) -> dict[str, object]:
        calls.append(url)
        if "geocoding-api.open-meteo.com" in url and "10024" in url:
            return {
                "results": [
                    {
                        "name": "New York",
                        "admin1": "New York",
                        "country": "United States",
                        "latitude": 40.71427,
                        "longitude": -74.00597,
                        "timezone": "America/New_York",
                    }
                ]
            }
        if "api.open-meteo.com" in url:
            return {
                "timezone": "America/New_York",
                "current": {"temperature_2m": 55.0, "weather_code": 3},
                "daily": {"time": ["2026-05-15"], "weather_code": [3]},
            }
        raise AssertionError(url)

    registry = ToolRegistry()
    register_weather_forecast_tool(registry, json_get=fake_json_get)

    result = registry.invoke(
        _invoke(
            "weather_forecast",
            {"location_text": "Upper West Side, NY 10024", "forecast_days": 1},
        )
    )

    assert result.status == "completed"
    assert result.output["location"]["name"] == "New York, New York, United States"
    geocode_calls = [call for call in calls if "geocoding-api.open-meteo.com" in call]
    assert len(geocode_calls) == 1
    assert "10024" in geocode_calls[0]


def test_weather_forecast_uses_saved_profile_location_when_arguments_omit_location(monkeypatch) -> None:
    from nullion import preferences

    monkeypatch.setattr(preferences, "load_profile", lambda: {"address": "Providence, RI 02903"})
    calls: list[str] = []

    def fake_json_get(url: str, timeout: int) -> dict[str, object]:
        calls.append(url)
        if "geocoding-api.open-meteo.com" in url:
            assert "02903" in url
            return {
                "results": [
                    {
                        "name": "Providence",
                        "admin1": "Rhode Island",
                        "country": "United States",
                        "latitude": 41.824,
                        "longitude": -71.4128,
                        "timezone": "America/New_York",
                    }
                ]
            }
        if "api.open-meteo.com" in url:
            return {
                "timezone": "America/New_York",
                "current": {"temperature_2m": 73.0, "weather_code": 1},
                "daily": {"time": ["2026-07-13"], "weather_code": [1]},
            }
        raise AssertionError(url)

    registry = ToolRegistry()
    register_weather_forecast_tool(registry, json_get=fake_json_get)

    result = registry.invoke(_invoke("weather_forecast", {"forecast_days": 1}))

    assert result.status == "completed"
    assert result.output["location"]["name"] == "Providence, Rhode Island, United States"
    assert result.output["location_source"] == "saved_profile"
    assert any("02903" in call for call in calls)


def test_market_quote_tool_returns_structured_rows() -> None:
    calls: list[tuple[str, int]] = []

    def fake_quote_fetcher(symbol: str, timeout: int) -> dict[str, object]:
        calls.append((symbol, timeout))
        return {
            "symbol": symbol,
            "name": f"{symbol} Incorporated",
            "price": 123.45,
            "currency": "USD",
            "exchange": "Nasdaq",
            "provider": "test-provider",
        }

    registry = ToolRegistry()
    register_market_quote_tool(registry, quote_fetcher=fake_quote_fetcher)

    spec = registry.get_spec("market_quote")
    assert "market_data" in spec.capability_tags
    assert spec.input_schema["properties"]["symbols"]["type"] == "array"

    result = registry.invoke(_invoke("market_quote", {"symbols": ["msft", "CRML"], "timeout_seconds": 99}))

    assert result.status == "completed"
    assert result.output["symbols"] == ["MSFT", "CRML"]
    assert [quote["symbol"] for quote in result.output["quotes"]] == ["MSFT", "CRML"]
    assert calls == [("MSFT", 12), ("CRML", 12)]


def test_market_quote_tool_fetches_bounded_symbol_batch_concurrently() -> None:
    symbols = [f"Q{index}" for index in range(16)]
    started = time.perf_counter()

    def slow_quote_fetcher(symbol: str, timeout: int) -> dict[str, object]:
        time.sleep(0.05)
        return {"symbol": symbol, "price": 100.0, "currency": "USD"}

    registry = ToolRegistry()
    register_market_quote_tool(registry, quote_fetcher=slow_quote_fetcher)

    result = registry.invoke(_invoke("market_quote", {"symbols": symbols}))
    elapsed = time.perf_counter() - started

    assert result.status == "completed"
    assert [quote["symbol"] for quote in result.output["quotes"]] == symbols
    assert elapsed < 0.35


def test_weather_json_get_retries_transient_url_errors(monkeypatch) -> None:
    attempts = 0

    class Response:
        def __enter__(self) -> "Response":
            return self

        def __exit__(self, exc_type, exc, tb) -> None:
            return None

        def read(self, max_bytes: int) -> bytes:
            assert max_bytes == 500_000
            return b'{"ok": true}'

    class Opener:
        def open(self, request, timeout: int) -> Response:
            nonlocal attempts
            attempts += 1
            assert timeout == 7
            if attempts == 1:
                raise urllib.error.URLError("temporary TLS reset")
            return Response()

    monkeypatch.setattr(tools.urllib.request, "build_opener", lambda *handlers: Opener())

    assert tools._json_get("https://api.open-meteo.com/v1/forecast", 7) == {"ok": True}
    assert attempts == 2


def test_weather_default_json_get_does_not_retry_slow_provider(monkeypatch) -> None:
    attempts = 0

    class Opener:
        def open(self, request, timeout: int):  # noqa: ANN001
            nonlocal attempts
            attempts += 1
            assert timeout == 4
            raise TimeoutError("weather provider stalled")

    monkeypatch.setattr(tools.urllib.request, "build_opener", lambda *handlers: Opener())

    with pytest.raises(TimeoutError):
        tools._weather_json_get("https://api.open-meteo.com/v1/forecast", 4)

    assert attempts == 1


def test_weather_forecast_uses_short_bounded_http_timeouts(monkeypatch) -> None:
    calls: list[tuple[str, int]] = []
    monkeypatch.setenv("NULLION_WEATHER_HTTP_TIMEOUT_SECONDS", "4")

    def fake_json_get(url: str, timeout: int) -> dict[str, object]:
        calls.append((url, timeout))
        if "geocoding-api.open-meteo.com" in url:
            return {
                "results": [
                    {
                        "name": "New York",
                        "admin1": "New York",
                        "country": "United States",
                        "latitude": 40.71427,
                        "longitude": -74.00597,
                        "timezone": "America/New_York",
                    }
                ]
            }
        if "api.open-meteo.com" in url:
            return {
                "timezone": "America/New_York",
                "current": {"temperature_2m": 55.0, "weather_code": 3},
                "daily": {"time": ["2026-05-15"], "weather_code": [3]},
            }
        raise AssertionError(url)

    registry = ToolRegistry()
    register_weather_forecast_tool(registry, json_get=fake_json_get)

    result = registry.invoke(
        _invoke(
            "weather_forecast",
            {"location_text": "10024, New York", "forecast_days": 1},
        )
    )

    assert result.status == "completed"
    assert calls
    assert {timeout for _, timeout in calls} == {4}


def test_weather_forecast_uses_public_locality_fallback_for_private_address() -> None:
    calls: list[str] = []

    def fake_json_get(url: str, timeout: int) -> object:
        calls.append(url)
        if "geocoding-api.open-meteo.com" in url:
            return {"results": []}
        if "nominatim.openstreetmap.org" in url and "Brooklyn%2C+United+States" in url:
            return [
                {
                    "lat": "40.6782",
                    "lon": "-73.9442",
                    "display_name": "Brooklyn, New York, United States",
                    "address": {"country": "United States"},
                }
            ]
        if "nominatim.openstreetmap.org" in url:
            return []
        if "api.open-meteo.com" in url:
            return {
                "timezone": "America/New_York",
                "current": {"temperature_2m": 61.0, "weather_code": 1},
                "daily": {"time": ["2026-05-15"], "weather_code": [1]},
            }
        raise AssertionError(url)

    registry = ToolRegistry()
    register_weather_forecast_tool(registry, json_get=fake_json_get)

    result = registry.invoke(
        _invoke(
            "weather_forecast",
            {"location_text": "Brooklyn, Apt 9, United States", "forecast_days": 1},
        )
    )

    assert result.status == "completed"
    assert result.output["location"]["name"] == "Brooklyn, New York, United States"
    assert any("Brooklyn%2C+United+States" in call for call in calls)


def test_weather_forecast_returns_hourly_rows_when_requested() -> None:
    forecast_calls: list[str] = []

    def fake_json_get(url: str, timeout: int) -> dict[str, object]:
        if "geocoding-api.open-meteo.com" in url:
            return {
                "results": [
                    {
                        "name": "Greentown",
                        "admin1": "Pennsylvania",
                        "country": "United States",
                        "latitude": 41.333,
                        "longitude": -75.305,
                        "timezone": "America/New_York",
                    }
                ]
            }
        if "api.open-meteo.com" in url:
            forecast_calls.append(url)
            return {
                "timezone": "America/New_York",
                "current": {"temperature_2m": 52.8, "weather_code": 61},
                "daily": {"time": ["2026-05-24"], "weather_code": [61]},
                "hourly": {
                    "time": [
                        "2026-05-24T00:00",
                        "2026-05-24T01:00",
                        "2026-05-24T02:00",
                        "2026-05-24T03:00",
                    ],
                    "weather_code": [61, 63, 3, 2],
                    "temperature_2m": [52.8, 53.0, 53.5, 54.1],
                    "apparent_temperature": [48.6, 49.0, 50.0, 51.0],
                    "precipitation_probability": [98, 90, 50, 20],
                    "precipitation": [0.02, 0.01, 0.0, 0.0],
                    "wind_speed_10m": [11.1, 10.2, 9.4, 8.0],
                    "wind_gusts_10m": [18.0, 17.0, 15.0, 12.0],
                },
            }
        raise AssertionError(url)

    registry = ToolRegistry()
    register_weather_forecast_tool(registry, json_get=fake_json_get)
    schema = tools._default_input_schema_for_tool("weather_forecast")
    assert {"include_hourly", "hourly_hours"} <= set(schema["properties"])

    result = registry.invoke(
        _invoke(
            "weather_forecast",
            {
                "location_text": "Greentown, PA 18426",
                "forecast_days": 1,
                "include_hourly": True,
                "hourly_hours": 3,
            },
        )
    )

    assert result.status == "completed"
    assert len(result.output["hourly"]) == 3
    assert result.output["hourly"][0] == {
        "time": "2026-05-24T00:00",
        "weather_code": 61,
        "temperature_f": 52.8,
        "apparent_temperature_f": 48.6,
        "precipitation_probability": 98,
        "precipitation_in": 0.02,
        "wind_speed_mph": 11.1,
        "wind_gusts_mph": 18.0,
        "summary": "slight rain",
    }
    query = parse_qs(urlparse(forecast_calls[0]).query)
    assert "hourly" in query
    assert "temperature_2m" in query["hourly"][0]
    assert "wind_gusts_10m" in query["hourly"][0]


def test_list_reminders_returns_local_display_fields(monkeypatch) -> None:
    import nullion.preferences as preferences

    monkeypatch.setattr(preferences, "load_preferences", lambda: SimpleNamespace(timezone="America/New_York"))
    monkeypatch.setattr(preferences, "resolve_timezone", lambda name: ZoneInfo(name))
    runtime = RuntimeDouble(store=RuntimeStore())
    runtime.store.add_reminder(
        ReminderRecord(
            "rem-1",
            "web:operator",
            "stand up",
            datetime(2026, 5, 15, 20, 0, tzinfo=timezone.utc),
        )
    )
    registry = ToolRegistry()
    register_reminder_tools(registry, runtime, default_chat_id="web:operator")

    result = registry.invoke(_invoke("list_reminders", {}))

    reminder = result.output["reminders"][0]
    assert reminder["due_at"] == "2026-05-15T20:00:00+00:00"
    assert reminder["due_at_local"] == "2026-05-15T16:00:00-04:00"
    assert reminder["timezone"] == "America/New_York"
    assert reminder["due_at_display"] == "15 May at 4:00 PM"


def test_delete_reminder_removes_reminder_and_scheduled_task(monkeypatch) -> None:
    import nullion.preferences as preferences

    class RuntimeWithReminderDelete:
        def __init__(self) -> None:
            self.store = RuntimeStore()
            self.checkpoints = 0

        def checkpoint(self) -> None:
            self.checkpoints += 1

    monkeypatch.setattr(preferences, "load_preferences", lambda: SimpleNamespace(timezone="America/New_York"))
    monkeypatch.setattr(preferences, "resolve_timezone", lambda name: ZoneInfo(name))
    runtime = RuntimeWithReminderDelete()
    task = ScheduledTask("rem-1", "cap", ScheduleKind.ONCE, 0, True, None, 0)
    runtime.store.add_scheduled_task(task)
    runtime.store.add_reminder(
        ReminderRecord(
            "rem-1",
            "web:operator",
            "stand up",
            datetime(2026, 5, 15, 20, 0, tzinfo=timezone.utc),
        )
    )
    registry = ToolRegistry()
    register_reminder_tools(registry, runtime, default_chat_id="web:operator")

    result = registry.invoke(_invoke("delete_reminder", {"task_id": "rem-1"}))

    assert result.status == "completed"
    assert result.output["task_id"] == "rem-1"
    assert result.output["text"] == "stand up"
    assert result.output["action_receipt"]["action"] == "deleted"
    assert runtime.store.get_reminder("rem-1") is None
    assert runtime.store.get_scheduled_task("rem-1") is None
    assert runtime.checkpoints == 1


def test_update_reminder_changes_text_time_and_reenables_task(monkeypatch) -> None:
    import nullion.preferences as preferences

    class RuntimeWithReminderUpdate:
        def __init__(self) -> None:
            self.store = RuntimeStore()
            self.checkpoints = 0

        def checkpoint(self) -> None:
            self.checkpoints += 1

    monkeypatch.setattr(preferences, "load_preferences", lambda: SimpleNamespace(timezone="America/New_York"))
    monkeypatch.setattr(preferences, "resolve_timezone", lambda name: ZoneInfo(name))
    runtime = RuntimeWithReminderUpdate()
    task = ScheduledTask("rem-1", "cap", ScheduleKind.ONCE, 0, False, None, 0)
    runtime.store.add_scheduled_task(task)
    runtime.store.add_reminder(
        ReminderRecord(
            "rem-1",
            "web:operator",
            "stand up",
            datetime(2026, 5, 15, 20, 0, tzinfo=timezone.utc),
            delivered_at=datetime(2026, 5, 15, 20, 0, tzinfo=timezone.utc),
        )
    )
    registry = ToolRegistry()
    register_reminder_tools(registry, runtime, default_chat_id="web:operator")

    result = registry.invoke(
        _invoke(
            "update_reminder",
            {
                "task_id": "rem-1",
                "text": "stretch",
                "due_at": "2026-05-16T11:00:00",
            },
        )
    )

    assert result.status == "completed"
    assert result.output["task_id"] == "rem-1"
    assert result.output["text"] == "stretch"
    assert result.output["due_at"] == "2026-05-16T15:00:00+00:00"
    assert result.output["action_receipt"]["action"] == "updated"
    reminder = runtime.store.get_reminder("rem-1")
    assert reminder is not None
    assert reminder.text == "stretch"
    assert reminder.due_at == datetime(2026, 5, 16, 15, 0, tzinfo=timezone.utc)
    assert reminder.delivered_at is None
    scheduled_task = runtime.store.get_scheduled_task("rem-1")
    assert scheduled_task is not None and scheduled_task.enabled is True
    assert runtime.checkpoints == 1


def test_set_reminder_uses_detected_local_timezone_when_saved_preference_is_default_utc(monkeypatch) -> None:
    import nullion.preferences as preferences

    class RuntimeWithReminder:
        def __init__(self):
            self.store = RuntimeStore()
            self.scheduled_due_at = None
            self.checkpoints = 0

        def schedule_reminder(self, *, chat_id, text, due_at):  # noqa: ANN001
            self.scheduled_due_at = due_at
            task = ScheduledTask("rem-1", "cap", ScheduleKind.ONCE, 0, True, due_at, 0)
            self.store.add_reminder(ReminderRecord(task.task_id, str(chat_id), str(text), due_at))
            return task

        def checkpoint(self, force: bool = False) -> None:
            _ = force
            self.checkpoints += 1

    monkeypatch.setattr(preferences, "load_preferences", lambda: SimpleNamespace(timezone="UTC"))
    monkeypatch.setattr(preferences, "detect_system_timezone", lambda default="UTC": "America/New_York")
    monkeypatch.setattr(preferences, "resolve_timezone", lambda name: ZoneInfo(name))
    runtime = RuntimeWithReminder()
    registry = ToolRegistry()
    register_reminder_tools(registry, runtime, default_chat_id="web:operator")

    result = registry.invoke(
        _invoke(
            "set_reminder",
            {
                "text": "Go out",
                "due_at": "2026-05-16T11:00:00",
            },
        )
    )

    assert result.status == "completed"
    assert runtime.scheduled_due_at == datetime(2026, 5, 16, 15, 0, tzinfo=timezone.utc)
    assert result.output["due_at"] == "2026-05-16T15:00:00+00:00"
    assert result.output["due_at_local"] == "2026-05-16T11:00:00-04:00"
    assert result.output["timezone"] == "America/New_York"
    assert runtime.checkpoints == 1


def test_cached_skill_pack_index_stays_compact_and_omits_full_docs(monkeypatch) -> None:
    from nullion import runtime_cache
    from nullion.skill_pack_installer import format_cached_enabled_skill_pack_index_for_prompt

    runtime_cache.invalidate_namespace("skill_pack.compact_index")

    text = format_cached_enabled_skill_pack_index_for_prompt(
        ("nullion/connector-skills",),
        max_total_chars=420,
    )
    again = format_cached_enabled_skill_pack_index_for_prompt(
        ("nullion/connector-skills",),
        max_total_chars=420,
    )

    assert text == again
    assert "Installed skill index for routing only" in text
    assert "nullion/connector-skills" in text
    assert "For custom email bridges" not in text
    assert len(text) < 700


def test_cron_schema_generation_does_not_depend_on_private_helper(monkeypatch) -> None:
    monkeypatch.delattr(tools, "_cron_tool_properties")

    create_schema = tools._default_input_schema_for_tool("create_cron")
    update_schema = tools._default_input_schema_for_tool("update_cron")

    assert create_schema["required"] == ["name", "schedule", "task"]
    assert update_schema["required"] == ["id"]
    assert set(create_schema["properties"]) == {
        "id",
        "name",
        "schedule",
        "task",
        "enabled",
        "workspace_id",
        "delivery_channel",
        "delivery_target",
        "html_image_delivery_mode",
        "artifact_delivery_options",
    }
    assert create_schema["properties"] is not update_schema["properties"]
    assert create_schema["properties"]["delivery_channel"]["type"] == "string"
    options = create_schema["properties"]["artifact_delivery_options"]
    assert options["type"] == "object"
    assert options["properties"]["requires_attachment_delivery"]["type"] == "boolean"
    assert options["properties"]["required_artifact_extensions"]["items"]["type"] == "string"
    options["properties"]["requires_attachment_delivery"]["type"] = "changed"
    assert update_schema["properties"]["artifact_delivery_options"]["properties"]["requires_attachment_delivery"]["type"] == "boolean"


def test_required_plugin_dependencies_fail_loudly(tmp_path) -> None:
    registry = ToolRegistry()

    with pytest.raises(ValueError, match="email_searcher"):
        register_email_plugin(registry)
    with pytest.raises(ValueError, match="calendar_lister"):
        register_calendar_plugin(registry)
    with pytest.raises(ValueError, match="workspace_root or allowed_roots"):
        register_workspace_plugin(registry)


def test_document_create_pdf_output_delegates_to_pdf_create(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("NULLION_FILE_ACCESS_ENABLED", "1")
    output_path = tmp_path / "croton_gorge_trip.pdf"
    registry = tools.create_plugin_tool_registry(allowed_roots=[tmp_path])

    result = registry.invoke(
        ToolInvocation(
            invocation_id="document-pdf",
            tool_name="document_create",
            principal_id="telegram_chat",
            arguments={
                "output_path": str(output_path),
                "title": "Croton Gorge Day Trip",
                "paragraphs": ["A public-transit day trip pamphlet."],
                "sections": [
                    {
                        "heading": "Trail",
                        "body": "Walk the dam overlook and nearby trails.",
                        "bullets": [
                            "Take Metro-North to Croton-Harmon.",
                            "Use local transit or a short ride to the park.",
                        ],
                    }
                ],
                "tables": [{"headers": ["Stop", "Plan"], "rows": [["NYC", "Train"], ["Croton Gorge", "Walk"]]}],
            },
        )
    )

    assert result.status == "completed"
    assert result.tool_name == "document_create"
    assert result.output["delegated_tool_name"] == "pdf_create"
    assert Path(str(result.output["artifact_path"])) == output_path
    assert output_path.is_file()
    assert output_path.read_bytes().startswith(b"%PDF")


def test_file_read_returns_structured_directory_listing(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("NULLION_FILE_ACCESS_ENABLED", "1")
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    (uploads / "alpha_beta.txt").write_text("alpha", encoding="utf-8")
    (uploads / "photo.png").write_bytes(b"\x89PNG\r\n")
    (uploads / "nested").mkdir()
    registry = tools.create_plugin_tool_registry(allowed_roots=[tmp_path])

    result = registry.invoke(_invoke("file_read", {"path": str(uploads)}))

    assert result.status == "completed"
    assert result.output["type"] == "directory"
    assert result.output["is_directory"] is True
    assert result.output["entry_count"] == 3
    entries = {entry["name"]: entry for entry in result.output["entries"]}
    assert entries["alpha_beta.txt"]["type"] == "file"
    assert entries["photo.png"]["media_type"] == "image/png"
    assert entries["nested"]["type"] == "directory"


def test_file_read_reports_binary_files_without_handler_exception(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("NULLION_FILE_ACCESS_ENABLED", "1")
    image = tmp_path / "photo.png"
    image.write_bytes(b"\x89PNG\r\n\xff\xfe")
    registry = tools.create_plugin_tool_registry(allowed_roots=[tmp_path])

    result = registry.invoke(_invoke("file_read", {"path": str(image)}))

    assert result.status == "failed"
    assert result.output["media_type"] == "image/png"
    assert "Could not inspect image" in str(result.error)
    assert result.output["reason"] == "invalid_image"


def test_browser_image_collect_caps_candidate_fanout_and_fetch_timeout(monkeypatch) -> None:
    html = "\n".join(f'<img src="https://example.com/image-{index}.jpg" width="800" height="600">' for index in range(12))
    fetch_timeouts: list[int] = []

    def fake_fetch(url: str, timeout_seconds: int) -> dict[str, object]:
        fetch_timeouts.append(timeout_seconds)
        raise TimeoutError("simulated slow image host")

    monkeypatch.setattr(tools, "_fetch_browser_image_binary", fake_fetch)
    registry = ToolRegistry()
    tools.register_browser_image_collect_tool(registry)

    result = registry.invoke(
        ToolInvocation(
            invocation_id="image-budget",
            tool_name="browser_image_collect",
            principal_id="web:test",
            arguments={
                "html": html,
                "page_url": "https://example.com/page",
                "max_images": 2,
                "timeout_seconds": 2,
            },
        )
    )

    assert result.status == "failed"
    assert result.output["candidate_count"] == 6
    assert result.output["timeout_seconds"] == 2
    assert fetch_timeouts
    assert max(fetch_timeouts) <= 2
