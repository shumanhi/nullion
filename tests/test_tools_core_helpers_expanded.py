from __future__ import annotations

import io
import gzip
import sys
import tarfile
import time
import urllib.error
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest

from nullion import tools as tools_module
from nullion.approvals import BoundaryPermit
from nullion.policy import BoundaryFact, BoundaryKind
from nullion.runtime import invoke_tool
from nullion.runtime_store import RuntimeStore
from nullion.tools import (
    TerminalAttestationEvidence,
    TerminalBackendDescriptor,
    ToolInvocation,
    ToolRegistry,
    ToolResult,
    ToolRiskLevel,
    ToolSideEffectClass,
    ToolSpec,
    _boundary_approval_context_from_fact,
    _boundary_approval_match_key,
    _boundary_policy_principal_for_fact,
    _boundary_risk_score,
    _build_archive_create_handler,
    _build_file_download_handler,
    _build_file_search_handler,
    _build_file_patch_handler,
    _build_pdf_create_handler,
    _build_pdf_edit_handler,
    _build_spreadsheet_create_handler,
    _build_email_read_handler,
    _build_email_send_handler,
    _build_email_search_handler,
    _build_calendar_list_handler,
    _build_archive_extract_handler,
    _build_file_read_handler,
    _build_file_write_handler,
    _build_connector_request_handler,
    _build_workspace_summary_handler,
    _build_terminal_exec_handler,
    _email_attachment_text_preview,
    _connector_access_enabled,
    _connector_allowed_base_urls,
    _connector_credential_value,
    _connector_provider_id_looks_external,
    _connector_request_headers,
    _connector_request_url,
    create_core_tool_registry,
    _default_input_schema_for_tool,
    _effective_filesystem_roots,
    _env_flag,
    _is_global_literal_ip,
    _launcher_command_requires_probe_attestation,
    _missing_backend_attested_capabilities,
    _network_attempt_allowed,
    _network_is_approved,
    _normalize_network_mode,
    _outbound_network_approval_context_from_result,
    _path_within_any_root,
    _principal_workspace_file_roots,
    _record_wildcard_boundary_permit_accesses,
    _required_attested_capabilities_for_network_mode,
    _resolve_windows_terminal_shell,
    _requires_probe_derived_launcher_attestation,
    _selector_candidates_for_boundary_target,
    _selector_matches_target,
    _selector_matches_www_family,
    normalize_tool_result,
    normalize_tool_status,
    register_email_plugin,
    TerminalBackendDescriptor,
    TerminalExecutionResult,
    _WebFetchResolution,
)


def _minimal_text_pdf_bytes(lines: list[str]) -> bytes:
    text_ops = ["BT", "/F1 12 Tf", "72 720 Td"]
    for index, line in enumerate(lines):
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        if index:
            text_ops.append("0 -18 Td")
        text_ops.append(f"({escaped}) Tj")
    text_ops.append("ET")
    stream = "\n".join(text_ops).encode("ascii")
    objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        (
            b"3 0 obj\n<< /Type /Page /Parent 2 0 R "
            b"/Resources << /Font << /F1 4 0 R >> >> "
            b"/MediaBox [0 0 612 792] /Contents 5 0 R >>\nendobj\n"
        ),
        b"4 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
        b"5 0 obj\n<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream\nendobj\n",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(pdf))
        pdf.extend(obj)
    xref_start = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(f"trailer\n<< /Root 1 0 R /Size {len(offsets)} >>\nstartxref\n{xref_start}\n%%EOF\n".encode("ascii"))
    return bytes(pdf)


def test_web_fetch_tls_context_uses_certifi_when_available(monkeypatch, tmp_path) -> None:
    from nullion import tools

    ca_bundle = tmp_path / "cacert.pem"
    ca_bundle.write_text("", encoding="utf-8")
    context = object()
    calls: list[str | None] = []

    class FakeCertifi:
        @staticmethod
        def where() -> str:
            return str(ca_bundle)

    def fake_create_default_context(*, cafile=None):  # noqa: ANN001
        calls.append(cafile)
        return context

    monkeypatch.setitem(sys.modules, "certifi", FakeCertifi)
    monkeypatch.setattr(tools.ssl, "create_default_context", fake_create_default_context)

    assert tools._web_fetch_tls_context() is context
    assert calls == [str(ca_bundle)]


def fact(kind: BoundaryKind, target: str, **attrs) -> BoundaryFact:
    return BoundaryFact(kind=kind, tool_name="tool", operation="read", target=target, attributes={k: str(v) for k, v in attrs.items()})


def test_tool_status_registry_schemas_and_env_flags(monkeypatch, tmp_path) -> None:
    assert normalize_tool_status(None) == "unknown"
    assert normalize_tool_status("success") == "completed"
    assert normalize_tool_status("error") == "failed"
    assert normalize_tool_status("approval_required") == "denied"
    assert normalize_tool_status("pending") == "nonterminal"
    assert normalize_tool_status("???") == "unknown"
    assert normalize_tool_result(ToolResult("i", "t", "success", {"a": 1})).status == "completed"

    assert _default_input_schema_for_tool("file_read")["required"] == ["path"]
    assert _default_input_schema_for_tool("file_download")["required"] == ["url"]
    assert _default_input_schema_for_tool("archive_extract")["required"] == ["path"]
    assert _default_input_schema_for_tool("archive_create")["required"] == ["output_path"]
    assert _default_input_schema_for_tool("file_write")["required"] == ["path", "content"]
    assert "image_paths" in _default_input_schema_for_tool("pdf_create")["properties"]
    assert "html_pages" in _default_input_schema_for_tool("pdf_create")["properties"]
    assert "slides" in _default_input_schema_for_tool("presentation_create")["properties"]
    assert "report-quality layout profile" in _default_input_schema_for_tool("document_create")["properties"]["title"]["description"]
    assert "aspect-ratio-safe media placement" in _default_input_schema_for_tool("presentation_create")["properties"]["title"]["description"]
    assert "readable report content" in _default_input_schema_for_tool("pdf_create")["properties"]["text_pages"]["description"]
    assert "custom visual layout" in _default_input_schema_for_tool("pdf_create")["properties"]["html_pages"]["description"]
    assert _default_input_schema_for_tool("pdf_edit")["required"] == ["input_path"]
    terminal_schema = _default_input_schema_for_tool("terminal_exec")
    assert terminal_schema["properties"]["shell"]["enum"] == ["auto", "sh", "bash", "zsh", "powershell", "pwsh", "cmd"]
    assert terminal_schema["properties"]["timeout_seconds"]["maximum"] == 300
    assert _default_input_schema_for_tool("unknown")["type"] == "object"

    registry = ToolRegistry(plugin_registration_allowed=False, filesystem_allowed_roots=[tmp_path])
    spec = ToolSpec("alpha", "Alpha tool", ToolRiskLevel.LOW, ToolSideEffectClass.READ, False, 5)
    registry.register(spec, lambda invocation: ToolResult(invocation.invocation_id, invocation.tool_name, "completed", {"ok": True}))
    assert registry.get_spec("alpha") is spec
    assert registry.list_tool_definitions()[0]["name"] == "alpha"
    assert registry.invoke(ToolInvocation("inv", "alpha", "p", {})).output == {"ok": True}
    assert registry.filesystem_allowed_roots() == (tmp_path.resolve(),)
    registry.mark_plugin_installed(" search_plugin ")
    assert registry.is_plugin_installed("search_plugin") is True
    with pytest.raises(ValueError):
        registry.register(spec, lambda invocation: ToolResult("x", "alpha", "completed", {}))
    with pytest.raises(KeyError):
        registry.invoke(ToolInvocation("inv", "missing", "p", {}))
    with pytest.raises(ValueError):
        registry.require_plugin_registration_allowed()

    monkeypatch.delenv("FLAG", raising=False)
    assert _env_flag("FLAG", default=False) is False
    monkeypatch.setenv("FLAG", "off")
    assert _env_flag("FLAG") is False
    monkeypatch.setenv("NULLION_CONNECTOR_ACCESS_ENABLED", "true")
    monkeypatch.setenv("NULLION_CONNECTOR_GATEWAY", "acme")
    assert _connector_access_enabled() is True
    monkeypatch.delenv("NULLION_CONNECTOR_ACCESS_ENABLED", raising=False)
    monkeypatch.delenv("NULLION_CONNECTOR_GATEWAY", raising=False)
    monkeypatch.delenv("NULLION_ENABLED_SKILL_PACKS", raising=False)
    monkeypatch.setattr(
        "nullion.connections.load_connection_registry",
        lambda: SimpleNamespace(connections=[SimpleNamespace(provider_id="skill_pack_connector_acme", active=True)]),
    )
    assert _connector_access_enabled() is True


def test_read_only_tool_timeout_returns_terminal_failure() -> None:
    registry = ToolRegistry()
    registry.register(
        ToolSpec(
            "slow_read",
            "Slow read",
            ToolRiskLevel.LOW,
            ToolSideEffectClass.READ,
            False,
            0.05,
        ),
        lambda invocation: (time.sleep(1.0) or ToolResult(invocation.invocation_id, invocation.tool_name, "completed", {})),
    )

    started = time.perf_counter()
    result = registry.invoke(ToolInvocation("inv-timeout", "slow_read", "principal", {}))

    assert time.perf_counter() - started < 0.5
    assert result.status == "failed"
    assert result.output["reason"] == "handler_timeout"
    assert "timed out" in str(result.error)


def test_write_tool_timeout_does_not_move_handler_to_background() -> None:
    registry = ToolRegistry()
    registry.register(
        ToolSpec(
            "quick_write",
            "Quick write",
            ToolRiskLevel.MEDIUM,
            ToolSideEffectClass.WRITE,
            False,
            0.01,
        ),
        lambda invocation: ToolResult(invocation.invocation_id, invocation.tool_name, "completed", {"ok": True}),
    )

    result = registry.invoke(ToolInvocation("inv-write", "quick_write", "principal", {}))

    assert result.status == "completed"
    assert result.output == {"ok": True}


def test_connector_access_requires_explicit_access_or_active_connection(monkeypatch, tmp_path) -> None:
    monkeypatch.delenv("NULLION_CONNECTOR_ACCESS_ENABLED", raising=False)
    monkeypatch.setenv("NULLION_ENABLED_SKILL_PACKS", "nullion/connector-skills")
    monkeypatch.setattr(
        "nullion.connections.load_connection_registry",
        lambda: SimpleNamespace(connections=[]),
    )

    assert _connector_access_enabled() is False
    assert _connector_provider_id_looks_external("skill_pack_connector_acme") is True
    assert _connector_provider_id_looks_external("acme_connector_provider") is True
    assert _connector_provider_id_looks_external("email_plugin") is False

    registry = create_core_tool_registry(workspace_root=tmp_path)
    assert registry.is_plugin_installed("connector_plugin") is False
    assert "connector_request" not in {spec.name for spec in registry.list_specs()}

    monkeypatch.setattr(
        "nullion.connections.load_connection_registry",
        lambda: SimpleNamespace(connections=[SimpleNamespace(provider_id="acme_connector_provider", active=True)]),
    )

    assert _connector_access_enabled() is True


def test_failed_account_tool_advertises_connected_connector_fallback(monkeypatch) -> None:
    connection = SimpleNamespace(
        provider_id="zephyr_connector_provider",
        display_name="Zephyr mail connector",
        credential_scope="workspace",
        active=True,
    )
    monkeypatch.setattr(
        "nullion.connections.load_connection_registry",
        lambda: SimpleNamespace(connections=[connection]),
    )
    monkeypatch.setattr(
        "nullion.connections.connection_for_principal",
        lambda principal_id, provider_id: connection if provider_id == "zephyr_connector_provider" else None,
    )

    def fail_search(query: str, limit: int) -> list[dict[str, object]]:
        raise RuntimeError("Invalid credentials")

    registry = ToolRegistry()
    register_email_plugin(registry, email_searcher=fail_search)

    result = registry.invoke(ToolInvocation("inv", "email_search", "workspace:workspace_admin", {"query": "inbox", "limit": 1}))

    assert result.status == "failed"
    assert result.output["query"] == "inbox"
    assert result.output["available_connector_providers"] == [
        {
            "provider_id": "zephyr_connector_provider",
            "display_name": "Zephyr mail connector",
            "credential_scope": "workspace",
        }
    ]
    assert "try connector_request" in result.output["next_step"]


def test_failed_account_tool_includes_installed_connector_skill_pack(monkeypatch) -> None:
    pack = SimpleNamespace(pack_id="zephyr/api-gateway-skill")
    connection = SimpleNamespace(
        provider_id="skill_pack_connector_zephyr_api_gateway_skill",
        display_name="Zephyr mail",
        credential_ref="ZEPHYR_TOKEN",
        credential_scope="workspace",
        active=True,
    )
    monkeypatch.setattr("nullion.skill_pack_installer.list_installed_skill_packs", lambda: [pack])
    monkeypatch.setattr(
        "nullion.connections.load_connection_registry",
        lambda: SimpleNamespace(connections=[connection]),
    )
    monkeypatch.setattr(
        "nullion.connections.connection_for_principal",
        lambda principal_id, provider_id: connection if provider_id == "skill_pack_connector_zephyr_api_gateway_skill" else None,
    )

    def fail_search(query: str, limit: int) -> list[dict[str, object]]:
        raise RuntimeError("Invalid credentials")

    registry = ToolRegistry()
    register_email_plugin(registry, email_searcher=fail_search)

    result = registry.invoke(ToolInvocation("inv", "email_search", "workspace:workspace_admin", {"query": "newer:7d", "limit": 3}))

    assert result.status == "failed"
    assert result.output["available_connector_providers"][0]["provider_id"] == "skill_pack_connector_zephyr_api_gateway_skill"
    assert result.output["available_connector_providers"][0]["skill_pack_id"] == "zephyr/api-gateway-skill"
    assert "consult that skill's instructions" in result.output["next_step"]
    assert "suggested_connector_requests" not in result.output


def test_connector_credentials_accept_generic_gateway_prefixes(monkeypatch) -> None:
    monkeypatch.setenv("ACME_API_KEY", "token-1")

    assert _connector_credential_value(None, "acme_connector_provider") == "token-1"
    assert _connector_credential_value(None, "skill_pack_connector_acme") == "token-1"

    monkeypatch.delenv("ACME_API_KEY", raising=False)
    monkeypatch.setenv("ACME_SECRET_KEY", "secret-1")
    assert _connector_credential_value(SimpleNamespace(credential_ref="env:ACME_SECRET_KEY"), "custom_connector_provider") == "secret-1"


def test_connector_request_headers_force_connection_close(monkeypatch) -> None:
    monkeypatch.setenv("ACME_API_KEY", "token-1")

    headers = _connector_request_headers(None, "acme_connector_provider")

    assert headers["Authorization"] == "Bearer token-1"
    assert headers["Connection"] == "close"


def test_connector_base_urls_are_loaded_from_connection_config(monkeypatch) -> None:
    provider_id = "custom_connector_provider"
    connection = SimpleNamespace(
        provider_id=provider_id,
        provider_profile="https://www.example.com/profile-root",
        credential_ref="NULLION_CUSTOM_CONNECTOR_TOKEN",
    )
    monkeypatch.setenv("NULLION_CUSTOM_CONNECTOR_BASE_URL", "https://example.com/api")

    assert _connector_allowed_base_urls(connection, provider_id) == (
        "https://www.example.com/profile-root/",
        "https://example.com/api/",
    )
    assert _connector_request_url(
        "https://example.com/api/mail/v1/messages",
        {"q": "inbox"},
        connection,
        provider_id,
    ).endswith("?q=inbox")
    assert _connector_request_url(
        "/mail/v1/messages",
        {"q": "inbox"},
        connection,
        provider_id,
    ) == "https://www.example.com/profile-root/mail/v1/messages?q=inbox"
    assert _connector_request_url(
        "mail/v1/messages",
        {"q": "inbox"},
        connection,
        provider_id,
    ) == "https://www.example.com/profile-root/mail/v1/messages?q=inbox"
    with pytest.raises(ValueError, match="not under configured connector base URL"):
        _connector_request_url(
            "https://www.google.com/gmail",
            {},
            connection,
            provider_id,
        )


def test_connector_base_urls_are_loaded_from_installed_skill_pack(monkeypatch, tmp_path) -> None:
    provider_id = "skill_pack_connector_zephyr_api_gateway_skill"
    pack_root = tmp_path / "pack"
    pack_root.mkdir()
    (pack_root / "SKILL.md").write_text(
        "Base URL: https://example.com/{app}/{native-api-path}\n"
        "Manage connections at https://www.example.com/connections.\n",
        encoding="utf-8",
    )
    pack = SimpleNamespace(pack_id="zephyr/api-gateway-skill", path=str(pack_root))
    connection = SimpleNamespace(provider_id=provider_id, credential_ref="ZEPHYR_TOKEN")
    monkeypatch.setattr("nullion.skill_pack_installer.list_installed_skill_packs", lambda: [pack])
    monkeypatch.setattr("nullion.skill_pack_installer.get_installed_skill_pack", lambda pack_id: pack)

    assert "https://example.com/" in _connector_allowed_base_urls(connection, provider_id)
    assert "https://www.example.com/connections/" in _connector_allowed_base_urls(connection, provider_id)
    assert _connector_request_url(
        "https://example.com/mail/v1/messages",
        {"q": "inbox"},
        connection,
        provider_id,
    ).endswith("?q=inbox")
    assert _connector_request_url(
        "/google-drive/drive/v3/files",
        {"pageSize": 10},
        connection,
        provider_id,
    ) == "https://example.com/google-drive/drive/v3/files?pageSize=10"
    with pytest.raises(ValueError, match="not under configured connector base URL"):
        _connector_request_url(
            "https://www.google.com/gmail",
            {},
            connection,
            provider_id,
        )


def test_connector_write_requires_connection_permission(monkeypatch) -> None:
    provider_id = "skill_pack_connector_acme_mail"
    connection = SimpleNamespace(
        provider_id=provider_id,
        provider_profile="https://api.example.com/",
        credential_ref="ACME_TOKEN",
        permission_mode="read",
        active=True,
    )
    monkeypatch.setenv("ACME_TOKEN", "secret")
    monkeypatch.setattr("nullion.connections.require_workspace_connection_for_principal", lambda principal_id, candidate: connection)

    handler = _build_connector_request_handler()
    result = handler(
        ToolInvocation(
            "inv",
            "connector_request",
            "operator",
            {
                "provider_id": provider_id,
                "url": "https://api.example.com/gmail/send",
                "method": "POST",
                "json": {"to": "you@example.com"},
            },
        )
    )

    assert result.status == "failed"
    assert "read-only" in str(result.error)
    assert result.output["method"] == "POST"


def test_connector_request_surfaces_remote_connection_state(monkeypatch) -> None:
    from nullion import tools

    provider_id = "skill_pack_connector_maton_ai_api_gateway_skill"
    connection = SimpleNamespace(
        provider_id=provider_id,
        provider_profile="https://api.maton.ai/",
        credential_ref="MATON_API_KEY",
        permission_mode="write",
        active=True,
    )

    class FailingOpener:
        def open(self, request, timeout):  # noqa: ANN001
            body = (
                b'{"message":"Your connections for `google-mail` are either PENDING or FAILED.",'
                b'"type":"Bad Request","code":400}'
            )
            raise urllib.error.HTTPError(
                request.full_url,
                400,
                "Bad Request",
                {"content-type": "application/json"},
                io.BytesIO(body),
            )

    monkeypatch.setenv("MATON_API_KEY", "test-token")
    monkeypatch.setattr(tools, "_connector_connection_for_invocation", lambda invocation, candidate: connection)
    monkeypatch.setattr(tools, "_resolve_web_fetch_resolution", lambda url: SimpleNamespace(host="api.maton.ai", address_infos=()))
    monkeypatch.setattr(tools, "_build_web_fetch_opener", lambda: FailingOpener())

    result = _build_connector_request_handler()(
        ToolInvocation(
            "inv",
            "connector_request",
            "telegram:1",
            {
                "provider_id": provider_id,
                "method": "GET",
                "url": "https://api.maton.ai/google-mail/gmail/v1/users/me/profile",
            },
        )
    )

    assert result.status == "failed"
    assert result.error == "google-mail connection is pending or failed. Reconnect the account before using this connector."
    assert result.output["status_code"] == 400
    assert result.output["connection_state"] == "pending_or_failed"
    assert result.output["connector_app_id"] == "google-mail"


def test_connector_write_posts_payload_when_permission_enabled(monkeypatch) -> None:
    provider_id = "skill_pack_connector_acme_mail"
    connection = SimpleNamespace(
        provider_id=provider_id,
        provider_profile="https://api.example.com/",
        credential_ref="ACME_TOKEN",
        permission_mode="write",
        active=True,
    )
    captured = {}

    class FakeHeaders:
        def get_content_type(self):
            return "application/json"

    class FakeResponse:
        headers = FakeHeaders()
        status = 202

        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, tb):
            return False

        def read(self, size):
            return b'{"ok": true}'

    class FakeOpener:
        def open(self, request, timeout):
            captured["method"] = request.get_method()
            captured["data"] = request.data
            captured["headers"] = dict(request.header_items())
            captured["url"] = request.full_url
            return FakeResponse()

    monkeypatch.setenv("ACME_TOKEN", "secret")
    monkeypatch.setattr("nullion.connections.require_workspace_connection_for_principal", lambda principal_id, candidate: connection)
    monkeypatch.setattr("nullion.tools._resolve_web_fetch_resolution", lambda url: SimpleNamespace(host="api.example.com", address_infos=()))
    monkeypatch.setattr("nullion.tools._build_web_fetch_opener", lambda: FakeOpener())

    handler = _build_connector_request_handler()
    result = handler(
        ToolInvocation(
            "inv",
            "connector_request",
            "operator",
            {
                "provider_id": provider_id,
                "url": "https://api.example.com/gmail/send",
                "method": "POST",
                "json": {"to": "you@example.com"},
            },
        )
    )

    assert result.status == "completed"
    assert result.output["method"] == "POST"
    assert captured["method"] == "POST"
    assert captured["data"] == b'{"to": "you@example.com"}'
    assert captured["headers"]["Content-type"] == "application/json"


def test_filesystem_and_network_helper_edges(monkeypatch, tmp_path) -> None:
    monkeypatch.setattr("nullion.workspace_storage.workspace_file_roots_for_principal", lambda principal: [tmp_path / "workspace"])
    assert _principal_workspace_file_roots("user:1") == ((tmp_path / "workspace").resolve(),)
    assert _principal_workspace_file_roots("anonymous") == ()

    invocation = ToolInvocation("inv", "file_read", "user:1", {})
    roots = _effective_filesystem_roots(invocation=invocation, resolved_root=tmp_path, resolved_allowed_roots=None, include_principal_workspace=True)
    assert tmp_path.resolve() in roots
    inside = tmp_path / "child.txt"
    outside = tmp_path.parent / "outside.txt"
    assert _path_within_any_root(inside.resolve(), (tmp_path.resolve(),)) is True
    assert _path_within_any_root(outside.resolve(), (tmp_path.resolve(),)) is False

    assert _is_global_literal_ip("8.8.8.8") is True
    assert _is_global_literal_ip("127.0.0.1") is False
    assert _normalize_network_mode("full") == "full"
    assert _normalize_network_mode("bad") is None
    assert _required_attested_capabilities_for_network_mode("none") == ("network_policy_enforced", "network_policy_enforced.none")
    assert _required_attested_capabilities_for_network_mode("full") == ()
    assert _network_attempt_allowed(attempt={"address_class": "localhost"}, network_mode="localhost_only", approved_targets=[]) is True
    assert _network_attempt_allowed(attempt={"target": "https://example.com/path"}, network_mode="approved_only", approved_targets=["https://example.com/*"]) is True
    assert _network_attempt_allowed(attempt={"target": "https://evil.com"}, network_mode="approved_only", approved_targets=["https://example.com/*"]) is False
    assert _network_attempt_allowed(attempt={}, network_mode="none", approved_targets=[]) is False
    assert _network_is_approved(target="https://www.example.com/a", approved_targets=["https://example.com/*"]) is True


def test_virtual_workspace_path_resolves_to_principal_workspace_without_approval(monkeypatch, tmp_path) -> None:
    monkeypatch.setenv("NULLION_WORKSPACE_STORAGE_ROOT", str(tmp_path / "workspaces"))
    store = RuntimeStore()
    registry = create_core_tool_registry()

    result = invoke_tool(
        store,
        invocation=ToolInvocation(
            "inv",
            "file_write",
            "telegram_chat",
            {"path": "/workspace/bird1.png", "content": "png-ish"},
        ),
        registry=registry,
    )

    expected_path = tmp_path / "workspaces" / "workspace_admin" / "bird1.png"
    assert result.status == "completed"
    assert result.output["path"] == str(expected_path.resolve())
    assert expected_path.read_text(encoding="utf-8") == "png-ish"
    assert store.list_approval_requests() == []


def test_boundary_approval_contexts_and_match_keys() -> None:
    assert _selector_candidates_for_boundary_target("https://example.com/path")["allow_once"] == "example.com"
    assert _selector_candidates_for_boundary_target("/tmp/file")["allow_once"] == "/tmp/file"

    result = ToolResult(
        "inv",
        "terminal_exec",
        "denied",
        {
            "reason": "network_denied",
            "network_mode": "approved_only",
            "egress_attempts": [{"kind": "outbound_network", "target": "https://example.com", "address_class": "public"}],
        },
    )
    assert _outbound_network_approval_context_from_result(result)["target"] == "https://example.com"
    assert _outbound_network_approval_context_from_result(ToolResult("i", "web_fetch", "denied", {})) is None

    network = fact(BoundaryKind.OUTBOUND_NETWORK, "https://example.com", address_class="public")
    private_network = fact(BoundaryKind.OUTBOUND_NETWORK, "http://127.0.0.1", address_class="loopback")
    filesystem = fact(BoundaryKind.FILESYSTEM_ACCESS, "/tmp/a.txt", path="/tmp/a.txt")
    account = fact(BoundaryKind.ACCOUNT_ACCESS, "gmail:me", account_type="gmail")
    assert _boundary_approval_context_from_fact(network)["boundary_kind"] == "outbound_network"
    assert _boundary_approval_context_from_fact(private_network) is None
    assert _boundary_approval_context_from_fact(filesystem)["path"] == "/tmp/a.txt"
    assert _boundary_approval_context_from_fact(account)["selector_candidates"]["always_allow"] == "gmail:*"
    assert _boundary_risk_score(account) == 7
    assert _boundary_risk_score(filesystem) == 6
    assert _boundary_risk_score(private_network) == 9
    assert _boundary_risk_score(network) == 4
    assert _boundary_policy_principal_for_fact("user:1", filesystem)
    assert _boundary_policy_principal_for_fact("user:1", network)

    context = _boundary_approval_context_from_fact(network)
    assert _boundary_approval_match_key(context=context, fallback_target="fallback") == ("outbound_network", "example.com")
    assert _boundary_approval_match_key(context={}, fallback_target="fallback") is None


def test_selector_matching_and_terminal_attestation_helpers() -> None:
    assert _selector_matches_target(selector="https://example.com/*", target="https://www.example.com/a") is True
    assert _selector_matches_target(selector="/tmp/*", target="/tmp/file") is True
    assert _selector_matches_target(selector="/tmp/*", target="/tmp") is True
    assert _selector_matches_target(selector="*.txt", target="report.txt") is True
    assert _selector_matches_www_family(selector="https://example.com/path", target="https://www.example.com/path/") is True
    assert _selector_matches_www_family(selector="https://evil.com/*", target="https://example.com") is False

    descriptor = TerminalBackendDescriptor(
        mode="subprocess",
        attested_capabilities=("network_policy_enforced",),
        metadata={},
        evidence=None,
    )
    assert _missing_backend_attested_capabilities(descriptor, required_capabilities=["network_policy_enforced", "x"]) == ("x",)
    assert _requires_probe_derived_launcher_attestation(["network_policy_enforced.none"]) is True
    assert _requires_probe_derived_launcher_attestation(["other"]) is False
    assert _launcher_command_requires_probe_attestation("/usr/bin/nullion-launcher") is True
    assert _launcher_command_requires_probe_attestation("nullion-launcher") is False
    assert _launcher_command_requires_probe_attestation(None) is False


def test_wildcard_boundary_permit_access_records_events(monkeypatch) -> None:
    store = RuntimeStore()
    permit = BoundaryPermit(
        permit_id="permit",
        approval_id="approval",
        principal_id="global",
        boundary_kind=BoundaryKind.OUTBOUND_NETWORK,
        selector="*",
        granted_by="admin",
        granted_at=datetime.now(UTC),
    )
    network_fact = fact(BoundaryKind.OUTBOUND_NETWORK, "https://example.com", address_class="public")
    monkeypatch.setattr("nullion.tools.extract_boundary_facts", lambda invocation: [network_fact])

    invocation = ToolInvocation("inv", "terminal_exec", "user:1", {}, capsule_id="cap")
    _record_wildcard_boundary_permit_accesses(store, invocation, [permit, permit])

    events = store.list_events()
    assert len(events) == 1
    assert events[0].event_type == "boundary_permit.wildcard_access"
    assert store.list_audit_records()[0].action == "boundary_permit.wildcard_access"


def test_file_read_write_patch_and_workspace_summary_handlers(tmp_path) -> None:
    read_handler = _build_file_read_handler(workspace_root=tmp_path, include_principal_workspace=False)
    write_handler = _build_file_write_handler(workspace_root=tmp_path, include_principal_workspace=False)
    patch_handler = _build_file_patch_handler(workspace_root=tmp_path, include_principal_workspace=False)
    summary_handler = _build_workspace_summary_handler(workspace_root=tmp_path, include_principal_workspace=False)

    outside = tmp_path.parent / "outside-tools-test.txt"
    assert read_handler(ToolInvocation("inv", "file_read", "operator", {})).error == "Missing required argument: path"
    assert "outside workspace" in read_handler(ToolInvocation("inv", "file_read", "operator", {"path": str(outside)})).error
    assert "File not found" in read_handler(ToolInvocation("inv", "file_read", "operator", {"path": str(tmp_path / "missing.txt")})).error

    target = tmp_path / "folder" / "note.txt"
    assert write_handler(ToolInvocation("inv", "file_write", "operator", {"path": str(target)})).error == "Missing required argument: content"
    written = write_handler(ToolInvocation("inv", "file_write", "operator", {"path": str(target), "content": "alpha beta alpha"}))
    assert written.status == "completed"
    assert target.read_text(encoding="utf-8") == "alpha beta alpha"
    assert read_handler(ToolInvocation("inv", "file_read", "operator", {"path": str(target)})).output["content"] == "alpha beta alpha"

    uploaded = tmp_path / "uploads" / "Archive.zip"
    uploaded.parent.mkdir()
    with zipfile.ZipFile(uploaded, "w") as archive:
        archive.writestr("note.txt", "uploaded")
    wrong_bucket = tmp_path / "artifacts" / "Archive.zip"
    resolved_upload = read_handler(ToolInvocation("inv", "file_read", "operator", {"path": str(wrong_bucket)}))
    assert resolved_upload.status == "completed"
    assert resolved_upload.output["path"] == str(uploaded.resolve())
    uploaded.unlink()

    blocked_workbook = write_handler(
        ToolInvocation(
            "inv",
            "file_write",
            "operator",
            {"path": str(tmp_path / "fake-workbook.xlsx"), "content": "<html>not a workbook</html>"},
        )
    )
    assert blocked_workbook.status == "failed"
    assert "text-only" in blocked_workbook.error
    assert not (tmp_path / "fake-workbook.xlsx").exists()

    text_image_path = tmp_path / "note.png"
    image_like_text = write_handler(
        ToolInvocation(
            "inv",
            "file_write",
            "operator",
            {"path": str(text_image_path), "content": "plain text fixture"},
        )
    )
    assert image_like_text.status == "completed"
    assert text_image_path.read_text(encoding="utf-8") == "plain text fixture"
    text_image_path.unlink()

    assert patch_handler(ToolInvocation("inv", "file_patch", "operator", {"path": str(target), "old_string": "", "new_string": "x"})).error == "old_string must be non-empty"
    assert "replace_all" in patch_handler(ToolInvocation("inv", "file_patch", "operator", {"path": str(target), "old_string": "alpha", "new_string": "x", "replace_all": "yes"})).error
    assert "exactly once" in patch_handler(ToolInvocation("inv", "file_patch", "operator", {"path": str(target), "old_string": "alpha", "new_string": "x"})).error
    patched = patch_handler(ToolInvocation("inv", "file_patch", "operator", {"path": str(target), "old_string": "alpha", "new_string": "x", "replace_all": True}))
    assert patched.output["replacements"] == 2
    assert target.read_text(encoding="utf-8") == "x beta x"
    assert "Old string not found" in patch_handler(ToolInvocation("inv", "file_patch", "operator", {"path": str(target), "old_string": "missing", "new_string": "x"})).error

    summary = summary_handler(ToolInvocation("inv", "workspace_summary", "operator", {}))
    assert summary.status == "completed"
    assert summary.output["file_count"] == 1
    assert summary.output["extensions"] == [{"extension": ".txt", "count": 1}]

    csv_path = tmp_path / "archive_file_list.csv"
    csv_written = write_handler(
        ToolInvocation("inv", "file_write", "operator", {"path": str(csv_path), "content": "name,bytes\nreport.pdf,12\n"})
    )
    assert csv_written.status == "completed"
    assert csv_written.output["artifact_path"] == str(csv_path)
    assert csv_written.output["artifact_paths"] == [str(csv_path)]


def test_file_read_extracts_text_from_pdf_artifact(tmp_path) -> None:
    handler = _build_file_read_handler(workspace_root=tmp_path, include_principal_workspace=False)
    pdf_path = tmp_path / "i-751-checklist.pdf"
    pdf_path.write_bytes(
        _minimal_text_pdf_bytes(
            [
                "Document Checklist for Joint I-751",
                "Joint lease or deed",
                "Last 2 years of joint bank statements",
            ]
        )
    )

    result = handler(ToolInvocation("inv", "file_read", "operator", {"path": str(pdf_path)}))

    assert result.status == "completed"
    assert result.output["type"] == "pdf"
    assert result.output["media_type"] == "application/pdf"
    assert result.output["page_count"] == 1
    assert result.output["extraction_method"] == "pypdf"
    assert "Document Checklist for Joint I-751" in result.output["content"]
    assert "Joint lease or deed" in result.output["content"]


def test_file_read_extracts_structured_rows_and_formulas_from_xlsx_artifact(tmp_path) -> None:
    from openpyxl import Workbook

    workbook_path = tmp_path / "regional-values.xlsx"
    workbook = Workbook()
    values = workbook.active
    values.title = "Values"
    values.append(["Region", "Value"])
    values.append(["North", 12])
    values.append(["South", 19])
    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Total", "=SUM(Values!B2:B3)"])
    workbook.save(workbook_path)

    handler = _build_file_read_handler(workspace_root=tmp_path, include_principal_workspace=False)
    result = handler(ToolInvocation("inv", "file_read", "operator", {"path": str(workbook_path)}))

    assert result.status == "completed"
    assert result.output["type"] == "spreadsheet"
    assert result.output["sheet_count"] == 2
    assert result.output["extraction_method"] == "openpyxl"
    assert result.output["worksheets"][0]["rows"] == [
        ["Region", "Value"],
        ["North", 12],
        ["South", 19],
    ]
    assert result.output["worksheets"][1]["rows"] == [
        ["Metric", "Value"],
        ["Total", "=SUM(Values!B2:B3)"],
    ]
    assert "South\t19" in result.output["content"]
    assert result.output["truncated"] is False


def test_email_attachment_text_preview_extracts_pdf_checklist_text() -> None:
    data = _minimal_text_pdf_bytes(
        [
            "Document Checklist for Joint I-751",
            "Joint lease or deed",
            "Last 2 years of joint bank statements",
        ]
    )

    preview = _email_attachment_text_preview(data, mime_type="application/pdf", filename="i-751-checklist.pdf")

    assert "Document Checklist for Joint I-751" in preview
    assert "Joint lease or deed" in preview
    assert "Last 2 years of joint bank statements" in preview


def test_connector_json_payload_accepts_attachment_response_larger_than_generic_limit() -> None:
    attachment = b"%PDF-1.7\n" + (b"large-pdf-body" * 90_000)
    encoded = tools_module.base64.urlsafe_b64encode(attachment).decode("ascii").rstrip("=")
    response_body = tools_module.json.dumps({"data": encoded}).encode("utf-8")

    class Response:
        def read(self, limit: int) -> bytes:
            return response_body[:limit]

    assert len(response_body) > tools_module._CONNECTOR_JSON_MAX_BODY_BYTES

    payload = tools_module._connector_json_payload_from_response(
        Response(),
        max_body_bytes=tools_module._EMAIL_ATTACHMENT_CONNECTOR_MAX_BODY_BYTES,
    )

    assert tools_module._gmail_decode_attachment_bytes(payload["data"]) == attachment


def test_connector_json_payload_rejects_truncated_generic_response() -> None:
    response_body = tools_module.json.dumps(
        {"data": "x" * (tools_module._CONNECTOR_JSON_MAX_BODY_BYTES + 1)}
    ).encode("utf-8")

    class Response:
        def read(self, limit: int) -> bytes:
            return response_body[:limit]

    with pytest.raises(RuntimeError, match="connector response exceeded"):
        tools_module._connector_json_payload_from_response(Response())


def test_connector_email_attachment_read_source_role_without_delivery_contract(monkeypatch, tmp_path) -> None:
    from nullion.response_fulfillment_contract import artifact_paths_from_tool_results

    encoded = tools_module.base64.urlsafe_b64encode(b"alpha,beta\n1,2\n").decode("ascii").rstrip("=")
    monkeypatch.setattr(tools_module, "_connector_connection_for_invocation", lambda invocation, provider_id: object())
    monkeypatch.setattr(
        tools_module,
        "_email_messages_endpoint_for_provider",
        lambda connection, provider_id: "https://connector.example/google-mail/gmail/v1/users/me/messages",
    )
    monkeypatch.setattr(tools_module, "_connector_json_request", lambda *args, **kwargs: {"data": encoded})
    monkeypatch.setattr(
        "nullion.artifacts.artifact_path_for_generated_workspace_file",
        lambda *, principal_id, suffix, stem="email-attachment": tmp_path / f"{stem}{suffix}",
    )

    handler = tools_module._build_connector_email_attachment_read_handler()
    result = handler(
        ToolInvocation(
            "attachment",
            "email_attachment_read",
            "web:1",
            {
                "provider_id": "custom_connector_provider",
                "message_id": "msg-1",
                "attachment_id": "att-1",
                "filename": "stage.csv",
                "mime_type": "text/csv",
            },
        )
    )

    assert result.status == "completed"
    descriptors = result.output["artifact_descriptors"]
    assert descriptors[0]["role"] == "source"
    assert artifact_paths_from_tool_results([result]) == []


def test_connector_email_attachment_read_requests_attachment_sized_response(monkeypatch, tmp_path) -> None:
    captured: dict[str, object] = {}
    encoded = tools_module.base64.urlsafe_b64encode(b"pdf bytes").decode("ascii").rstrip("=")
    monkeypatch.setattr(tools_module, "_connector_connection_for_invocation", lambda invocation, provider_id: object())
    monkeypatch.setattr(
        tools_module,
        "_email_messages_endpoint_for_provider",
        lambda connection, provider_id: "https://connector.example/google-mail/gmail/v1/users/me/messages",
    )

    def request(*args, **kwargs):
        captured.update(kwargs)
        return {"data": encoded}

    monkeypatch.setattr(tools_module, "_connector_json_request", request)
    monkeypatch.setattr(
        "nullion.artifacts.artifact_path_for_generated_workspace_file",
        lambda *, principal_id, suffix, stem="email-attachment": tmp_path / f"{stem}{suffix}",
    )

    result = tools_module._build_connector_email_attachment_read_handler()(
        ToolInvocation(
            "attachment",
            "email_attachment_read",
            "web:1",
            {
                "provider_id": "custom_connector_provider",
                "message_id": "msg-1",
                "attachment_id": "att-1",
                "filename": "notice.pdf",
                "mime_type": "application/pdf",
            },
        )
    )

    assert result.status == "completed"
    assert captured["max_response_bytes"] == tools_module._EMAIL_ATTACHMENT_CONNECTOR_MAX_BODY_BYTES


def test_connector_email_attachment_read_deliverable_with_delivery_contract(monkeypatch, tmp_path) -> None:
    from nullion.response_fulfillment_contract import artifact_paths_from_tool_results

    encoded = tools_module.base64.urlsafe_b64encode(b"alpha,beta\n1,2\n").decode("ascii").rstrip("=")
    monkeypatch.setattr(tools_module, "_connector_connection_for_invocation", lambda invocation, provider_id: object())
    monkeypatch.setattr(
        tools_module,
        "_email_messages_endpoint_for_provider",
        lambda connection, provider_id: "https://connector.example/google-mail/gmail/v1/users/me/messages",
    )
    monkeypatch.setattr(tools_module, "_connector_json_request", lambda *args, **kwargs: {"data": encoded})
    monkeypatch.setattr(
        "nullion.artifacts.artifact_path_for_generated_workspace_file",
        lambda *, principal_id, suffix, stem="email-attachment": tmp_path / f"{stem}{suffix}",
    )

    handler = tools_module._build_connector_email_attachment_read_handler()
    result = handler(
        ToolInvocation(
            "attachment",
            "email_attachment_read",
            "web:1",
            {
                "provider_id": "custom_connector_provider",
                "message_id": "msg-1",
                "attachment_id": "att-1",
                "filename": "stage.csv",
                "mime_type": "text/csv",
            },
            flow_context={"requires_artifact_delivery": True},
        )
    )

    assert result.status == "completed"
    descriptors = result.output["artifact_descriptors"]
    assert descriptors[0]["role"] == "deliverable"
    assert artifact_paths_from_tool_results([result]) == [result.output["artifact_path"]]


def test_file_search_and_summary_prune_browser_profiles_and_symlinks(tmp_path) -> None:
    safe_match = tmp_path / "cron" / "daily-cron.json"
    safe_match.parent.mkdir()
    safe_match.write_text("{}", encoding="utf-8")
    browser_match = tmp_path / "brave-debug" / "Default" / "secret-cron.json"
    browser_match.parent.mkdir(parents=True)
    browser_match.write_text("{}", encoding="utf-8")
    link_match = tmp_path / "linked-cron.json"
    link_match.symlink_to(safe_match)

    search_handler = _build_file_search_handler(workspace_root=tmp_path, include_principal_workspace=False)
    found = search_handler(ToolInvocation("inv", "file_search", "operator", {"pattern": "cron"}))

    assert found.status == "completed"
    assert found.output["matches"] == [str(safe_match.resolve())]

    summary_handler = _build_workspace_summary_handler(workspace_root=tmp_path, include_principal_workspace=False)
    summary = summary_handler(ToolInvocation("inv", "workspace_summary", "operator", {}))

    assert summary.output["file_count"] == 1
    assert summary.output["sample_files"] == ["cron/daily-cron.json"]


def test_file_search_dedupes_overlapping_structured_roots(tmp_path) -> None:
    safe_match = tmp_path / "artifacts" / "report.xlsx"
    safe_match.parent.mkdir()
    safe_match.write_text("sheet", encoding="utf-8")

    search_handler = _build_file_search_handler(
        workspace_root=tmp_path,
        allowed_roots=[tmp_path, safe_match.parent],
        include_principal_workspace=False,
    )
    found = search_handler(ToolInvocation("inv", "file_search", "operator", {"pattern": "report"}))

    assert found.status == "completed"
    assert found.output["matches"] == [str(safe_match.resolve())]


def test_file_search_keeps_workspace_root_when_extra_allowed_roots_exist(tmp_path) -> None:
    workspace_match = tmp_path / "workspace" / "workspace-report.txt"
    workspace_match.parent.mkdir()
    workspace_match.write_text("workspace", encoding="utf-8")
    extra_root = tmp_path / "external"
    extra_root.mkdir()

    search_handler = _build_file_search_handler(
        workspace_root=workspace_match.parent,
        allowed_roots=[extra_root],
        include_principal_workspace=False,
    )
    found = search_handler(ToolInvocation("inv", "file_search", "operator", {"pattern": "workspace-report"}))

    assert found.status == "completed"
    assert found.output["matches"] == [str(workspace_match.resolve())]
    assert str(workspace_match.parent.resolve()) in found.output["searched_roots"]
    assert str(extra_root.resolve()) in found.output["searched_roots"]


def test_file_search_scopes_requested_allowed_root_and_can_search_contents(tmp_path) -> None:
    workspace_root = tmp_path / "workspace"
    allowed_root = tmp_path / "allowed"
    other_root = tmp_path / "other"
    workspace_root.mkdir()
    allowed_root.mkdir()
    other_root.mkdir()
    target = allowed_root / "notes.txt"
    target.write_text("The launch code is emerald.", encoding="utf-8")
    outside = other_root / "notes.txt"
    outside.write_text("The launch code is crimson.", encoding="utf-8")

    search_handler = _build_file_search_handler(
        workspace_root=workspace_root,
        allowed_roots=[workspace_root, allowed_root],
        include_principal_workspace=False,
    )
    found = search_handler(
        ToolInvocation(
            "inv",
            "file_search",
            "operator",
            {"pattern": "emerald", "root": str(allowed_root), "search_contents": True},
        )
    )

    assert found.status == "completed"
    assert found.output["matches"] == [str(target.resolve())]
    assert found.output["match_details"] == [{"path": str(target.resolve()), "match_type": "content"}]
    assert found.output["searched_roots"] == [str(allowed_root.resolve())]

    rejected = search_handler(
        ToolInvocation(
            "inv",
            "file_search",
            "operator",
            {"pattern": "crimson", "root": str(other_root), "search_contents": True},
        )
    )

    assert rejected.status == "failed"
    assert rejected.output["reason"] == "search_root_outside_allowed_roots"


def test_pdf_create_and_edit_handlers_are_local_artifact_tools(tmp_path) -> None:
    from PIL import Image
    from pypdf import PdfReader

    first_image = tmp_path / "first.png"
    second_image = tmp_path / "second.png"
    Image.new("RGB", (80, 60), "blue").save(first_image)
    Image.new("RGB", (60, 80), "green").save(second_image)

    create_handler = _build_pdf_create_handler(workspace_root=tmp_path, include_principal_workspace=False)
    created_path = tmp_path / "combined.pdf"
    created = create_handler(
        ToolInvocation(
            "inv",
            "pdf_create",
            "operator",
            {
                "output_path": str(created_path),
                "image_paths": [str(first_image), str(second_image)],
                "title": "Birds",
            },
        )
    )

    assert created.status == "completed"
    assert created.output["artifact_path"] == str(created_path)
    assert created.output["source_image_paths"] == [str(first_image), str(second_image)]
    assert created_path.read_bytes().startswith(b"%PDF")
    assert len(PdfReader(str(created_path)).pages) == 2

    remote_image = create_handler(
        ToolInvocation(
            "inv",
            "pdf_create",
            "operator",
            {
                "output_path": str(tmp_path / "remote-image.pdf"),
                "image_paths": ["https://example.test/product.jpg"],
                "title": "Remote",
            },
        )
    )
    assert remote_image.status == "failed"
    assert remote_image.output["reason"] == "artifact_media_inputs_failed"
    assert remote_image.output["image_paths"] == ["https://example.test/product.jpg"]

    text_created_path = tmp_path / "single-text.pdf"
    text_created = create_handler(
        ToolInvocation(
            "inv",
            "pdf_create",
            "operator",
            {
                "output_path": str(text_created_path),
                "text_pages": "A single text page should be accepted.",
                "title": "Text",
            },
        )
    )
    assert text_created.status == "completed"
    assert text_created.output["page_count"] == 1

    edit_handler = _build_pdf_edit_handler(workspace_root=tmp_path, include_principal_workspace=False)
    edited_path = tmp_path / "edited.pdf"
    edited = edit_handler(
        ToolInvocation(
            "inv",
            "pdf_edit",
            "operator",
            {
                "input_path": str(created_path),
                "output_path": str(edited_path),
                "page_numbers": [2, 1],
                "append_text_pages": ["A local text appendix."],
            },
        )
    )

    assert edited.status == "completed"
    assert edited.output["artifact_path"] == str(edited_path)
    assert len(PdfReader(str(edited_path)).pages) == 3

    remote_append = edit_handler(
        ToolInvocation(
            "inv",
            "pdf_edit",
            "operator",
            {
                "input_path": str(created_path),
                "output_path": str(tmp_path / "remote-append.pdf"),
                "append_image_paths": ["https://example.test/product.jpg"],
            },
        )
    )
    assert remote_append.status == "failed"
    assert remote_append.output["reason"] == "artifact_media_inputs_failed"
    assert remote_append.output["image_paths"] == ["https://example.test/product.jpg"]


def test_file_handlers_allow_trusted_filesystem_selector_without_workspace_root(tmp_path) -> None:
    target = tmp_path / "trusted.txt"
    target.write_text("trusted", encoding="utf-8")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    handler = _build_file_read_handler(workspace_root=workspace, include_principal_workspace=False)

    denied = handler(ToolInvocation("inv", "file_read", "operator", {"path": str(target)}))
    assert "outside workspace" in denied.error

    allowed = handler(
        ToolInvocation(
            "inv",
            "file_read",
            "operator",
            {"path": str(target)},
            trusted_filesystem_selectors=(str(tmp_path / "*"),),
        )
    )
    assert allowed.output["content"] == "trusted"


def test_terminal_exec_handler_validates_network_policy_and_backend_results() -> None:
    class Backend:
        def __init__(self, *, descriptor=None, result=None, exc=None):
            self.descriptor = descriptor or TerminalBackendDescriptor("subprocess", (), {})
            self.result = result or TerminalExecutionResult(0, "ok", "")
            self.exc = exc
            self.calls = []

        def describe(self):
            return self.descriptor

        def run(self, command, *, cwd, timeout, policy, shell=None):
            self.calls.append((command, cwd, timeout, policy, shell))
            if self.exc:
                raise self.exc
            return self.result

    handler = _build_terminal_exec_handler(workspace_root=None, terminal_executor_backend=Backend())
    assert handler(ToolInvocation("inv", "terminal_exec", "operator", {})).error == "Missing required argument: command"
    assert handler(ToolInvocation("inv", "terminal_exec", "operator", {"command": "echo hi", "network_mode": "bad"})).output["reason"] == "invalid_network_mode"
    assert handler(ToolInvocation("inv", "terminal_exec", "operator", {"command": "curl https://example.com", "network_mode": "none"})).output["reason"] == "network_denied"

    failed_backend = Backend(result=TerminalExecutionResult(2, "", "bad"))
    failed = _build_terminal_exec_handler(terminal_executor_backend=failed_backend)(
        ToolInvocation("inv", "terminal_exec", "operator", {"command": "false"})
    )
    assert failed.status == "failed"
    assert failed.error == "Command failed with exit code 2"

    ok_backend = Backend(result=TerminalExecutionResult(0, "hello", ""))
    ok = _build_terminal_exec_handler(workspace_root=Path("/tmp"), terminal_executor_backend=ok_backend)(
        ToolInvocation(
            "inv",
            "terminal_exec",
            "operator",
            {"command": "echo hello", "network_mode": "full", "shell": "pwsh", "timeout_seconds": 45},
        )
    )
    assert ok.status == "completed"
    assert ok.output["stdout"] == "hello"
    assert ok.output["timeout_seconds"] == 45
    assert ok_backend.calls[0][4] == "pwsh"


def test_terminal_exec_handler_reports_verified_created_artifacts(tmp_path) -> None:
    artifact = tmp_path / "random_cat_names.xlsx"

    class Backend:
        def describe(self):
            return TerminalBackendDescriptor("subprocess", (), {})

        def run(self, command, *, cwd, timeout, policy, shell=None):
            artifact.write_bytes(b"PK\x03\x04\x14dummy-xlsx")
            return TerminalExecutionResult(0, "created workbook", "")

    result = _build_terminal_exec_handler(
        workspace_root=tmp_path,
        terminal_executor_backend=Backend(),
    )(ToolInvocation("inv", "terminal_exec", "operator", {"command": "create workbook"}))

    assert result.status == "completed"
    assert result.output["artifact_paths"] == [str(artifact.resolve())]


def test_terminal_exec_handler_reports_principal_workspace_artifacts_with_empty_stdout(
    tmp_path: Path,
    monkeypatch,
) -> None:
    from nullion import workspace_storage

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    artifact_root = tmp_path / "principal-artifacts"
    artifact_root.mkdir()
    artifact = artifact_root / "stage_random_media_archive_embedded_report.html"
    monkeypatch.setattr(
        workspace_storage,
        "workspace_storage_roots_for_principal",
        lambda principal_id: SimpleNamespace(artifacts=artifact_root, scratch=tmp_path / "scratch"),
    )

    class Backend:
        def describe(self):
            return TerminalBackendDescriptor("subprocess", (), {})

        def run(self, command, *, cwd, timeout, policy, shell=None):
            artifact.write_text("<html><body>STAGE RANDOM WEBSITE QA</body></html>", encoding="utf-8")
            return TerminalExecutionResult(0, "", "", shell=shell or "sh")

    result = _build_terminal_exec_handler(
        workspace_root=workspace,
        terminal_executor_backend=Backend(),
    )(ToolInvocation("inv", "terminal_exec", "workspace_admin", {"command": "create embedded html"}))

    assert result.status == "completed"
    assert result.output["stdout"] == ""
    assert result.output["artifact_paths"] == [str(artifact.resolve())]


def test_terminal_exec_handler_reports_runtime_home_artifacts_with_empty_stdout(
    tmp_path: Path,
    monkeypatch,
) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    runtime_home = tmp_path / "runtime-home"
    artifact_root = runtime_home / "artifacts"
    artifact_root.mkdir(parents=True)
    artifact = artifact_root / "stage_random_website_qa_self_contained_report.html"
    monkeypatch.setenv("NULLION_HOME", str(runtime_home))
    monkeypatch.delenv("NULLION_DATA_DIR", raising=False)

    class Backend:
        def describe(self):
            return TerminalBackendDescriptor("subprocess", (), {})

        def run(self, command, *, cwd, timeout, policy, shell=None):
            artifact.write_text("<html><body>STAGE RANDOM WEBSITE QA</body></html>", encoding="utf-8")
            return TerminalExecutionResult(0, "", "", shell=shell or "sh")

    result = _build_terminal_exec_handler(
        workspace_root=workspace,
        terminal_executor_backend=Backend(),
    )(ToolInvocation("inv", "terminal_exec", "workspace_admin", {"command": "create runtime html"}))

    assert result.status == "completed"
    assert result.output["stdout"] == ""
    assert result.output["artifact_paths"] == [str(artifact.resolve())]


def test_zip_archives_are_inspected_without_utf8_decoding_or_macos_sidecars(tmp_path) -> None:
    archive_path = tmp_path / "Archive.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("report.html", "<h1>Report</h1>")
        archive.writestr("folder/data.csv", "name,value\nA,1\n")
        archive.writestr("__MACOSX/._report.html", "mac metadata")
        archive.writestr("folder/.DS_Store", "mac metadata")
        archive.writestr("folder/._data.csv", "mac metadata")

    read_handler = _build_file_read_handler(workspace_root=tmp_path, include_principal_workspace=False)
    result = read_handler(ToolInvocation("read", "file_read", "operator", {"path": str(archive_path)}))

    assert result.status == "completed"
    assert result.output["type"] == "archive"
    assert result.output["archive_format"] == "zip"
    assert result.output["entry_count"] == 2
    assert result.output["metadata_entries_ignored"] == 3
    assert [entry["name"] for entry in result.output["entries"]] == ["report.html", "folder/data.csv"]
    assert "__MACOSX" not in result.output["content"]
    assert "._data.csv" not in result.output["content"]


def test_archive_extract_skips_metadata_and_blocks_unsafe_members(tmp_path) -> None:
    archive_path = tmp_path / "Archive.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("report.html", "<h1>Report</h1>")
        archive.writestr("nested/data.csv", "name,value\nA,1\n")
        archive.writestr("__MACOSX/._report.html", "mac metadata")
        archive.writestr("../escape.txt", "nope")

    extract_handler = _build_archive_extract_handler(workspace_root=tmp_path, include_principal_workspace=False)
    result = extract_handler(ToolInvocation("extract", "archive_extract", "operator", {"path": str(archive_path)}))

    output_dir = tmp_path / "Archive_extracted"
    assert result.status == "completed"
    assert result.output["output_dir"] == str(output_dir.resolve())
    assert result.output["file_count"] == 2
    assert result.output["metadata_entries_ignored"] == 1
    assert result.output["unsafe_entries_ignored"] == 1
    assert result.output["manifest_row_count"] == 2
    assert Path(result.output["manifest_csv_path"]).is_file()
    assert Path(result.output["manifest_json_path"]).is_file()
    assert Path(result.output["manifest_xlsx_path"]).is_file()
    assert result.output["manifest_xlsx_path"] in result.output["manifest_artifact_paths"]
    assert result.output["artifact_paths"] == []
    assert result.output["artifact_path"] == ""
    roles_by_path = {item["path"]: item["role"] for item in result.output["artifact_descriptors"]}
    assert roles_by_path[str(archive_path)] == "source"
    assert roles_by_path[result.output["manifest_xlsx_path"]] == "intermediate"
    assert roles_by_path[result.output["manifest_csv_path"]] == "intermediate"
    assert roles_by_path[result.output["manifest_json_path"]] == "intermediate"
    assert result.output["manifest_csv_path"] not in result.output["artifact_paths"]
    assert result.output["manifest_json_path"] not in result.output["artifact_paths"]
    assert not any(str(column).startswith("Part ") for column in result.output["manifest_columns"])
    assert (output_dir / "report.html").read_text(encoding="utf-8") == "<h1>Report</h1>"
    assert (output_dir / "nested" / "data.csv").read_text(encoding="utf-8") == "name,value\nA,1\n"
    assert not (output_dir / "__MACOSX").exists()
    assert not (tmp_path.parent / "escape.txt").exists()

    import openpyxl

    workbook = openpyxl.load_workbook(result.output["manifest_xlsx_path"])
    sheet = workbook.active
    assert sheet.max_row == 3
    assert [sheet.cell(row=row, column=2).value for row in range(2, 4)] == ["report.html", "nested/data.csv"]


def test_archive_extract_writes_requested_manifest_to_artifacts_from_files_source(tmp_path) -> None:
    files_dir = tmp_path / "files"
    artifact_dir = tmp_path / "artifacts"
    files_dir.mkdir()
    artifact_dir.mkdir()
    archive_path = files_dir / "Archive.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("report.html", "<h1>Report</h1>")
        archive.writestr("nested/data.csv", "name,value\nA,1\n")
        archive.writestr("__MACOSX/._report.html", "mac metadata")

    extract_handler = _build_archive_extract_handler(workspace_root=tmp_path, include_principal_workspace=False)
    result = extract_handler(
        ToolInvocation(
            "extract",
            "archive_extract",
            "operator",
            {
                "path": str(archive_path),
                "output_dir": str(files_dir / "Archive_extracted"),
                "manifest_output_path": "stage_archive_manifest_rebuilt.xlsx",
            },
        )
    )

    assert result.status == "completed"
    assert result.output["file_count"] == 2
    assert result.output["manifest_xlsx_path"] == str((artifact_dir / "stage_archive_manifest_rebuilt.xlsx").resolve())
    assert result.output["artifact_paths"] == [result.output["manifest_xlsx_path"]]
    assert result.output["manifest_csv_path"].startswith(str(artifact_dir.resolve()))
    assert result.output["manifest_json_path"].startswith(str(artifact_dir.resolve()))
    assert not any("__MACOSX" in entry["name"] for entry in result.output["entries"])

    import openpyxl

    workbook = openpyxl.load_workbook(result.output["manifest_xlsx_path"])
    sheet = workbook.active
    assert sheet.max_row == 3
    assert [sheet.cell(row=row, column=2).value for row in range(2, 4)] == ["report.html", "nested/data.csv"]


def test_archive_extract_manifest_embeds_images_without_filename_transform_logic(tmp_path) -> None:
    from PIL import Image
    import openpyxl

    archive_path = tmp_path / "Archive.zip"
    image_name = "images/field-alpha-sunrise-product-photo.jpg"
    image_buffer = io.BytesIO()
    Image.new("RGB", (240, 180), color=(32, 96, 160)).save(image_buffer, format="JPEG", quality=95)
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr(image_name, image_buffer.getvalue())
        archive.writestr("notes/readme.txt", "inventory notes")

    extract_handler = _build_archive_extract_handler(workspace_root=tmp_path, include_principal_workspace=False)
    result = extract_handler(ToolInvocation("extract", "archive_extract", "operator", {"path": str(archive_path)}))

    assert result.status == "completed"
    manifest_path = Path(result.output["manifest_xlsx_path"])
    assert manifest_path.is_file()
    assert result.output["artifact_paths"] == []
    assert result.output["manifest_embedded_image_count"] == 1
    assert len(result.output["embedded_images"]) == 1
    assert result.output["source_image_paths"] == result.output["embedded_images"]
    assert "Image" in result.output["manifest_columns"]
    assert not any(str(column).startswith("Part ") for column in result.output["manifest_columns"])

    with zipfile.ZipFile(manifest_path) as workbook_zip:
        assert any(name.startswith("xl/media/") for name in workbook_zip.namelist())

    workbook = openpyxl.load_workbook(manifest_path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    image_row = 2
    assert headers[1] == "Name"
    assert sheet.cell(row=image_row, column=headers.index("Name") + 1).value == image_name
    assert sheet.cell(row=image_row, column=headers.index("Filename") + 1).value == Path(image_name).name
    assert all(not str(header).startswith("Part ") for header in headers)


def test_archive_extract_defaults_manifest_to_artifacts_for_files_source(tmp_path) -> None:
    files_dir = tmp_path / "files"
    artifact_dir = tmp_path / "artifacts"
    files_dir.mkdir()
    artifact_dir.mkdir()
    archive_path = files_dir / "Archive.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("report.html", "<h1>Report</h1>")

    extract_handler = _build_archive_extract_handler(workspace_root=tmp_path, include_principal_workspace=False)
    result = extract_handler(ToolInvocation("extract", "archive_extract", "operator", {"path": str(archive_path)}))

    assert result.status == "completed"
    assert Path(result.output["manifest_xlsx_path"]).parent == artifact_dir.resolve()
    assert Path(result.output["manifest_csv_path"]).parent == artifact_dir.resolve()
    assert Path(result.output["manifest_json_path"]).parent == artifact_dir.resolve()


def test_spreadsheet_create_rejects_partial_archive_manifest_overwrite(tmp_path) -> None:
    artifact_dir = tmp_path / "artifacts"
    files_dir = tmp_path / "files"
    artifact_dir.mkdir()
    files_dir.mkdir()
    archive_path = artifact_dir / "Archive.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("one.txt", "1")
        archive.writestr("two.txt", "2")
        archive.writestr("three.txt", "3")

    extract_handler = _build_archive_extract_handler(workspace_root=tmp_path, include_principal_workspace=False)
    extracted = extract_handler(
        ToolInvocation(
            "extract",
            "archive_extract",
            "operator",
            {"path": str(archive_path), "output_dir": str(files_dir / "Archive_extracted")},
        )
    )
    manifest_path = Path(extracted.output["manifest_xlsx_path"])
    assert extracted.status == "completed"
    assert manifest_path.parent == artifact_dir.resolve()

    spreadsheet_handler = _build_spreadsheet_create_handler(workspace_root=tmp_path, include_principal_workspace=False)
    partial = spreadsheet_handler(
        ToolInvocation(
            "sheet",
            "spreadsheet_create",
            "operator",
            {
                "output_path": str(manifest_path),
                "columns": ["#", "Name", "Media Type", "Bytes", "Extracted Path"],
                "rows": [
                    {"#": 1, "Name": "one.txt", "Media Type": "text/plain", "Bytes": 1, "Extracted Path": "one"},
                    {"#": 2, "Name": "two.txt", "Media Type": "text/plain", "Bytes": 1, "Extracted Path": "two"},
                ],
            },
        )
    )

    assert partial.status == "failed"
    assert partial.output["reason"] == "archive_manifest_row_count_mismatch"
    assert partial.output["expected_rows"] == 3

    import openpyxl

    workbook = openpyxl.load_workbook(manifest_path)
    assert workbook.active.max_row == 4


def test_spreadsheet_create_requires_shipped_openpyxl_for_xlsx(monkeypatch, tmp_path) -> None:
    import builtins

    real_import = builtins.__import__

    def fake_import(name, globals=None, locals=None, fromlist=(), level=0):  # noqa: ANN001
        if name == "openpyxl" or name.startswith("openpyxl."):
            raise ModuleNotFoundError("No module named 'openpyxl'")
        return real_import(name, globals, locals, fromlist, level)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    spreadsheet_handler = _build_spreadsheet_create_handler(workspace_root=tmp_path, include_principal_workspace=False)
    output_path = tmp_path / "manifest.xlsx"
    result = spreadsheet_handler(
        ToolInvocation(
            "sheet",
            "spreadsheet_create",
            "operator",
            {
                "output_path": str(output_path),
                "columns": ["Name", "Bytes", "Formula"],
                "rows": [
                    {"Name": "one.txt", "Bytes": 1, "Formula": "=B2+1"},
                    {"Name": "two.txt", "Bytes": 2, "Formula": "=B3+1"},
                ],
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "missing_dependency"
    assert result.output["dependency_id"] == "openpyxl"


def test_spreadsheet_create_csv_does_not_require_openpyxl(monkeypatch, tmp_path) -> None:
    import builtins

    real_import = builtins.__import__

    def fake_import(name, globals=None, locals=None, fromlist=(), level=0):  # noqa: ANN001
        if name == "openpyxl" or name.startswith("openpyxl."):
            raise ModuleNotFoundError("No module named 'openpyxl'")
        return real_import(name, globals, locals, fromlist, level)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    spreadsheet_handler = _build_spreadsheet_create_handler(workspace_root=tmp_path, include_principal_workspace=False)
    output_path = tmp_path / "manifest.csv"
    result = spreadsheet_handler(
        ToolInvocation(
            "sheet",
            "spreadsheet_create",
            "operator",
            {
                "output_path": str(output_path),
                "columns": ["Name", "Bytes"],
                "rows": [
                    {"Name": "one.txt", "Bytes": 1},
                    {"Name": "two.txt", "Bytes": 2},
                ],
            },
        )
    )

    assert result.status == "completed"
    assert result.output["artifact_path"] == str(output_path)
    assert result.output["artifact_descriptors"][0]["role"] == "deliverable"
    assert output_path.read_text(encoding="utf-8").splitlines() == [
        "Name,Bytes",
        "one.txt,1",
        "two.txt,2",
    ]


def test_archive_create_skips_metadata_and_delivers_zip_artifact(tmp_path) -> None:
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    (source_dir / "report.html").write_text("<h1>Report</h1>", encoding="utf-8")
    (source_dir / ".DS_Store").write_text("metadata", encoding="utf-8")
    macosx = source_dir / "__MACOSX"
    macosx.mkdir()
    (macosx / "._report.html").write_text("metadata", encoding="utf-8")

    create_handler = _build_archive_create_handler(workspace_root=tmp_path, include_principal_workspace=False)
    output_path = tmp_path / "bundle.zip"
    result = create_handler(
        ToolInvocation(
            "create",
            "archive_create",
            "operator",
            {"output_path": str(output_path), "source_dir": str(source_dir)},
        )
    )

    assert result.status == "completed"
    assert result.output["artifact_path"] == str(output_path)
    assert result.output["artifact_paths"] == [str(output_path)]
    assert result.output["artifact_descriptors"][0]["role"] == "deliverable"
    assert result.output["entry_count"] == 1
    with zipfile.ZipFile(output_path) as archive:
        assert archive.namelist() == ["report.html"]


def test_tar_and_single_file_compression_archives_are_supported(tmp_path) -> None:
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    (source_dir / "report.html").write_text("<h1>Report</h1>", encoding="utf-8")
    (source_dir / ".DS_Store").write_text("metadata", encoding="utf-8")

    create_handler = _build_archive_create_handler(workspace_root=tmp_path, include_principal_workspace=False)
    read_handler = _build_file_read_handler(workspace_root=tmp_path, include_principal_workspace=False)
    extract_handler = _build_archive_extract_handler(workspace_root=tmp_path, include_principal_workspace=False)

    tar_path = tmp_path / "bundle.tar.gz"
    created_tar = create_handler(
        ToolInvocation("create-tar", "archive_create", "operator", {"output_path": str(tar_path), "source_dir": str(source_dir)})
    )
    assert created_tar.status == "completed"
    assert created_tar.output["archive_format"] == "tar.gz"

    inspected_tar = read_handler(ToolInvocation("read-tar", "file_read", "operator", {"path": str(tar_path)}))
    assert inspected_tar.status == "completed"
    assert inspected_tar.output["archive_format"] == "tar.gz"
    assert inspected_tar.output["entry_count"] == 1
    assert inspected_tar.output["metadata_entries_ignored"] == 0
    assert [entry["name"] for entry in inspected_tar.output["entries"]] == ["report.html"]

    extracted_tar = extract_handler(
        ToolInvocation("extract-tar", "archive_extract", "operator", {"path": str(tar_path), "output_dir": str(tmp_path / "tar_out")})
    )
    assert extracted_tar.status == "completed"
    assert (tmp_path / "tar_out" / "report.html").read_text(encoding="utf-8") == "<h1>Report</h1>"

    text_path = tmp_path / "notes.txt"
    text_path.write_text("compressed text", encoding="utf-8")
    gzip_path = tmp_path / "notes.txt.gz"
    created_gzip = create_handler(
        ToolInvocation("create-gzip", "archive_create", "operator", {"output_path": str(gzip_path), "source_paths": [str(text_path)]})
    )
    assert created_gzip.status == "completed"
    assert created_gzip.output["archive_format"] == "gzip"
    assert gzip.decompress(gzip_path.read_bytes()) == b"compressed text"

    inspected_gzip = read_handler(ToolInvocation("read-gzip", "file_read", "operator", {"path": str(gzip_path)}))
    assert inspected_gzip.status == "completed"
    assert inspected_gzip.output["type"] == "compressed_file"
    assert inspected_gzip.output["entries"][0]["name"] == "notes.txt"

    extracted_gzip = extract_handler(
        ToolInvocation("extract-gzip", "archive_extract", "operator", {"path": str(gzip_path), "output_dir": str(tmp_path / "gzip_out")})
    )
    assert extracted_gzip.status == "completed"
    assert (tmp_path / "gzip_out" / "notes.txt").read_text(encoding="utf-8") == "compressed text"


def test_file_download_saves_remote_bytes_without_terminal_exec(monkeypatch, tmp_path) -> None:
    archive_buffer = io.BytesIO()
    with zipfile.ZipFile(archive_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("data.txt", "zipdata")
    archive_bytes = archive_buffer.getvalue()

    class Headers(dict):
        def get_content_type(self):
            return self.get("Content-Type", "application/octet-stream")

    class Response:
        headers = Headers({"Content-Type": "application/zip", "Content-Length": str(len(archive_bytes))})

        def __init__(self) -> None:
            self._body = io.BytesIO(archive_bytes)

        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, tb):
            return False

        def read(self, size=-1):
            return self._body.read(size)

        def geturl(self):
            return "https://93.184.216.34/archive.zip"

    class Opener:
        def open(self, request, timeout):  # noqa: ANN001
            assert request.full_url == "https://93.184.216.34/archive.zip"
            assert timeout == 10
            return Response()

    monkeypatch.setattr("nullion.tools._build_web_fetch_opener", lambda: Opener())

    handler = _build_file_download_handler(workspace_root=tmp_path, include_principal_workspace=False)
    output_path = tmp_path / "Archive.zip"
    result = handler(
        ToolInvocation(
            "download",
            "file_download",
            "operator",
            {"url": "https://93.184.216.34/archive.zip", "output_path": str(output_path)},
        )
    )

    assert result.status == "completed"
    assert zipfile.is_zipfile(output_path)
    assert result.output["path"] == str(output_path)
    assert result.output["artifact_paths"] == [str(output_path)]
    assert result.output["artifact_descriptors"][0]["role"] == "source"
    assert result.output["content_type"] == "application/zip"


def test_file_download_retries_truncated_archive(monkeypatch, tmp_path) -> None:
    archive_buffer = io.BytesIO()
    with zipfile.ZipFile(archive_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("manifest.csv", "name,count\nhtml,2\n")
    archive_bytes = archive_buffer.getvalue()
    response_bodies = [archive_bytes[:32], archive_bytes]

    class Headers(dict):
        def get_content_type(self):
            return self.get("Content-Type", "application/octet-stream")

    class Response:
        headers = Headers({"Content-Type": "application/zip"})

        def __init__(self, body: bytes) -> None:
            self._body = io.BytesIO(body)

        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, tb):
            return False

        def read(self, size=-1):
            return self._body.read(size)

        def geturl(self):
            return "https://93.184.216.34/archive.zip"

    class Opener:
        calls = 0

        def open(self, request, timeout):  # noqa: ANN001
            assert request.full_url == "https://93.184.216.34/archive.zip"
            body = response_bodies[self.calls]
            self.calls += 1
            return Response(body)

    opener = Opener()
    monkeypatch.setattr("nullion.tools._build_web_fetch_opener", lambda: opener)

    handler = _build_file_download_handler(workspace_root=tmp_path, include_principal_workspace=False)
    output_path = tmp_path / "Archive.zip"
    result = handler(
        ToolInvocation(
            "download-retry",
            "file_download",
            "operator",
            {"url": "https://93.184.216.34/archive.zip", "output_path": str(output_path)},
        )
    )

    assert result.status == "completed"
    assert opener.calls == 2
    assert zipfile.is_zipfile(output_path)
    assert output_path.read_bytes() == archive_bytes


def test_file_download_returns_timeout_and_removes_partial_file(monkeypatch, tmp_path) -> None:
    class Headers(dict):
        def get_content_type(self):
            return self.get("Content-Type", "application/octet-stream")

    class Response:
        headers = Headers({"Content-Type": "application/zip"})

        def __init__(self) -> None:
            self._reads = 0

        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, tb):
            return False

        def read(self, size=-1):  # noqa: ARG002
            self._reads += 1
            return b"partial" if self._reads == 1 else b""

        def geturl(self):
            return "https://93.184.216.34/archive.zip"

    class Opener:
        def open(self, request, timeout):  # noqa: ANN001
            assert request.full_url == "https://93.184.216.34/archive.zip"
            assert timeout == 2
            return Response()

    perf_values = iter([0.0, 0.0, 3.0])
    monkeypatch.setattr("nullion.tools.perf_counter", lambda: next(perf_values))
    monkeypatch.setattr("nullion.tools._build_web_fetch_opener", lambda: Opener())

    handler = _build_file_download_handler(workspace_root=tmp_path, include_principal_workspace=False)
    output_path = tmp_path / "Archive.zip"
    result = handler(
        ToolInvocation(
            "download-timeout",
            "file_download",
            "operator",
            {
                "url": "https://93.184.216.34/archive.zip",
                "output_path": str(output_path),
                "timeout_seconds": 2,
            },
        )
    )

    assert result.status == "failed"
    assert result.output["reason"] == "download_timeout"
    assert result.output["bytes_read"] == len(b"partial")
    assert not output_path.exists()


def test_terminal_exec_blocks_unapproved_filesystem_mutations_outside_workspace(tmp_path: Path) -> None:
    class Backend:
        def __init__(self) -> None:
            self.calls = []

        def describe(self):
            return TerminalBackendDescriptor("subprocess", (), {})

        def run(self, command, *, cwd, timeout, policy, shell=None):  # noqa: ANN001
            self.calls.append(command)
            return TerminalExecutionResult(0, "mutated", "")

    backend = Backend()
    handler = _build_terminal_exec_handler(workspace_root=tmp_path, terminal_executor_backend=backend)

    result = handler(
        ToolInvocation(
            "inv",
            "terminal_exec",
            "operator",
            {"command": "ln -sfn /opt/homebrew/bin/python3.11 /usr/local/bin/python3"},
        )
    )

    assert result.status == "denied"
    assert result.output["reason"] == "filesystem_mutation_approval_required"
    assert Path(result.output["target"]).resolve() == Path("/usr/local/bin/python3").resolve()
    assert backend.calls == []


def test_terminal_exec_allows_read_only_diagnostics_and_trusted_mutations(tmp_path: Path) -> None:
    class Backend:
        def __init__(self) -> None:
            self.calls = []

        def describe(self):
            return TerminalBackendDescriptor("subprocess", (), {})

        def run(self, command, *, cwd, timeout, policy, shell=None):  # noqa: ANN001
            self.calls.append(command)
            return TerminalExecutionResult(0, "ok", "")

    backend = Backend()
    handler = _build_terminal_exec_handler(workspace_root=tmp_path, terminal_executor_backend=backend)

    read_only = handler(
        ToolInvocation(
            "read",
            "terminal_exec",
            "operator",
            {"command": "which python3 && python3 --version"},
        )
    )
    read_only_with_dev_null = handler(
        ToolInvocation(
            "read-dev-null",
            "terminal_exec",
            "operator",
            {"command": "command -v python3 2> /dev/null && python3 --version"},
        )
    )
    trusted_target = tmp_path / "python-version.txt"
    trusted_write = handler(
        ToolInvocation(
            "write",
            "terminal_exec",
            "operator",
            {"command": f"python3 --version > {trusted_target}"},
        )
    )
    approved_symlink = handler(
        ToolInvocation(
            "approved",
            "terminal_exec",
            "operator",
            {"command": "ln -sfn /opt/homebrew/bin/python3.11 /usr/local/bin/python3"},
            trusted_filesystem_selectors=("/usr/local/bin/*",),
        )
    )

    assert read_only.status == "completed"
    assert read_only_with_dev_null.status == "completed"
    assert trusted_write.status == "completed"
    assert approved_symlink.status == "completed"
    assert backend.calls == [
        "which python3 && python3 --version",
        "command -v python3 2> /dev/null && python3 --version",
        f"python3 --version > {trusted_target}",
        "ln -sfn /opt/homebrew/bin/python3.11 /usr/local/bin/python3",
    ]


def test_terminal_exec_blocks_unknown_workspace_directory_creation(tmp_path: Path, monkeypatch) -> None:
    class Backend:
        def __init__(self) -> None:
            self.calls = []

        def describe(self):
            return TerminalBackendDescriptor("subprocess", (), {})

        def run(self, command, *, cwd, timeout, policy):  # noqa: ANN001
            self.calls.append(command)
            return TerminalExecutionResult(0, "created", "")

    storage_root = tmp_path / "workspaces"
    (storage_root / "workspace_nate").mkdir(parents=True)
    monkeypatch.setenv("NULLION_WORKSPACE_STORAGE_ROOT", str(storage_root))
    backend = Backend()
    handler = _build_terminal_exec_handler(workspace_root=tmp_path, terminal_executor_backend=backend)

    denied = handler(
        ToolInvocation(
            "inv",
            "terminal_exec",
            "operator",
            {"command": f"mkdir -p {storage_root / 'workspace_nathan' / 'files'}"},
        )
    )
    allowed = handler(
        ToolInvocation(
            "inv2",
            "terminal_exec",
            "operator",
            {"command": f"ls {storage_root / 'workspace_nate'}"},
        )
    )

    assert denied.status == "failed"
    assert denied.output["reason"] == "unknown_workspace_denied"
    assert denied.output["workspace_id"] == "workspace_nathan"
    assert backend.calls == [f"ls {storage_root / 'workspace_nate'}"]
    assert allowed.status == "completed"


def test_terminal_exec_handler_passes_shell_timeout_and_discovers_artifacts(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    artifact = workspace / "report.pdf"

    class Backend:
        def __init__(self) -> None:
            self.calls = []

        def describe(self):
            return TerminalBackendDescriptor("subprocess", (), {})

        def run(self, command, *, cwd, timeout, policy, shell=None):
            self.calls.append((command, cwd, timeout, policy, shell))
            artifact.write_bytes(b"%PDF-1.4\n")
            return TerminalExecutionResult(
                exit_code=0,
                stdout="done",
                stderr="",
                shell=shell or "unknown",
                argv=("fake-shell", "-lc", command),
            )

    backend = Backend()
    result = _build_terminal_exec_handler(workspace_root=workspace, terminal_executor_backend=backend)(
        ToolInvocation(
            "inv",
            "terminal_exec",
            "operator",
            {"command": "make report", "shell": "zsh", "timeout_seconds": 3},
        )
    )

    assert result.status == "completed"
    assert backend.calls[0][:3] == ("make report", str(workspace.resolve()), 3)
    assert backend.calls[0][4] == "zsh"
    assert result.output["shell"] == "zsh"
    assert result.output["timeout_seconds"] == 3
    assert result.output["artifact_paths"] == [str(artifact.resolve())]


def test_terminal_exec_handler_falls_back_when_workspace_root_is_missing() -> None:
    class Backend:
        def __init__(self) -> None:
            self.calls = []

        def describe(self):
            return TerminalBackendDescriptor("subprocess", (), {})

        def run(self, command, *, cwd, timeout, policy, shell=None):
            self.calls.append((command, cwd, timeout, policy, shell))
            return TerminalExecutionResult(0, "found", "")

    backend = Backend()
    missing_workspace = Path("/definitely/missing/nullion/workspace")
    result = _build_terminal_exec_handler(
        workspace_root=missing_workspace,
        terminal_executor_backend=backend,
    )(ToolInvocation("inv", "terminal_exec", "operator", {"command": "echo found"}))

    assert result.status == "completed"
    assert result.output["stdout"] == "found"
    assert backend.calls[0][1] is None


def test_terminal_exec_resolves_windows_shell_without_unix_default(monkeypatch) -> None:
    monkeypatch.setattr(
        "nullion.tools.shutil.which",
        lambda executable: f"C:/Windows/System32/{executable}" if executable == "powershell.exe" else None,
    )

    invocation = _resolve_windows_terminal_shell("Write-Output hello", "auto")

    assert invocation.shell == "powershell"
    assert invocation.argv[0].endswith("powershell.exe")
    assert "/bin/sh" not in invocation.argv
    assert invocation.argv[-1] == "Write-Output hello"


def test_terminal_exec_uses_structured_shell_and_bounded_timeout() -> None:
    class Backend:
        def __init__(self) -> None:
            self.calls = []

        def describe(self):
            return TerminalBackendDescriptor("subprocess", (), {})

        def run(self, command, *, cwd, timeout, policy, shell=None):
            self.calls.append((command, cwd, timeout, policy, shell))
            return TerminalExecutionResult(0, "ok", "", shell=shell or "auto")

    backend = Backend()
    result = _build_terminal_exec_handler(terminal_executor_backend=backend)(
        ToolInvocation(
            "inv",
            "terminal_exec",
            "operator",
            {"command": "Write-Output hello", "shell": "powershell", "timeout_seconds": 999},
        )
    )

    assert result.status == "completed"
    assert result.output["shell"] == "powershell"
    assert result.output["timeout_seconds"] == 300
    assert backend.calls[0][2] == 300
    assert backend.calls[0][4] == "powershell"


def test_weather_location_resolution_is_cached_for_same_location_text() -> None:
    with tools_module._WEATHER_LOCATION_CACHE_LOCK:
        tools_module._WEATHER_LOCATION_CACHE.clear()
    calls: list[str] = []

    def fake_json_get(url: str, timeout_seconds: int) -> object:
        calls.append(url)
        return {
            "results": [
                {
                    "name": "New York",
                    "admin1": "New York",
                    "country": "United States",
                    "latitude": 40.7128,
                    "longitude": -74.006,
                    "timezone": "America/New_York",
                }
            ]
        }

    first = tools_module._resolve_open_meteo_location(
        {"location_text": "New York, NY"},
        json_get=fake_json_get,
    )
    second = tools_module._resolve_open_meteo_location(
        {"location_text": "  New York,   NY  "},
        json_get=fake_json_get,
    )

    assert first == second
    assert first["latitude"] == 40.7128
    assert len(calls) == 1
